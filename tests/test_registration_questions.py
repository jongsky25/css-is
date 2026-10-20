"""Custom registration questions (PR-A) — the pure rules, the storage layer and the
routes (manage card + public form).

Three layers, mirroring test_registration_dedup:
  * reg_questions.py — pure rules, no database.
  * storage on an isolated temp DB — questions, answers, ordering, archiving.
  * routes on the shared test DB — the organizer's manage routes and the public
    registration flow with answers.
"""
from __future__ import annotations
import tzutil as _test_tz
import json
import re
from datetime import date, timedelta
from urllib.parse import urlencode

import reg_questions as rq
import storage as st

_FUTURE_DATE = (_test_tz.today_local() + timedelta(days=365)).isoformat()


# --------------------------------------------------------------------------
# reg_questions — pure
# --------------------------------------------------------------------------

def _q(qid, kind, options=(), required=False, label="Q"):
    return {"id": qid, "kind": kind, "options": list(options),
            "required": required, "label": label}


def test_parse_options_strips_dedupes_caps():
    assert rq.parse_options(" A \n\n b\nB\n  \nA\nc ") == ["A", "b", "c"]
    many = "\n".join(f"opt{i}" for i in range(50))
    assert len(rq.parse_options(many)) == rq.MAX_OPTIONS
    assert len(rq.parse_options("x" * 500)[0]) == rq.MAX_OPTION_LEN


def test_clean_question_errors_and_shape():
    assert rq.clean_question({"label": "  ", "kind": "text"}) == (None, "label")
    assert rq.clean_question({"label": "L", "kind": "nope"}) == (None, "kind")
    assert rq.clean_question({"label": "L", "kind": "single", "options": "only"}) == (None, "options")
    q, err = rq.clean_question({"label": " L ", "kind": "text", "options": "a\nb",
                                "required": "on", "help_text": " h "})
    assert err is None
    assert q == {"label": "L", "help_text": "h", "kind": "text", "options": [],
                 "required": True}
    q, _ = rq.clean_question({"label": "L", "kind": "multi", "options": "a\nb"})
    assert q["options"] == ["a", "b"] and q["required"] is False


def test_collect_required_blank_names_the_question():
    qs = [_q(1, "text", required=True, label="Dietary")]
    answers, errors = rq.collect_answers(qs, [("name", "x")])
    assert answers == {} and len(errors) == 1 and "Dietary" in errors[0]


def test_collect_forged_single_is_blank():
    qs = [_q(1, "single", ["A", "B"], required=True, label="Pick")]
    answers, errors = rq.collect_answers(qs, [("q_1", "<script>")])
    assert answers == {} and errors
    answers, errors = rq.collect_answers([_q(1, "single", ["A", "B"])], [("q_1", "<script>")])
    assert answers == {} and errors == []
    assert rq.collect_answers([_q(1, "single", ["A", "B"])], [("q_1", "B")])[0] == {1: "B"}


def test_collect_multi_keeps_valid_in_option_order_as_json():
    qs = [_q(5, "multi", ["None", "Vegan", "Halal"])]
    answers, errors = rq.collect_answers(
        qs, [("q_5", "Halal"), ("q_5", "bogus"), ("q_5", "Vegan")])
    assert errors == []
    assert json.loads(answers[5]) == ["Vegan", "Halal"]
    assert rq.display_value(qs[0], answers[5]) == "Vegan, Halal"
    # Only forged values -> blank (and an error when required).
    qs[0]["required"] = True
    assert rq.collect_answers(qs, [("q_5", "bogus")])[1]


def test_collect_yesno_date_time_validation():
    qs = [_q(1, "yesno"), _q(2, "date"), _q(3, "time")]
    ok, _ = rq.collect_answers(qs, [("q_1", "Yes"), ("q_2", "2026-12-31"), ("q_3", "09:30")])
    assert ok == {1: "Yes", 2: "2026-12-31", 3: "09:30"}
    bad, errs = rq.collect_answers(qs, [("q_1", "yes"), ("q_2", "2026-13-40"), ("q_3", "25:99")])
    assert bad == {} and errs == []
    bad, _ = rq.collect_answers(qs, [("q_2", "12/31/2026"), ("q_3", "9:30")])
    assert bad == {}


def test_collect_text_truncated():
    qs = [_q(1, "text"), _q(2, "longtext")]
    answers, _ = rq.collect_answers(qs, [("q_1", "x" * 500), ("q_2", "y" * 5000)])
    assert len(answers[1]) == 200 and len(answers[2]) == rq.MAX_ANSWER


def test_display_value_edges():
    q = _q(1, "multi", ["a"])
    assert rq.display_value(q, None) == "" and rq.display_value(q, "") == ""
    assert rq.display_value(q, "not json") == "not json"
    assert rq.display_value(_q(2, "text"), "hi") == "hi"


def test_presets_all_clean():
    for key in rq.PRESETS:
        got = rq.preset_questions(key)
        assert len(got) == len(rq.PRESETS[key])
    assert rq.preset_questions("nope") == []
    assert set(rq.PRESETS) == set(rq.PRESET_LABELS)


# --------------------------------------------------------------------------
# storage — isolated temp DB
# --------------------------------------------------------------------------

def _fresh(tmp_path):
    conn = st.connect(str(tmp_path / "rq.db"))
    st.init_db(conn)
    return conn


def _event(conn, event_date="2099-01-01"):
    f2 = st.create_form_instance(conn, 2, "F2", event_date, topic="Event review")
    ev = st.create_event(conn, "Forum", event_date,
                         form_instance_ids=[f2], created_by="manager@test.local")
    return ev["id"]


def _add(conn, eid, label="Q", kind="text", options="", required=False):
    q, err = rq.clean_question({"label": label, "kind": kind, "options": options,
                                "required": "on" if required else ""})
    assert err is None, err
    return st.add_event_question(conn, eid, q, created_by="t")


def test_add_and_list_in_order(tmp_path):
    conn = _fresh(tmp_path)
    eid = _event(conn)
    a = _add(conn, eid, "First")
    b = _add(conn, eid, "Second", "single", "x\ny", required=True)
    got = st.list_event_questions(conn, eid)
    assert [q["id"] for q in got] == [a, b]
    assert got[1]["options"] == ["x", "y"] and got[1]["required"] is True
    assert got[0]["options"] == [] and got[0]["n_answers"] == 0


def test_add_stops_at_max(tmp_path):
    conn = _fresh(tmp_path)
    eid = _event(conn)
    for i in range(rq.MAX_QUESTIONS):
        assert _add(conn, eid, f"Q{i}") is not None
    q, _ = rq.clean_question({"label": "one too many", "kind": "text"})
    assert st.add_event_question(conn, eid, q) is None
    assert len(st.list_event_questions(conn, eid)) == rq.MAX_QUESTIONS


def test_move_swaps_and_noop_at_ends(tmp_path):
    conn = _fresh(tmp_path)
    eid = _event(conn)
    a, b, c = (_add(conn, eid, n) for n in "abc")
    ids = lambda: [q["id"] for q in st.list_event_questions(conn, eid)]  # noqa: E731
    st.move_event_question(conn, eid, a, "up")        # already first
    assert ids() == [a, b, c]
    st.move_event_question(conn, eid, a, "down")
    assert ids() == [b, a, c]
    st.move_event_question(conn, eid, c, "down")      # already last
    assert ids() == [b, a, c]
    st.move_event_question(conn, eid, c, "up")
    assert ids() == [b, c, a]
    st.move_event_question(conn, eid, c, "sideways")
    assert ids() == [b, c, a]


def test_remove_deletes_or_archives(tmp_path):
    conn = _fresh(tmp_path)
    eid = _event(conn)
    unanswered = _add(conn, eid, "Unanswered")
    answered = _add(conn, eid, "Answered")
    aid = st.add_attendee(conn, eid, name="Ana", email="a@x.com", consent_text="x")
    st.save_attendee_answers(conn, aid, {answered: "hello"})
    assert st.remove_event_question(conn, unanswered) == "deleted"
    assert st.get_event_question(conn, eid, unanswered) is None
    assert st.remove_event_question(conn, answered) == "archived"
    assert st.list_event_questions(conn, eid) == []
    arch = st.list_event_questions(conn, eid, include_archived=True)
    assert [q["id"] for q in arch] == [answered] and arch[0]["archived_at"]
    assert st.get_attendee_answers(conn, aid) == {answered: "hello"}   # kept


def test_update_ignores_kind_change_once_answered(tmp_path):
    conn = _fresh(tmp_path)
    eid = _event(conn)
    qid = _add(conn, eid, "Pick", "single", "a\nb")
    new, _ = rq.clean_question({"label": "Pick again", "kind": "text", "required": "on"})
    st.update_event_question(conn, qid, new)                      # unanswered -> changes
    got = st.get_event_question(conn, eid, qid)
    assert got["kind"] == "text" and got["label"] == "Pick again" and got["options"] == []
    aid = st.add_attendee(conn, eid, name="Ana", email="a@x.com", consent_text="x")
    st.save_attendee_answers(conn, aid, {qid: "v"})
    again, _ = rq.clean_question({"label": "Renamed", "kind": "multi", "options": "p\nq"})
    st.update_event_question(conn, qid, again)                    # answered -> kind kept
    got = st.get_event_question(conn, eid, qid)
    assert got["kind"] == "text" and got["label"] == "Renamed" and got["options"] == []


def test_save_answers_upserts_and_keeps_absent(tmp_path):
    conn = _fresh(tmp_path)
    eid = _event(conn)
    q1, q2 = _add(conn, eid, "one"), _add(conn, eid, "two")
    aid = st.add_attendee(conn, eid, name="Ana", email="a@x.com", consent_text="x")
    st.save_attendee_answers(conn, aid, {q1: "A", q2: "B"})
    st.save_attendee_answers(conn, aid, {q1: "A2"})               # q2 absent -> kept
    assert st.get_attendee_answers(conn, aid) == {q1: "A2", q2: "B"}
    assert st.event_answers(conn, eid) == {aid: {q1: "A2", q2: "B"}}
    other = _event(conn)
    assert st.event_answers(conn, other) == {}


def test_get_event_question_is_event_scoped(tmp_path):
    conn = _fresh(tmp_path)
    ea, eb = _event(conn), _event(conn)
    qa = _add(conn, ea, "mine")
    assert st.get_event_question(conn, ea, qa)["label"] == "mine"
    assert st.get_event_question(conn, eb, qa) is None


# --------------------------------------------------------------------------
# routes — shared test DB + the manager client
# --------------------------------------------------------------------------

def _make_event(manager, db_path):
    conn = st.connect(db_path)
    rp = conn.execute("SELECT id FROM resource_person WHERE is_internal=1 LIMIT 1").fetchone()
    count = conn.execute("SELECT COUNT(*) FROM event").fetchone()[0]
    # Distinct dates represent distinct activities; duplicate protection stays active.
    from datetime import date, timedelta
    event_day = (date.fromisoformat(_FUTURE_DATE) + timedelta(days=count)).isoformat()
    conn.close()
    r = manager.post("/create", data={
        "event_type": "bureau", "activity_name": "Questions Forum",
        "event_date": event_day, "venue": "Hall A", "qr_mode": "combined",
        "rp_id": str(rp[0]), "topic": "Keynote",
    }, follow_redirects=True)
    assert r.status_code == 200
    return re.search(r"/e/([0-9a-f]{12})", r.text).group(1)


def _questions(db_path, tok):
    conn = st.connect(db_path)
    try:
        ev = st.get_event(conn, tok)
        return st.list_event_questions(conn, ev["id"], include_archived=True)
    finally:
        conn.close()


def _add_q(manager, tok, **fields):
    data = {"label": "Q", "kind": "text"}
    data.update(fields)
    return manager.post(f"/activities/{tok}/questions", data=data, follow_redirects=False)


def _reg(**kw):
    d = {"name": "Rey Dela Cruz", "email": "rey@example.com", "sex": "M",
         "affiliation": "DOH", "consent": "on"}
    d.update(kw)
    return d


def _post_pairs(c, url, pairs):
    """POST a form whose keys repeat (a checkbox group), as a browser sends it —
    the test client's `data=` only takes a dict."""
    return c.post(url, content=urlencode(pairs),
                  headers={"content-type": "application/x-www-form-urlencoded"},
                  follow_redirects=False)


def _attendee_answers(db_path, tok, email):
    conn = st.connect(db_path)
    try:
        ev = st.get_event(conn, tok)
        rows = conn.execute("SELECT id FROM attendee WHERE event_id=? AND LOWER(email)=? "
                            "ORDER BY id", (ev["id"], email.lower())).fetchall()
        return [st.get_attendee_answers(conn, r[0]) for r in rows]
    finally:
        conn.close()


def _n_attendees(db_path, tok):
    conn = st.connect(db_path)
    try:
        ev = st.get_event(conn, tok)
        return conn.execute("SELECT COUNT(*) FROM attendee WHERE event_id=?",
                            (ev["id"],)).fetchone()[0]
    finally:
        conn.close()


def test_manager_adds_question_shown_on_both_pages(client, manager, db_path):
    tok = _make_event(manager, db_path)
    r = _add_q(manager, tok, label="Dietary needs")
    assert r.status_code == 303 and r.headers["location"].endswith("?qok=1#questions")
    assert "Dietary needs" in manager.get(f"/activities/{tok}/manage").text
    page = client.get(f"/e/{tok}")
    assert page.status_code == 200 and "Dietary needs" in page.text


def test_food_preset_adds_two(manager, db_path):
    tok = _make_event(manager, db_path)
    r = manager.post(f"/activities/{tok}/questions/preset", data={"preset": "food"},
                     follow_redirects=False)
    assert r.status_code == 303 and "qok=1" in r.headers["location"]
    assert len(_questions(db_path, tok)) == 2
    # An unknown preset is refused and adds nothing.
    r = manager.post(f"/activities/{tok}/questions/preset", data={"preset": "zzz"},
                     follow_redirects=False)
    assert "qerr=kind" in r.headers["location"] and len(_questions(db_path, tok)) == 2
    # The food button is gone once its first question exists.
    assert 'name="preset" value="food"' not in manager.get(f"/activities/{tok}/manage").text


def test_bad_question_forms_redirect_with_code(manager, db_path):
    tok = _make_event(manager, db_path)
    assert "qerr=label" in _add_q(manager, tok, label="").headers["location"]
    assert "qerr=kind" in _add_q(manager, tok, kind="zzz").headers["location"]
    assert "qerr=options" in _add_q(manager, tok, kind="single", options="one").headers["location"]
    assert _questions(db_path, tok) == []


def test_viewer_cannot_add_question(viewer, manager, db_path):
    tok = _make_event(manager, db_path)
    r = _add_q(viewer, tok, label="Sneaky")
    assert r.status_code == 303 and r.headers["location"] == f"/e/{tok}"
    assert _questions(db_path, tok) == []


def test_question_id_from_other_event_is_notfound(manager, db_path):
    tok_a, tok_b = _make_event(manager, db_path), _make_event(manager, db_path)
    _add_q(manager, tok_b, label="B's question")
    qb = _questions(db_path, tok_b)[0]["id"]
    for path, data in ((f"/activities/{tok_a}/questions/{qb}", {"label": "Hijack", "kind": "text"}),
                       (f"/activities/{tok_a}/questions/{qb}/remove", {}),
                       (f"/activities/{tok_a}/questions/{qb}/move", {"direction": "down"})):
        r = manager.post(path, data=data, follow_redirects=False)
        assert r.status_code == 303 and "qerr=notfound" in r.headers["location"]
    got = _questions(db_path, tok_b)
    assert len(got) == 1 and got[0]["label"] == "B's question" and not got[0]["archived_at"]


def test_edit_remove_and_locked_kind_via_routes(client, manager, db_path):
    tok = _make_event(manager, db_path)
    _add_q(manager, tok, label="Pick", kind="single", options="a\nb")
    qid = _questions(db_path, tok)[0]["id"]
    client.post(f"/e/{tok}", data=_reg(**{f"q_{qid}": "a"}))
    # Answered: the kind select is disabled in the form (not posted); the route pins it.
    r = manager.post(f"/activities/{tok}/questions/{qid}",
                     data={"label": "Pick one", "options": "a\nb\nc"}, follow_redirects=False)
    assert "qok=1" in r.headers["location"]
    q = _questions(db_path, tok)[0]
    assert q["label"] == "Pick one" and q["kind"] == "single" and q["options"] == ["a", "b", "c"]
    assert "Type is locked" in manager.get(f"/activities/{tok}/manage").text
    # Removing an answered question archives it: gone from the form, answers kept.
    manager.post(f"/activities/{tok}/questions/{qid}/remove")
    assert _questions(db_path, tok)[0]["archived_at"]
    assert f'name="q_{qid}"' not in client.get(f"/e/{tok}").text
    assert _attendee_answers(db_path, tok, "rey@example.com") == [{qid: "a"}]


def test_move_route_reorders(manager, db_path):
    tok = _make_event(manager, db_path)
    _add_q(manager, tok, label="First")
    _add_q(manager, tok, label="Second")
    first = _questions(db_path, tok)[0]["id"]
    manager.post(f"/activities/{tok}/questions/{first}/move", data={"direction": "down"})
    assert [q["label"] for q in _questions(db_path, tok)] == ["Second", "First"]


def test_required_question_unanswered_blocks_registration(client, manager, db_path):
    tok = _make_event(manager, db_path)
    _add_q(manager, tok, label="Shirt size please", required="on")
    before = _n_attendees(db_path, tok)
    r = client.post(f"/e/{tok}", data=_reg(), follow_redirects=False)
    assert r.status_code == 200 and "Please answer" in r.text and "Shirt size please" in r.text
    assert _n_attendees(db_path, tok) == before


def test_registration_stores_answers_including_multi(client, manager, db_path):
    tok = _make_event(manager, db_path)
    manager.post(f"/activities/{tok}/questions/preset", data={"preset": "food"})
    qs = {q["label"]: q["id"] for q in _questions(db_path, tok)}
    multi, note = qs["Food restrictions"], qs[
        "If you have a food allergy or other restriction, please specify"]
    # A list of tuples so the checkbox key repeats, as a browser sends it.
    r = _post_pairs(client, f"/e/{tok}", [
        ("name", "Rey"), ("email", "rey@example.com"), ("sex", "M"),
        ("affiliation", "DOH"), ("consent", "on"),
        (f"q_{multi}", "Vegan"), (f"q_{multi}", "Halal"), (f"q_{multi}", "forged"),
        (f"q_{note}", "no peanuts")])
    assert r.status_code == 303
    (got,) = _attendee_answers(db_path, tok, "rey@example.com")
    assert json.loads(got[multi]) == ["Vegan", "Halal"] and got[note] == "no peanuts"


def test_error_rerender_keeps_checked_boxes(client, manager, db_path):
    tok = _make_event(manager, db_path)
    manager.post(f"/activities/{tok}/questions/preset", data={"preset": "food"})
    multi = next(q["id"] for q in _questions(db_path, tok) if q["kind"] == "multi")
    r = _post_pairs(client, f"/e/{tok}", [
        ("name", ""), ("email", "x@example.com"), ("sex", "M"), ("affiliation", "DOH"),
        ("consent", "on"), (f"q_{multi}", "Vegan"), (f"q_{multi}", "Halal")])
    assert r.status_code == 200 and "Your name is required." in r.text
    assert re.search(r'value="Vegan"\s+checked', r.text)
    assert re.search(r'value="Halal"\s+checked', r.text)
    assert not re.search(r'value="Vegetarian"\s+checked', r.text)


def test_confirm_interstitial_carries_answers_through(client, manager, db_path):
    tok = _make_event(manager, db_path)
    manager.post(f"/activities/{tok}/questions/preset", data={"preset": "food"})
    multi = next(q["id"] for q in _questions(db_path, tok) if q["kind"] == "multi")
    client.post(f"/e/{tok}", data=_reg(name="First Person", email="qshared@office.example",
                                       **{f"q_{multi}": "None"}))
    second = [("name", "Second Person"), ("email", "qshared@office.example"), ("sex", "F"),
              ("affiliation", "DOH"), ("consent", "on"),
              (f"q_{multi}", "Vegan"), (f"q_{multi}", "Halal")]
    r = _post_pairs(client, f"/e/{tok}", second)
    assert r.status_code == 200 and "Is this you?" in r.text
    assert len(re.findall(rf'<input type="hidden" name="q_{multi}"', r.text)) == 2
    # "I'm a different person" -> the new attendee gets the carried answers.
    r = _post_pairs(client, f"/e/{tok}", second + [("resolve", "new")])
    assert r.status_code == 303
    first, new = _attendee_answers(db_path, tok, "qshared@office.example")
    assert json.loads(first[multi]) == ["None"]
    assert json.loads(new[multi]) == ["Vegan", "Halal"]


def test_consent_text_extra_sentence_only_with_questions(client, manager, db_path):
    tok = _make_event(manager, db_path)
    extra = "additional registration"
    assert extra not in client.get(f"/e/{tok}").text
    client.post(f"/e/{tok}", data=_reg(name="Plain Person", email="plain@example.com"))
    _add_q(manager, tok, label="Anything else?")
    # ?fresh=1: the earlier registration left this client's recognition cookie set.
    assert extra in client.get(f"/e/{tok}?fresh=1").text
    client.post(f"/e/{tok}", data=_reg(name="Asked Person", email="asked@example.com"))
    conn = st.connect(db_path)
    ev = st.get_event(conn, tok)
    texts = dict(conn.execute("SELECT email, consent_text FROM attendee WHERE event_id=?",
                              (ev["id"],)).fetchall())
    conn.close()
    assert extra not in texts["plain@example.com"]
    assert texts["asked@example.com"].endswith("deleted together with your personal details.")
    assert texts["asked@example.com"].startswith(texts["plain@example.com"])


def test_every_kind_renders_on_both_pages(client, manager, db_path):
    tok = _make_event(manager, db_path)
    for kind, opts in (("text", ""), ("longtext", ""), ("single", "a\nb"),
                       ("multi", "c\nd"), ("yesno", ""), ("date", ""), ("time", "")):
        assert _add_q(manager, tok, label=f"K-{kind}", kind=kind, options=opts,
                      help_text="a hint" if kind == "text" else "",
                      ).headers["location"].endswith("?qok=1#questions")
    qs = {q["kind"]: q["id"] for q in _questions(db_path, tok)}
    manage = manager.get(f"/activities/{tok}/manage")
    assert manage.status_code == 200
    assert "Registration questions" in manage.text and "K-multi" in manage.text
    page = client.get(f"/e/{tok}")
    assert page.status_code == 200
    h = page.text
    assert f'<input type="text" id="q_{qs["text"]}"' in h
    assert f'<textarea id="q_{qs["longtext"]}"' in h
    assert f'<select id="q_{qs["single"]}"' in h
    assert f'<input type="checkbox" name="q_{qs["multi"]}" value="c"' in h
    assert f'<select id="q_{qs["yesno"]}"' in h and 'value="Yes"' in h and 'value="No"' in h
    assert f'<input type="date" id="q_{qs["date"]}"' in h
    assert f'<input type="time" id="q_{qs["time"]}"' in h
    assert "a hint" in h
    # Multi checkboxes must NOT be `required` (app.js would demand all be ticked).
    assert not re.search(rf'name="q_{qs["multi"]}"[^>]*required', h)


def test_organizer_text_is_escaped(client, manager, db_path):
    tok = _make_event(manager, db_path)
    _add_q(manager, tok, label="<script>alert(1)</script>", kind="single",
           options="<b>x</b>\ny", help_text="<i>h</i>")
    for body in (client.get(f"/e/{tok}").text, manager.get(f"/activities/{tok}/manage").text):
        assert "<script>alert(1)</script>" not in body
        assert "&lt;script&gt;alert(1)&lt;/script&gt;" in body
        assert "<b>x</b>" not in body


def test_no_questions_activity_is_unchanged(client, manager, db_path):
    tok = _make_event(manager, db_path)
    page = client.get(f"/e/{tok}")
    assert "A few more details" not in page.text and "additional registration" not in page.text
    r = client.post(f"/e/{tok}", data=_reg(), follow_redirects=False)
    assert r.status_code == 303
    assert _attendee_answers(db_path, tok, "rey@example.com") == [{}]


# --------------------------------------------------------------------------
# Removing / merging attendees that hold answers (foreign keys are ON)
# --------------------------------------------------------------------------

def _answered_attendee(conn, eid, email):
    qid = st.add_event_question(conn, eid, {"label": "Diet", "kind": "text",
                                            "options": [], "required": False})
    aid, _ = st.register_attendee(conn, eid, name="Ana", email=email, sex="F",
                                  affiliation="DOH", consent_text="x")
    st.save_attendee_answers(conn, aid, {qid: "none"})
    return aid, qid


def test_delete_attendee_entry_removes_its_answers(tmp_path):
    conn = st.connect(str(tmp_path / "del.db")); st.init_db(conn)
    f2 = st.create_form_instance(conn, 2, "F2", "2099-01-01", topic="Review")
    eid = st.create_event(conn, "Forum", "2099-01-01", form_instance_ids=[f2],
                          created_by="m@test.local")["id"]
    aid, _ = _answered_attendee(conn, eid, "ana@x.com")
    assert st.delete_attendee_entry(conn, eid, aid) == ("ok", None)
    assert conn.execute("SELECT COUNT(*) FROM attendee_answer").fetchone()[0] == 0


def test_merge_group_keeps_survivor_answer_and_drops_duplicate(tmp_path):
    import merge_duplicate_attendees as merge
    conn = st.connect(str(tmp_path / "mrg.db")); st.init_db(conn)
    f2 = st.create_form_instance(conn, 2, "F2", "2099-01-01", topic="Review")
    eid = st.create_event(conn, "Forum", "2099-01-01", form_instance_ids=[f2],
                          created_by="m@test.local")["id"]
    a1, qid = _answered_attendee(conn, eid, "ana@x.com")
    a2 = st.add_attendee(conn, eid, name="Ana", email="ana@x.com", sex="F",
                         affiliation="DOH", consent_text="x")
    st.save_attendee_answers(conn, a2, {qid: "vegan"})
    merge.merge_group(conn, a1, [a2])
    conn.commit()
    assert st.get_attendee_answers(conn, a1) == {qid: "none"}
    assert conn.execute("SELECT COUNT(*) FROM attendee_answer").fetchone()[0] == 1


# ==========================================================================
# PR-B — roster, summary, export, purge
# ==========================================================================

def test_display_helpers_and_tally_counts():
    diet = _q(1, "multi", ["None", "Vegan", "Halal"], label="Diet")
    stay = _q(2, "yesno", label="Stay?")
    pick = _q(3, "single", ["S", "M"], label="Size")
    note = _q(4, "text", label="Note")
    gone = dict(_q(5, "single", ["A", "B"], label="Old"), archived_at="2026-01-01")
    qs = [diet, stay, pick, note, gone]
    ans = {
        10: {1: json.dumps(["Vegan", "Halal"]), 2: "Yes", 3: "M", 4: "hi", 5: "A"},
        11: {1: json.dumps(["Vegan"]), 2: "No", 3: "M"},
        12: {1: json.dumps(["Old option"]), 2: "Yes"},
    }
    t = {x["label"]: x for x in rq.tally(qs, ans)}
    assert set(t) == {"Diet", "Stay?", "Size"}                 # no text, no archived
    assert t["Diet"]["counts"] == [("Vegan", 2), ("Halal", 1), ("Old option", 1)]
    assert t["Stay?"]["counts"] == [("Yes", 2), ("No", 1)]
    assert t["Stay?"]["text"] == "Yes 2 · No 1"
    assert t["Size"]["counts"] == [("M", 2)]
    assert rq.tally([_q(1, "single", ["A", "B"], label="Z")], {})[0]["text"] == "No answers yet"
    assert rq.tally([_q(1, "text")], ans) == []
    # answer_rows: question order, blanks skipped, archived flagged.
    assert rq.answer_rows(qs, ans[10]) == [
        ("Diet", "Vegan, Halal"), ("Stay?", "Yes"), ("Size", "M"), ("Note", "hi"),
        ("Old (removed)", "A")]
    assert rq.answer_rows(qs, None) == []


def test_purge_expired_attendees_deletes_only_purged_answers(tmp_path):
    conn = _fresh(tmp_path)
    eid = _event(conn)
    old, qid = _answered_attendee(conn, eid, "old@x.com")
    new, _ = st.register_attendee(conn, eid, name="New", email="new@x.com", sex="F",
                                  affiliation="DOH", consent_text="x")
    st.save_attendee_answers(conn, new, {qid: "keep"})
    conn.execute("UPDATE attendee SET consent_at='2000-01-01T00:00:00' WHERE id=?", (old,))
    conn.commit()
    assert st.purge_expired_attendees(conn, "2001-01-01") == 1
    assert st.get_attendee_answers(conn, old) == {}
    assert st.get_attendee_answers(conn, new) == {qid: "keep"}
    # The attendee row survives as the consent record.
    assert conn.execute("SELECT purged_at FROM attendee WHERE id=?", (old,)).fetchone()[0]


def test_retention_run_reports_registration_answers(tmp_path, capsys):
    import retention
    path = str(tmp_path / "ret.db")
    conn = st.connect(path); st.init_db(conn)
    eid = _event(conn)
    old, qid = _answered_attendee(conn, eid, "old@x.com")
    conn.execute("UPDATE attendee SET consent_at='2000-01-01T00:00:00' WHERE id=?", (old,))
    conn.commit(); conn.close()
    assert retention.run(path, dry_run=True) == 1
    assert "registration answers" in capsys.readouterr().out
    assert retention.run(path) == 1
    conn = st.connect(path)
    assert conn.execute("SELECT COUNT(*) FROM attendee_answer").fetchone()[0] == 0
    assert "registration answers" in conn.execute(
        "SELECT summary FROM audit_log WHERE action='retention_purge'").fetchone()[0]


def _roster_event(client, manager, db_path):
    """An activity with a multi, an archived and a yes/no question, and three
    registrants — one answered everything, one answered nothing, one purged."""
    tok = _make_event(manager, db_path)
    _add_q(manager, tok, label="Food", kind="multi", options="None\nVegan\nHalal")
    _add_q(manager, tok, label="Old Q", kind="text")
    _add_q(manager, tok, label="Stay", kind="yesno")
    qs = {q["label"]: q["id"] for q in _questions(db_path, tok)}
    r = _post_pairs(client, f"/e/{tok}", [
        ("name", "Rey Dela Cruz"), ("email", "rey@example.com"), ("sex", "M"),
        ("affiliation", "DOH"), ("consent", "on"),
        (f"q_{qs['Food']}", "Vegan"), (f"q_{qs['Food']}", "Halal"),
        (f"q_{qs['Old Q']}", "<b>bold</b>"), (f"q_{qs['Stay']}", "Yes")])
    assert r.status_code == 303
    client.post(f"/e/{tok}", data=_reg(name="Bea Silang", email="bea@example.com"))
    client.post(f"/e/{tok}", data=_reg(name="Cy Tan", email="cy@example.com"))
    conn = st.connect(db_path)
    ev = st.get_event(conn, tok)
    cy = conn.execute("SELECT id FROM attendee WHERE event_id=? AND email='cy@example.com'",
                      (ev["id"],)).fetchone()[0]
    st.save_attendee_answers(conn, cy, {qs["Food"]: json.dumps(["None"])})
    conn.close()
    # Archive "Old Q" (answered, so it is hidden rather than deleted).
    manager.post(f"/activities/{tok}/questions/{qs['Old Q']}/remove")
    # Purge Cy only.
    conn = st.connect(db_path)
    conn.execute("UPDATE attendee SET consent_at='2000-01-01T00:00:00' WHERE id=?", (cy,))
    conn.commit()
    st.purge_expired_attendees(conn, "2001-01-01")
    conn.close()
    return tok, qs


def test_roster_shows_answers_and_summary(client, manager, db_path):
    tok, _ = _roster_event(client, manager, db_path)
    html = manager.get(f"/monitor/{tok}").text
    assert "Registration answers" in html
    assert "Vegan, Halal" in html
    assert "Old Q (removed)" in html                       # archived: kept in details
    assert "&lt;b&gt;bold&lt;/b&gt;" in html and "<b>bold</b>" not in html   # escaped
    glance = html.split("Registration answers at a glance")[1].split("</div>")[0]
    assert "Food:" in glance and "Vegan 1 · Halal 1" in glance
    assert "Stay:" in glance and "Yes 1" in glance
    assert "Old Q" not in glance                           # archived: not tallied
    assert "None 1" not in glance                          # Cy's answer was purged


def test_roster_without_questions_is_unchanged(client, manager, db_path):
    tok = _make_event(manager, db_path)
    client.post(f"/e/{tok}", data=_reg())
    html = manager.get(f"/monitor/{tok}").text
    assert "Rey Dela Cruz" in html
    assert "Registration answers" not in html and '<details class="regans"' not in html


def test_export_includes_question_columns(client, manager, db_path):
    import io
    from openpyxl import load_workbook
    tok, qs = _roster_event(client, manager, db_path)
    r = manager.get(f"/monitor/{tok}/export")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Registration"]
    header = [c.value for c in ws[6]]
    assert header[:8] == ["No.", "Name", "Email", "Contact", "Sex", "Affiliation",
                          "Registered at", "CSS status"]
    assert header[8:] == ["Food", "Old Q (removed)", "Stay"]
    rows = {ws.cell(r_, 2).value: [c.value for c in ws[r_]] for r_ in range(7, ws.max_row + 1)}
    assert rows["Rey Dela Cruz"][8:] == ["Vegan, Halal", "<b>bold</b>", "Yes"]
    assert rows["Bea Silang"][8:] in (["", "", ""], [None, None, None])
    purged = rows["(identity removed)"]
    assert all(v in ("", None) for v in purged[8:])         # purged: answer cells blank
    # Letterhead merges span every column, not just the original eight.
    assert any(str(m) == "A1:K1" for m in ws.merged_cells.ranges)
    assert ws.column_dimensions["K"].width == 24
    about = " ".join(str(c.value) for row in wb["About this report"].iter_rows() for c in row if c.value)
    assert "registration questions" in about


def test_export_without_questions_keeps_eight_columns(client, manager, db_path):
    import io
    from openpyxl import load_workbook
    tok = _make_event(manager, db_path)
    client.post(f"/e/{tok}", data=_reg())
    wb = load_workbook(io.BytesIO(manager.get(f"/monitor/{tok}/export").content))
    ws = wb["Registration"]
    assert [c.value for c in ws[6] if c.value is not None][-1] == "CSS status"
    assert ws.max_column == 8
    assert any(str(m) == "A1:H1" for m in ws.merged_cells.ranges)
    about = " ".join(str(c.value) for row in wb["About this report"].iter_rows() for c in row if c.value)
    assert "registration questions" not in about


def test_export_with_questions_and_no_registrants(manager, db_path):
    import io
    from openpyxl import load_workbook
    tok = _make_event(manager, db_path)
    _add_q(manager, tok, label="Diet")
    ws = load_workbook(io.BytesIO(manager.get(f"/monitor/{tok}/export").content))["Registration"]
    assert ws.cell(6, 9).value == "Diet"
    assert any(str(m) == "A7:I7" for m in ws.merged_cells.ranges)   # "No one has registered"
