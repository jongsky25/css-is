"""Office calendar Word import screens (increment B2): upload -> review -> import,
admin only, all-or-nothing, idempotent."""
from __future__ import annotations
import io
import json
from urllib.parse import urlencode

import docx
import pytest

import calendar_routes
import storage as st
from test_calendar_import import _color_run, month_table

FORM = {"Content-Type": "application/x-www-form-urlencoded"}


def make_docx(extra=None):
    doc = docx.Document()
    days = {5: [("CALIM TA Palawan", "hl:cyan")], 6: [("CALIM TA Palawan", "hl:cyan"),
                                                       ("CALIM Section meeting (PK)", "hl:yellow")],
            7: [("CALIM TA Palawan", "hl:cyan")], 12: [("CALIM Strat Planning", "hl:green")],
            13: [("CALIM Loose end", None)]}
    days.update(extra or {})
    t = month_table(doc, 10, days)
    t.rows[4].cells[3].paragraphs[0].text = "14  Non-working Holiday"      # Wed Oct 14 (row 4 = week of Oct 11)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def conn(app_mod):
    c = st.connect(app_mod.DB)
    st.init_db(c)
    yield c
    for (i,) in c.execute("SELECT id FROM calendar_entry WHERE title LIKE 'CALIM %'").fetchall():
        st.cal_delete(c, i)
    c.execute("DELETE FROM calendar_import_batch")
    c.execute("DELETE FROM app_setting WHERE key IN ('calendar_import_owner_units', 'calendar_unit_tones')")
    c.execute("DELETE FROM org_unit WHERE path LIKE 'CALIM%'")
    c.commit()
    c.close()


@pytest.fixture()
def units(conn):
    out = {}
    for name, level in (("Bureau", "bureau"), ("Section 1", "section"), ("Section 4", "section")):
        out[name] = conn.execute("INSERT INTO org_unit(name, short_name, level, path) VALUES (?,?,?,?)",
                                 (name, name, level, f"CALIM/{name}")).lastrowid
    conn.commit()
    return out


def upload(client, data=None, name="CY-2026.docx"):
    return client.post("/calendar/import", files={"file": (name, data if data is not None else make_docx(),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                       follow_redirects=False)


def batch_id(r):
    assert r.status_code == 303, r.text
    return int(r.headers["location"].rsplit("/", 1)[1])


def form_for(conn, bid, **over):
    """The review form exactly as the page would post it (everything ticked that is
    not already imported), with `over` applied on top."""
    b = st.cal_import_get(conn, bid)
    existing = st.cal_existing_import_keys(conn, [d["import_key"] for d in b["payload"]["drafts"]])
    f = {"audience": "bureau"}
    for i, d in enumerate(b["payload"]["drafts"]):
        if d["import_key"] not in existing:
            f[f"inc_{i}"] = "1"
        f.update({f"title_{i}": d["title"], f"kind_{i}": d["kind"], f"status_{i}": d["status"],
                  f"owner_{i}": d["owner"] or "", f"start_{i}": d["start_date"],
                  f"end_{i}": d["end_date"] or ""})
    f.update(over)
    return f


def post_import(client, bid, form):
    return client.post(f"/calendar/import/{bid}", content=urlencode(form), headers=FORM,
                       follow_redirects=False)


def titles(conn):
    return sorted(r[0] for r in conn.execute("SELECT title FROM calendar_entry WHERE title LIKE 'CALIM %'"))


# ------------------------------------------------------------------- access ---

@pytest.mark.parametrize("who", ["viewer", "encoder", "manager"])
def test_only_admins_may_import(request, who, conn):
    c = request.getfixturevalue(who)
    assert c.get("/calendar/import").status_code == 403
    assert upload(c).status_code == 403
    assert c.post("/calendar/import/1", content="x=1", headers=FORM).status_code == 403
    assert c.post("/calendar/import/1/discard").status_code == 403
    assert conn.execute("SELECT COUNT(*) FROM calendar_import_batch").fetchone()[0] == 0


def test_anonymous_is_sent_to_login(client):
    assert client.get("/calendar/import", follow_redirects=False).status_code in (302, 303, 401)


def test_the_import_link_is_shown_to_admins_only(admin, viewer):
    assert "Import from Word" in admin.get("/calendar").text
    assert "Import from Word" not in viewer.get("/calendar").text


# ------------------------------------------------------------------- upload ---

def test_upload_page_renders(admin):
    r = admin.get("/calendar/import")
    assert r.status_code == 200 and 'type="file"' in r.text


@pytest.mark.parametrize("data, name, message", [
    (b"", "x.docx", "Word (.docx)"),
    (b"not a zip at all", "x.docx", "look like the calendar Word file"),
    (b"%PDF-1.4", "x.pdf", "Word (.docx)"),
])
def test_bad_uploads_get_a_friendly_error_and_save_nothing(admin, conn, data, name, message):
    r = upload(admin, data, name)
    assert r.status_code == 400 and message in r.text
    assert conn.execute("SELECT COUNT(*) FROM calendar_import_batch").fetchone()[0] == 0


def test_a_docx_with_no_activities_is_refused(admin, conn):
    doc = docx.Document(); month_table(doc, 10, {})
    buf = io.BytesIO(); doc.save(buf)
    r = upload(admin, buf.getvalue())
    assert r.status_code == 400 and "No activities were found" in r.text


def test_an_oversized_upload_is_refused(admin, monkeypatch):
    monkeypatch.setattr(calendar_routes, "MAX_UPLOAD", 100)
    assert "too large" in upload(admin).text


def test_upload_stores_a_batch_and_shows_the_review(admin, conn, units):
    bid = batch_id(upload(admin))
    b = st.cal_import_get(conn, bid)
    assert b["filename"] == "CY-2026.docx" and b["created_by"] == "admin@test.local" and b["imported_at"] is None
    assert [d["title"] for d in b["payload"]["drafts"]] == [
        "CALIM TA Palawan", "CALIM Section meeting (PK)", "CALIM Strat Planning", "CALIM Loose end"]
    html = admin.get(f"/calendar/import/{bid}").text
    for needle in ("CALIM TA Palawan", "CALIM Loose end", "2026-10", "Section 4", "1 holiday lines",
                   "Which unit does each colour belong to", "Import the ticked items"):
        assert needle in html, needle
    assert "name=\"map_s1\"" in html and "name=\"inc_0\"" in html
    assert admin.get("/calendar/import/99999", follow_redirects=False).status_code == 303


def test_the_review_guesses_units_by_name(admin, conn, units):
    html = admin.get(f"/calendar/import/{batch_id(upload(admin))}").text
    for key, name in (("s1", "Section 1"), ("s4", "Section 4")):        # "bureau" is ambiguous in a real DB
        sel = html.split(f'name="map_{key}"')[1].split("</select>")[0]
        assert f'value="{units[name]}" selected' in sel, key


# ------------------------------------------------------------------- import ---

def test_import_creates_the_ticked_items_with_owner_tone_and_one_audit_line(admin, conn, units):
    bid = batch_id(upload(admin))
    form = form_for(conn, bid, map_bureau=units["Bureau"], map_s1=units["Section 1"],
                    map_s4=units["Section 4"])
    last_audit = conn.execute("SELECT COALESCE(MAX(id), 0) FROM audit_log").fetchone()[0]
    r = post_import(admin, bid, form)
    assert r.status_code == 303 and titles(conn) == [
        "CALIM Loose end", "CALIM Section meeting (PK)", "CALIM Strat Planning", "CALIM TA Palawan"]
    rows = {e["title"]: e for e in st.cal_list_between(conn, "2026-10-01", "2026-10-31")
            if e["title"].startswith("CALIM ")}
    ta = rows["CALIM TA Palawan"]
    assert (ta["start_date"], ta["end_date"], ta["kind"], ta["source"], ta["audience"]) == (
        "2026-10-05", "2026-10-07", "travel", "import", "bureau")
    assert ta["owner_org_id"] == units["Section 4"] and ta["import_key"].startswith("docx:")
    assert ta["created_by"] == "admin@test.local"
    assert rows["CALIM Section meeting (PK)"]["tags"] == ["PK"] and rows["CALIM Section meeting (PK)"]["owner_org_id"] == units["Section 1"]
    assert rows["CALIM Loose end"]["owner_org_id"] is None
    tones = json.loads(st.get_setting(conn, "calendar_unit_tones"))
    assert tones[str(units["Section 4"])] == "cyan" and tones[str(units["Section 1"])] == "yellow"
    assert tones[str(units["Bureau"])] == "green"
    acts = [r[0] for r in conn.execute("SELECT action FROM audit_log WHERE id > ? AND action LIKE 'calendar_%'",
                                       (last_audit,))]       # the audit log is append-only: look only at ours
    assert acts.count("calendar_import") == 1 and "calendar_create" not in acts
    page = admin.get(f"/calendar/import/{bid}").text
    assert "Imported 4 items" in page and "Import the ticked items" not in page
    assert 'data-tone="cyan"' in admin.get("/calendar?day=2026-10-05&scope=bureau").text      # the legend colour is used


def test_the_mapping_is_remembered_for_the_next_upload(admin, conn, units):
    bid = batch_id(upload(admin))
    post_import(admin, bid, form_for(conn, bid, map_s1=units["Section 1"]))
    html = admin.get(f"/calendar/import/{batch_id(upload(admin, make_docx({20: [('CALIM Next one', 'hl:yellow')]})))}").text
    assert f'value="{units["Section 1"]}" selected' in html.split('name="map_s1"')[1].split("</select>")[0]


def test_unticked_rows_are_not_imported_and_edits_are_used(admin, conn, units):
    bid = batch_id(upload(admin))
    form = form_for(conn, bid, map_s4=units["Section 4"])
    for k in [k for k in form if k.startswith("inc_")][1:]:
        del form[k]                                           # keep only the first row
    form.update(title_0="CALIM TA Palawan (edited)", kind_0="meeting", status_0="tentative", owner_0="",
                end_0="")
    assert post_import(admin, bid, form).status_code == 303
    (e,) = [e for e in st.cal_list_between(conn, "2026-10-01", "2026-10-31") if e["title"].startswith("CALIM ")]
    assert (e["title"], e["kind"], e["status"], e["owner_org_id"], e["end_date"]) == (
        "CALIM TA Palawan (edited)", "meeting", "tentative", None, None)
    assert st.cal_get(conn, e["id"])["import_key"]
    assert "3 were not selected" in admin.get(f"/calendar/import/{bid}").text


def test_unit_audience_applies_only_where_a_unit_was_mapped(admin, conn, units):
    bid = batch_id(upload(admin))
    post_import(admin, bid, form_for(conn, bid, audience="unit", map_s4=units["Section 4"]))
    by = {e["title"]: e for e in st.cal_list_between(conn, "2026-10-01", "2026-10-31")
          if e["title"].startswith("CALIM ")}
    assert by["CALIM TA Palawan"]["audience"] == "unit"
    assert by["CALIM Loose end"]["audience"] == "bureau"          # no unit -> never invisible
    assert by["CALIM Strat Planning"]["audience"] == "bureau"     # Bureau colour left unmapped
    assert "no unit chosen" in admin.get(f"/calendar/import/{bid}").text


def test_importing_the_same_file_twice_adds_nothing_and_overwrites_nothing(admin, conn, units):
    first = batch_id(upload(admin))
    post_import(admin, first, form_for(conn, first))
    e = next(e for e in st.cal_list_between(conn, "2026-10-01", "2026-10-31") if e["title"] == "CALIM TA Palawan")
    st.cal_update(conn, e["id"], title="CALIM TA Palawan (changed since)", by="mae")
    second = batch_id(upload(admin))
    html = admin.get(f"/calendar/import/{second}").text
    assert html.count("Already in the calendar") == 4 and "4</strong> are already in the calendar" in html
    form = form_for(conn, second)
    form.update({f"inc_{i}": "1" for i in range(4)})              # even if someone ticks them all
    assert post_import(admin, second, form).status_code == 303
    assert len(titles(conn)) == 4 and "CALIM TA Palawan (changed since)" in titles(conn)
    assert "Imported 0 items" in admin.get(f"/calendar/import/{second}").text


def test_a_bad_row_blocks_the_whole_import_and_keeps_what_was_typed(admin, conn, units):
    bid = batch_id(upload(admin))
    form = form_for(conn, bid, start_1="2026-10-20", end_1="2026-10-01", title_0="CALIM typed text")
    r = post_import(admin, bid, form)
    assert r.status_code == 400 and "Nothing was imported" in r.text and "end_date cannot be before" in r.text
    assert "CALIM typed text" in r.text and titles(conn) == []
    assert st.cal_import_get(conn, bid)["imported_at"] is None


@pytest.mark.parametrize("over, message", [
    ({"title_0": "  "}, "title is required"),
    ({"kind_0": "activity"}, "pick a type"),
    ({"status_0": "cancelled"}, "pick a status"),
    ({"owner_0": "s9"}, "unknown owner"),
    ({"start_0": "soon"}, "YYYY-MM-DD"),
    ({"audience": "world"}, "who can see"),
    ({"map_s1": "999999"}, "no longer exists"),
])
def test_validation_messages(admin, conn, units, over, message):
    bid = batch_id(upload(admin))
    r = post_import(admin, bid, form_for(conn, bid, **over))
    assert r.status_code == 400 and message.lower() in r.text.lower() and titles(conn) == []


def test_two_colours_cannot_share_a_unit(admin, conn, units):
    bid = batch_id(upload(admin))
    r = post_import(admin, bid, form_for(conn, bid, map_s1=units["Section 1"], map_s4=units["Section 1"]))
    assert r.status_code == 400 and "same unit" in r.text and titles(conn) == []


def test_an_imported_batch_cannot_be_imported_again_or_discarded(admin, conn, units):
    bid = batch_id(upload(admin))
    post_import(admin, bid, form_for(conn, bid))
    before = titles(conn)
    assert post_import(admin, bid, form_for(conn, bid)).status_code == 303 and titles(conn) == before
    admin.post(f"/calendar/import/{bid}/discard")
    assert st.cal_import_get(conn, bid) is not None


def test_discard_throws_an_unimported_batch_away(admin, conn):
    bid = batch_id(upload(admin))
    assert admin.post(f"/calendar/import/{bid}/discard", follow_redirects=False).status_code == 303
    assert st.cal_import_get(conn, bid) is None and titles(conn) == []


def test_the_upload_list_shows_earlier_batches(admin, conn):
    bid = batch_id(upload(admin, name="CY-2026-draft.docx"))
    html = admin.get("/calendar/import").text
    assert "CY-2026-draft.docx" in html and f"/calendar/import/{bid}" in html and "Waiting for review" in html
