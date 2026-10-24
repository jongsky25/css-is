"""A18 report selection, receipts at date boundaries and downloadable evidence."""
import io
from urllib.parse import parse_qs, urlsplit
from zipfile import ZipFile

import pandas as pd
import pytest
from lxml import etree, html
from openpyxl import load_workbook
from pypdf import PdfReader

import css_is_engine as eng
import storage as st
from tests.test_rp_report_filter import _make_bureau, _add_responses
from tests.test_csm_public_form import walkin
from tests.test_sparse_analytics import reference


def export_text(response, fmt):
    assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ''
    if fmt == 'pdf':
        return ' '.join(' '.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages).split())
    with ZipFile(io.BytesIO(response.content)) as z:
        if fmt == 'docx':
            names = ['word/document.xml']
        else:
            names = [n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')]
        return ' '.join(' '.join(etree.fromstring(z.read(n)).itertext()) for n in names)


@pytest.fixture
def rp_scope(manager, app_mod):
    _, first = _make_bureau(manager, app_mod.DB, 'Scope Alpha', '2026-06-10', 'A18 scope room')
    # A second form is a member of the same merged activity, not its representative.
    _, member = _make_bureau(manager, app_mod.DB, 'Scope Alpha second', '2026-06-10', 'A18 scope room')
    _, other = _make_bureau(manager, app_mod.DB, 'Scope Bravo', '2026-09-12', 'A18 other room')
    _add_responses(app_mod.DB, first, 2)
    _add_responses(app_mod.DB, member, 1)
    _add_responses(app_mod.DB, other, 4)
    app_mod.invalidate_atoms()
    return first, member, other


@pytest.mark.parametrize('fmt', ['docx', 'pdf', 'pptx', 'xlsx'])
def test_activity_member_and_stale_scope_are_identical_in_report_and_export(manager, app_mod, rp_scope, fmt):
    _, member, _ = rp_scope
    params = {'type': 3, 'activity': str(member), 'period': 'quarter:2026-Q2'}
    page = manager.get('/reports', params=params)
    context = html.fromstring(page.text).get_element_by_id('report-context').text_content()
    assert '3 total submission(s)' in context
    response = manager.get('/reports/export', params={**params, 'fmt': fmt})
    if fmt == 'xlsx':
        wb = load_workbook(io.BytesIO(response.content))
        assert wb['Responses'].max_row - 3 == 3
        assert wb['Summary']['B4'].value == 3
        text = ' '.join(str(c.value) for ws in wb for row in ws for c in row)
    else:
        text = export_text(response, fmt)
    assert '3 total submission(s)' in text and '3 rated respondent(s)' in text
    assert '2026-04-01 to 2026-06-30' in text
    assert 'activity start date' in text
    assert 'Scope Bravo' not in text

    params['activity'] = '99999999'
    page = manager.get('/reports', params=params)
    assert '0 total submission(s)' in page.text
    assert 'Unavailable activity (saved filter)' in page.text
    assert 'not been widened automatically' in page.text
    response = manager.get('/reports/export', params={**params, 'fmt': fmt})
    if fmt == 'xlsx':
        wb = load_workbook(io.BytesIO(response.content))
        assert 'Responses' not in wb.sheetnames
        text = ' '.join(str(c.value) for ws in wb for row in ws for c in row)
    else:
        text = export_text(response, fmt)
    assert '0 total submission(s)' in text
    assert 'Scope Alpha' not in text and 'Scope Bravo' not in text


def test_roster_member_stale_and_period_filters(manager, rp_scope):
    _, member, _ = rp_scope
    params = {'activity': str(member), 'period': 'quarter:2026-Q2'}
    text = export_text(manager.get('/reports/rp-roster-export', params=params), 'pptx')
    assert 'Scope Alpha' in text and 'Scope Bravo' not in text
    for changed in ({'activity': '99999999'}, {'period': 'quarter:2026-Q4'}):
        text = export_text(manager.get('/reports/rp-roster-export', params={**params, **changed}), 'pptx')
        assert 'No resource-person activities in this scope.' in text
        assert 'Scope Bravo' not in text


@pytest.mark.parametrize('url', ['/reports?type=2', '/reports/export?type=3&fmt=xlsx',
                                '/reports/rp-roster-export?', '/csm-blhsd/export?',
                                '/reports/verbatims?survey_type=1'])
@pytest.mark.parametrize('period', ['bad', 'semester:2026-H3', 'quarter:2026-Q0', 'year:0000'])
def test_malformed_period_does_not_widen_a_report(manager, url, period):
    response = manager.get(url + ('&' if '?' in url else '?') + 'period=' + period)
    assert response.status_code == 400
    assert 'Choose' in response.text


def test_unknown_saved_filters_remain_visible_and_reset_preserves_report_type(manager, rp_scope):
    page = manager.get('/reports?type=3&activity=99999999&org=unknown&period=year:2040')
    soup = html.fromstring(page.text)
    assert soup.xpath('//select[@id="period"]/option[@selected]')[0].get('value') == 'year:2040'
    assert soup.xpath('//select[@id="activity"]/option[@selected]')[0].get('value') == '99999999'
    assert soup.xpath('//select[@id="org"]/option[@selected]')[0].get('value') == 'unknown'
    assert soup.get_element_by_id('report-context').xpath('.//a')[0].get('href') == '/reports?type=3'
    assert 'No submissions match these filters.' in page.text
    assert soup.get_element_by_id('filterForm').xpath('.//button[@type="submit"]')[0].text_content() == 'Apply filters'


def test_org_is_exact_node_plus_descendants(app_mod):
    atoms = eng.assign_periods(pd.DataFrame({
        'response_id': [1, 2, 3, 4, 5], 'survey_type': [2]*5,
        'org_path': ['ROOT/DIV', 'ROOT/DIV/SEC', 'ROOT/DIV2', None, 'ROOT/DIVISION'],
        'event_date': ['2026-01-01']*5}))
    assert app_mod._report_scope_atoms(atoms, 'ROOT/DIV', 'all', 2)['response_id'].tolist() == [1, 2]


def test_walkin_submission_period_archives_and_last_fractional_second(manager, app_mod, walkin):
    conn = st.connect(app_mod.DB)
    fid = conn.execute('SELECT id FROM form_instance WHERE qr_token=?', (walkin,)).fetchone()[0]
    conn.execute("UPDATE form_instance SET event_date='2025-01-01' WHERE id=?", (fid,))
    conn.commit()
    rids = []
    for timestamp, remark in [('2026-06-30T23:59:59.999999', 'A18 included fraction'),
                              ('2026-06-30 23:59:59.999999', 'A18 included space'),
                              ('2026-07-01T00:00:00', 'A18 outside quarter'),
                              ('2026-06-30T12:00:00', 'A18 archived response')]:
        rids.append(st.record_response(conn, fid, [('SQD1', 5)], submitted_at=timestamp, remarks=remark))
    conn.execute('UPDATE response SET archived_at=? WHERE id=?', ('2026-07-02', rids[-1]))
    conn.commit()
    included = app_mod._csm_type1_response_ids(conn, '2026-04-01', '2026-06-30')
    assert set(included) == set(rids[:2])
    rows = app_mod._csm_blhsd_rows(conn, '2026-04-01', '2026-06-30')
    assert len(rows) == 2
    conn.close()
    app_mod.invalidate_atoms()
    params = {'type': 1, 'period': 'quarter:2026-Q2', 'org': 'irrelevant-office'}
    page = manager.get('/reports', params=params)
    soup = html.fromstring(page.text)
    assert not soup.xpath('//*[@id="org"]')
    context = soup.get_element_by_id('report-context').text_content()
    assert '2 total submission(s)' in context
    assert 'Submission dates' in context and '2026-04-01 to 2026-06-30' in context
    assert 'irrelevant-office' not in context
    assert soup.xpath('//option[@value="quarter:2026-Q3"]')
    text = export_text(manager.get('/reports/export', params={**params, 'fmt': 'docx'}), 'docx')
    assert 'A18 included fraction' in text and 'A18 included space' in text
    assert 'A18 archived response' not in text and 'A18 outside quarter' not in text
    comments = manager.get('/reports/verbatims', params={'survey_type': 1, 'period': 'quarter:2026-Q2', 'org': 'irrelevant-office'})
    assert 'A18 included fraction' in comments.text and 'A18 included space' in comments.text
    assert 'A18 archived response' not in comments.text and 'A18 outside quarter' not in comments.text
    wb = load_workbook(io.BytesIO(manager.get('/csm-blhsd/export', params={'period': 'quarter:2026-Q2'}).content))
    text = ' '.join(str(c.value) for ws in wb for row in ws for c in row)
    assert 'A18 included fraction' in text and 'A18 included space' in text
    assert 'A18 archived response' not in text and 'A18 outside quarter' not in text


def test_walkin_full_export_batches_response_lookups(app_mod, walkin):
    conn = st.connect(app_mod.DB)
    fid = conn.execute('SELECT id FROM form_instance WHERE qr_token=?', (walkin,)).fetchone()[0]
    with conn:
        for n in range(1100):
            rid = conn.execute(
                'INSERT INTO response(form_instance_id,submitted_at,remarks) VALUES (?,?,?)',
                (fid, '2026-10-20T09:00:00', f'Synthetic note {n}')).lastrowid
            conn.execute('INSERT INTO answer(response_id,item_id,value,is_na) VALUES (?,?,?,0)',
                         (rid, 'SQD1', (n % 5) + 1))
            conn.execute('INSERT INTO item_remark(response_id,item_id,text) VALUES (?,?,?)',
                         (rid, 'SERVICE', f'Synthetic service {n}'))
    rows = app_mod._csm_blhsd_rows(conn)
    assert len(rows) >= 1100
    by_note = {row['SUGGESTIONS']: row for row in rows}
    assert by_note['Synthetic note 0']['SERVICE'] == 'Synthetic service 0'
    assert by_note['Synthetic note 0']['SQD1'] == 1
    assert by_note['Synthetic note 1099']['SERVICE'] == 'Synthetic service 1099'
    assert by_note['Synthetic note 1099']['SQD1'] == 5
    conn.close()


def test_unrated_submission_kept_in_excel_and_distinguished_from_empty(manager, app_mod):
    _, fid = _make_bureau(manager, app_mod.DB, 'Unrated scope', '2026-08-09', 'No-score room')
    conn = st.connect(app_mod.DB)
    st.record_response(conn, fid, [('B1', None)])
    conn.close()
    app_mod.invalidate_atoms()
    params = {'type': 3, 'activity': str(fid)}
    page = manager.get('/reports', params=params)
    assert '1 total submission(s)' in page.text
    assert 'Submissions exist, but none has a valid dimension rating.' in page.text
    response = manager.get('/reports/export', params={**params, 'fmt': 'xlsx'})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert wb['Responses'].max_row - 3 == 1
    assert wb['Summary']['B4'].value == 1
    headers = [c.value for c in wb['Responses'][3]]
    assert wb['Responses'].cell(4, headers.index('Overall')+1).value is None
    comments = manager.get('/reports/verbatims', params={'survey_type': 3, 'activity': str(fid), 'period': 'year:2026', 'org': ''})
    soup = html.fromstring(comments.text)
    back = next(a.get('href') for a in soup.xpath('//nav[@aria-label="Breadcrumb"]//a[starts-with(@href,"/reports?")]') if 'type=3' in a.get('href'))
    query = parse_qs(urlsplit(back).query)
    assert query['period'] == ['year:2026'] and query['activity'] == [str(fid)]


@pytest.mark.parametrize('fmt', ['docx', 'pdf', 'pptx', 'xlsx'])
def test_reference_download_agrees_with_report(reference, manager, app_mod, monkeypatch, fmt):
    # Inject the independently calculated sparse/categorical dataset at the source
    # boundary, then use the real authenticated HTTP screen and document builders.
    monkeypatch.setattr(app_mod, 'get_atoms', lambda *a, **k: eng.assign_periods(reference))
    monkeypatch.setattr(st, 'demographics_for_responses', lambda *a: {})
    monkeypatch.setattr(st, 'response_context_for_ids', lambda *a: {})
    page = manager.get('/reports?type=2&period=year:2026')
    assert page.status_code == 200
    assert '8 total submission(s)' in page.text and '2 rated respondent(s)' in page.text
    assert '4.00' in page.text
    response = manager.get('/reports/export', params={'type': 2, 'period': 'year:2026', 'fmt': fmt})
    if fmt == 'xlsx':
        wb = load_workbook(io.BytesIO(response.content))
        assert wb['Responses'].max_row - 3 == 8
        assert wb['Summary']['B4'].value == 8
        assert wb['Summary']['B6'].value == '4.00'
        text = ' '.join(str(c.value) for ws in wb for row in ws for c in row)
    else:
        text = export_text(response, fmt)
    assert '8 total submission(s)' in text and '2 rated respondent(s)' in text
    assert '4.00' in text
