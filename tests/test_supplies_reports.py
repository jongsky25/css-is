"""S6.1 saved-transaction reports, export safety and scoped reads."""
from io import BytesIO

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

import app
import supplies as s
import supplies_stock as stock
import supplies_reports as reports
import supplies_replenishment as rep
import supplies_corrections as corrections
from supplies_storage import Denied, Invalid
from tests.test_supplies import supplies_db, key
from tests.test_supplies_requests import request_db
from tests.test_supplies_approvals import approval_db
from tests.test_supplies_replenishment_flow import flow_db
from tests.test_supplies_deliveries import approved, payload, post
from tests.test_supplies_backorders import ready
from tests.test_supplies_issues import issue
from tests.conftest import session_token


def test_inventory_replenishment_and_stock_card_reconcile_linked_receipt(flow_db, monkeypatch):
    c,u,_,_,item=flow_db
    rid=approved(c,u,item,10)
    delivery=post(c,u,rid,payload(c,u,rid,6,accepted=4,rejected=2,reason='Damaged packaging'))
    inventory=reports.report(c,u[1],'inventory')['rows']
    row=next(x for x in inventory if x['item_id']==item)
    assert (row['on_hand'],row['reserved'],row['available'],row['incoming'])==(14,0,14,6)
    pending=reports.report(c,u[1],'replenishment')['rows']
    assert (pending[0]['approved'],pending[0]['net_received'],pending[0]['incoming'])==(10,4,6)
    card=reports.report(c,u[1],'stock-card',item_id=item)['rows']
    assert [x['balance'] for x in card]==[10,14]
    assert card[-1]['reference']==f"RCV-{delivery['receipt_id']:06d}"
    assert card[-1]['supplier']=='Synthetic supplier' and card[-1]['dr_number']=='DR-SYN-1'
    assert reports.report(c,u[1],'stock-card',item_id=item,as_of='2026-10-19')['rows']==[]
    receipt=stock.get_receipt(c,u[1],delivery['receipt_id'])
    line=receipt['lines'][0]
    fields={'source_line_id':line['id'],'quantity':1,'effective_date':'2026-10-20',
            'reason':'Synthetic correction','reopen_incoming':'yes'}
    preview=corrections.preview(c,u[1],'receipt_reversal',receipt['id'],fields)
    monkeypatch.setattr(corrections.tz,'now_iso',lambda: '2026-10-21T09:00:00')
    corrections.reverse_receipt(c,u[1],receipt['id'],{**fields,'preview_token':preview['preview_token'],
        'confirmed':'yes'},operation_key=key())
    row=next(x for x in reports.report(c,u[1],'inventory')['rows'] if x['item_id']==item)
    assert (row['on_hand'],row['incoming'])==(13,7)
    assert [x['balance'] for x in reports.report(c,u[1],'stock-card',item_id=item)['rows']]==[10,14,13]
    assert [x['balance'] for x in reports.report(c,u[1],'stock-card',item_id=item,as_of='2026-10-20')['rows']]==[10,14]


def test_request_fulfillment_and_issue_consumption_net_reversal(flow_db, monkeypatch):
    c,u,_,_,item=flow_db
    rid=ready(c,u,item,6)
    released=issue(c,u,rid,4)
    fulfillment=reports.report(c,u[1],'fulfillment')['rows']
    row=next(x for x in fulfillment if x['request_id']==rid)
    assert (row['approved'],row['issued'],row['reserved'],row['awaiting'])==(6,4,2,0)
    consumed=reports.report(c,u[1],'consumption')['rows']
    assert len(consumed)==1 and consumed[0]['net_issued']==4
    assert consumed[0]['release_number']==released['release_number']
    assert reports.report(c,u[1],'consumption',as_of='2026-10-19')['rows']==[]
    source=c.execute('SELECT id FROM supply_issue_line WHERE issue_id=?',(released['id'],)).fetchone()[0]
    fields={'source_line_id':source,'quantity':1,'effective_date':'2026-10-20',
            'reason':'Synthetic issue correction','reopen_demand':'yes'}
    preview=corrections.preview(c,u[1],'issue_reversal',released['id'],fields)
    monkeypatch.setattr(corrections.tz,'now_iso',lambda: '2026-10-21T09:00:00')
    corrections.reverse_issue(c,u[1],released['id'],{**fields,'preview_token':preview['preview_token'],
        'confirmed':'yes'},operation_key=key())
    row=next(x for x in reports.report(c,u[1],'fulfillment')['rows'] if x['request_id']==rid)
    assert (row['issued'],row['reserved'])==(3,3)
    assert reports.report(c,u[1],'consumption')['rows'][0]['net_issued']==3
    assert reports.report(c,u[1],'consumption')['summary'][0]['net_issued']==3
    assert reports.report(c,u[1],'consumption',as_of='2026-10-20')['rows'][0]['net_issued']==4


def test_export_escapes_formula_text_and_restricts_report_access(flow_db):
    c,u,_,_,_=flow_db
    item=s.create_item(c,u[1],{'code':'FORMULA','description':'=SUM(1,2)',
        'unit':'Piece','critical_level':'0','target_level':'5'},operation_key=key())['id']
    result=reports.report(c,u[1],'inventory')
    content=reports.xlsx_bytes(result)
    sheet=load_workbook(BytesIO(content)).active
    cell=next(cell for row in sheet.iter_rows(min_row=7) for cell in row if cell.value=="'=SUM(1,2)")
    assert cell.data_type=='s'
    assert reports._safe_cell('   +HYPERLINK("x")').startswith("'")
    with pytest.raises(Denied):reports.report(c,u[3],'inventory')
    with pytest.raises(Invalid):reports.report(c,u[1],'inventory',as_of='2026-10-20')
    with pytest.raises(Invalid):reports.report(c,u[1],'stock-card')
    assert next(r for r in result['rows'] if r['item_id']==item)['description']=='=SUM(1,2)'


def test_report_html_print_xlsx_and_scoped_http(flow_db):
    c,u,_,path,item=flow_db
    old=app.DB;app.DB=path
    try:
        client=TestClient(app.app,raise_server_exceptions=True)
        import auth
        client.cookies.set(auth.COOKIE_NAME,session_token(u[1]))
        assert client.get('/supplies/reports/inventory').status_code==200
        assert client.get(f'/supplies/reports/stock-card/print?item_id={item}').status_code==200
        assert client.get('/supplies/reports/consumption/print').status_code==200
        download=client.get('/supplies/reports/inventory/xlsx')
        assert download.status_code==200 and load_workbook(BytesIO(download.content)).active['A6'].value=='Item code'
        client.cookies.set(auth.COOKIE_NAME,session_token(u[3]))
        assert client.get('/supplies/reports/inventory').status_code==403
        assert client.get('/supplies/reports/inventory/xlsx').status_code==403
    finally:app.DB=old
