"""
Work arrangement module (ANNEX C) — the semi-monthly internal work schedule:
period and deadline arithmetic, field validation, the staff-initiated change
flow, approver routing through the shared org chain, the advisory (never
blocking) coverage warnings, in-handler RBAC, and both annex exports.
"""
import io
import sqlite3
from datetime import date, timedelta

import pytest

import mailer
import storage as st
import tzutil as tz
import work_arrangement_annex as wanx


# --------------------------------------------------------------- fixtures ---

@pytest.fixture()
def wa_org(app_mod):
    """A section with a chief, SIX staff (so the '4 of 6' coverage copy is
    exercised with real numbers), a division chief above it, and a loner with no
    directory link. Fresh rows per test; work_arrangement is wiped so period
    occupancy and coverage start from zero."""
    conn = st.connect(app_mod.DB)
    st.init_db(conn)
    conn.execute("DELETE FROM work_arrangement_noawa")
    conn.execute("DELETE FROM work_arrangement")
    conn.execute("DELETE FROM work_arrangement_reopen")
    conn.execute("DELETE FROM admin_user WHERE username LIKE 'wa-%@test.local'")
    conn.execute("DELETE FROM resource_person WHERE name LIKE 'WA-Test %'")
    for path in ("TEST/WA/WASec", "TEST/WA"):
        conn.execute("DELETE FROM org_unit WHERE path = ?", (path,))
    div = conn.execute(
        "INSERT INTO org_unit(name, short_name, level, path) "
        "VALUES ('WA Test Division','WADiv','division','TEST/WA')").lastrowid
    unit = conn.execute(
        "INSERT INTO org_unit(name, short_name, level, path) "
        "VALUES ('WA Test Section','WASec','section','TEST/WA/WASec')").lastrowid

    def person(label, org, position, designation=None, role="viewer"):
        rp = conn.execute(
            "INSERT INTO resource_person(name, is_internal, org_unit_id, position, "
            "designation) VALUES (?, 1, ?, ?, ?)",
            (f"WA-Test {label}", org, position, designation)).lastrowid
        return st.create_admin_user(conn, f"wa-{label.lower()}@test.local",
                                    "x", "y", role, rp_id=rp)

    dchief = person("DivChief", div, "Division Chief", "Division Chief",
                    "section_manager")
    chief = person("Chief", unit, "CHPO", "Section Head", "section_manager")
    staff = [person(f"Staff{i}", unit, "HPO II") for i in range(6)]
    loner = st.create_admin_user(conn, "wa-loner@test.local", "x", "y", "viewer")
    conn.commit()
    # A REAL admin account, not a synthetic {"id": 0} — decided_by is a foreign
    # key, so a decision has to be attributable to an account that exists.
    admin_row = st.get_admin_user(conn, "admin@test.local")
    yield {"conn": conn, "unit": unit, "div": div, "path": "TEST/WA/WASec",
           "dchief": dchief, "chief": chief, "staff": staff, "loner": loner,
           "admin": {"id": admin_row["id"], "role": "admin"}}
    conn.execute("DELETE FROM work_arrangement_noawa")
    conn.execute("DELETE FROM work_arrangement")
    conn.execute("DELETE FROM work_arrangement_reopen")
    conn.execute("DELETE FROM admin_user WHERE username LIKE 'wa-%@test.local'")
    conn.execute("DELETE FROM resource_person WHERE name LIKE 'WA-Test %'")
    for path in ("TEST/WA/WASec", "TEST/WA"):
        conn.execute("DELETE FROM org_unit WHERE path = ?", (path,))
    conn.commit()
    conn.close()


def _client_for_id(app_mod, uid):
    from starlette.testclient import TestClient
    import auth
    conn = st.connect(app_mod.DB)
    u = st.get_admin_user_by_id(conn, uid)
    conn.close()
    c = TestClient(app_mod.app, raise_server_exceptions=True)
    c.cookies.set(auth.COOKIE_NAME, auth.make_token(
        {"id": u["id"], "username": u["username"], "role": u["role"]}))
    return c


def _open(n=0):
    return st.work_arrangement_open_periods()[n]


def _u(uid, role="viewer"):
    return {"id": uid, "role": role}


# --------------------------------------------------- periods and deadlines ---

def test_period_key_splits_the_month_at_the_16th():
    assert st.work_arrangement_period(date(2026, 8, 1)) == "2026-08-01"
    assert st.work_arrangement_period(date(2026, 8, 15)) == "2026-08-01"
    assert st.work_arrangement_period(date(2026, 8, 16)) == "2026-08-16"
    assert st.work_arrangement_period(date(2026, 8, 31)) == "2026-08-16"


@pytest.mark.parametrize("ps,last,label", [
    ("2026-08-01", 15, "August 1-15, 2026"),
    ("2026-08-16", 31, "August 16-31, 2026"),      # the source sheets' titles
    ("2026-09-16", 30, "September 16-30, 2026"),
    ("2026-02-16", 28, "February 16-28, 2026"),
    ("2028-02-16", 29, "February 16-29, 2028"),    # leap year
])
def test_period_bounds_and_label(ps, last, label):
    start, end = st.work_arrangement_period_bounds(ps)
    assert start.isoformat() == ps
    assert end.day == last
    assert st.work_arrangement_period_label(ps) == label


def test_periods_step_across_a_year_boundary():
    assert st.work_arrangement_next_period("2026-12-16") == "2027-01-01"
    assert st.work_arrangement_prev_period("2027-01-01") == "2026-12-16"
    for ps in ("2026-08-01", "2026-08-16", "2026-12-16", "2027-01-01"):
        assert st.work_arrangement_prev_period(
            st.work_arrangement_next_period(ps)) == ps


@pytest.mark.parametrize("ps,deadline", [
    ("2026-08-16", "2026-08-14"),      # period starts Sunday -> Friday before
    ("2026-09-01", "2026-08-28"),      # period starts Tuesday
    ("2027-01-01", "2026-12-25"),      # period starts ON a Friday -> Friday PRIOR
])
def test_deadline_is_the_friday_before_the_period(ps, deadline):
    d = st.work_arrangement_deadline(ps)
    assert d.isoformat() == deadline
    assert d.weekday() == 4                                  # always a Friday
    assert d < date.fromisoformat(ps)                        # always strictly before


def test_filing_window_closes_the_day_after_the_deadline(monkeypatch):
    """The deadline Friday itself is still open; the Saturday after is not, and
    a period that has already begun is never open."""
    def at(d):
        monkeypatch.setattr(tz, "today_local", lambda: d)
        return st.work_arrangement_open_periods()

    assert "2026-08-16" in at(date(2026, 8, 13))             # Thursday before
    assert "2026-08-16" in at(date(2026, 8, 14))             # the deadline itself
    assert "2026-08-16" not in at(date(2026, 8, 15))         # the day after
    # ...and once inside the period it stays shut.
    assert at(date(2026, 8, 20))[0] == "2026-09-01"


@pytest.mark.parametrize("today,shows", [
    (date(2026, 8, 13), "2026-08-01"),   # Thursday, mid-period
    (date(2026, 8, 14), "2026-08-01"),   # the period's LAST workday — still it
    (date(2026, 8, 15), "2026-08-16"),   # Saturday, nothing left to work -> next
    (date(2026, 8, 16), "2026-08-16"),   # Sunday, already the new period
    (date(2026, 8, 29), "2026-08-16"),   # Saturday, but Mon the 31st is to come
    (date(2026, 8, 30), "2026-08-16"),   # Sunday, likewise
    (date(2026, 8, 31), "2026-08-16"),   # Monday, the period's last workday
    (date(2026, 9, 1), "2026-09-01"),
])
def test_the_schedule_shown_rolls_forward_once_a_period_is_worked_out(today, shows):
    """A period stops being the one to show the moment it has no workdays left
    — on Saturday the 15th nobody works another day under the 1st-15th
    arrangement. The CALENDAR period is deliberately unmoved by this."""
    assert st.work_arrangement_schedule_period(today) == shows


def test_the_calendar_period_is_not_rolled_forward(monkeypatch):
    """Filing deadlines, "No AWA" intervals and the annex all key on the
    calendar period; rolling it early would change what people may file."""
    _at(monkeypatch, date(2026, 8, 15))
    assert st.work_arrangement_period() == "2026-08-01"
    assert st.work_arrangement_schedule_period() == "2026-08-16"


# ----------------------------------------------------- field validation ------

def test_reg_takes_no_weekday(wa_org):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    wa = st.create_work_arrangement(conn, emp, _open(), "REG", weekday=3)
    assert st.get_work_arrangement(conn, wa)["weekday"] is None
    assert st.get_work_arrangement(conn, wa)["remarks"] == ""


def test_day_types_require_a_weekday(wa_org):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    for arrangement in ("CWW", "WFH"):
        with pytest.raises(ValueError, match="Monday to Friday"):
            st.create_work_arrangement(conn, emp, _open(), arrangement)


@pytest.mark.parametrize("weekday", [5, 6, -1, 99])
def test_weekday_must_be_within_the_work_week(wa_org, weekday):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    with pytest.raises(ValueError, match="work week"):
        st.create_work_arrangement(conn, emp, _open(), "CWW", weekday)


def test_unknown_arrangement_is_rejected(wa_org):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    with pytest.raises(ValueError, match="work arrangement"):
        st.create_work_arrangement(conn, emp, _open(), "HYBRID", 1)


@pytest.mark.parametrize("arrangement,weekday,note,expected", [
    ("CWW", 4, None, "Friday Day off"),
    ("WFH", 2, None, "Wednesday WFH"),
    ("REG", None, None, ""),
    ("CWW", 4, "for approval from OSEC", "Friday Day off — for approval from OSEC"),
])
def test_remarks_match_the_annex_wording(arrangement, weekday, note, expected):
    assert st._wa_remarks(arrangement, weekday, note) == expected


def test_filing_for_a_closed_period_names_the_next_one(wa_org):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    closed = st.work_arrangement_prev_period(_open())
    with pytest.raises(ValueError) as e:
        st.create_work_arrangement(conn, emp, closed, "REG")
    msg = str(e.value)
    assert "deadline" in msg
    assert st.work_arrangement_period_label(closed) in msg
    # a refusal has to say what to do instead
    assert st.work_arrangement_period_label(_open()) in msg


def test_filing_too_far_ahead_is_refused(wa_org):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    far = _open(1)
    for _ in range(2):
        far = st.work_arrangement_next_period(far)
    with pytest.raises(ValueError, match="isn't open for filing yet"):
        st.create_work_arrangement(conn, emp, far, "REG")


# ------------------------------------------------------------ uniqueness ----

def test_one_live_arrangement_per_person_per_period(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    st.create_work_arrangement(conn, emp, ps, "WFH", 4)
    with pytest.raises(ValueError, match="waiting for approval"):
        st.create_work_arrangement(conn, emp, ps, "CWW", 0)


def test_the_partial_index_backstops_a_double_submit(wa_org):
    """The Python pre-check says it in English; the index is what actually
    holds if two submits race."""
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    st.create_work_arrangement(conn, emp, ps, "WFH", 4)
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute(
            "INSERT INTO work_arrangement(user_id, period_start, arrangement, "
            "weekday, status) VALUES (?,?,'CWW',0,'pending')", (emp, ps))
    conn.rollback()
    # ...but a decided row is outside the constraint, so history accumulates.
    conn.execute(
        "INSERT INTO work_arrangement(user_id, period_start, arrangement, "
        "weekday, status) VALUES (?,?,'CWW',0,'denied')", (emp, ps))
    conn.commit()


def test_withdrawing_or_being_denied_frees_the_period(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    first = st.create_work_arrangement(conn, emp, ps, "WFH", 4)
    st.cancel_work_arrangement(conn, first, emp)
    assert st.my_pending_work_arrangement(conn, emp, ps) is None

    second = st.create_work_arrangement(conn, emp, ps, "CWW", 0)
    st.decide_work_arrangement(conn, second, _u(wa_org["chief"], "section_manager"), "deny")
    third = st.create_work_arrangement(conn, emp, ps, "REG")
    assert st.get_work_arrangement(conn, third)["status"] == "pending"


def test_only_a_pending_request_can_be_withdrawn(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    wa = st.create_work_arrangement(conn, emp, ps, "REG")
    with pytest.raises(ValueError, match="isn't your work arrangement"):
        st.cancel_work_arrangement(conn, wa, wa_org["staff"][1])
    st.decide_work_arrangement(conn, wa, _u(wa_org["chief"], "section_manager"), "approve")
    with pytest.raises(ValueError, match="waiting for approval"):
        st.cancel_work_arrangement(conn, wa, emp)


# ------------------------------------------- one AWA day per calendar week ---
#
# Half-month periods end mid-week, so the week of September 14-18, 2026 is split
# down the middle: Monday the 14th is the last week of September 1-15 and
# Wednesday the 16th is the first of September 16-30. One filing per period is
# therefore NOT one arrangement per week, and these cover the gap.

_SEPT_BOTH_OPEN = date(2026, 8, 28)      # both September periods accepting filings
_SEPT_1_15, _SEPT_16_30 = "2026-09-01", "2026-09-16"


@pytest.fixture()
def sept(wa_org, monkeypatch):
    """Positioned so both halves of September can be filed for at once."""
    _at(monkeypatch, _SEPT_BOTH_OPEN)
    assert st.work_arrangement_open_periods() == [_SEPT_1_15, _SEPT_16_30]
    return wa_org


def test_the_two_september_periods_share_one_working_week():
    """The premise the rest of this section rests on: Monday the 14th and
    Wednesday the 16th are different periods but the same week."""
    assert st.work_arrangement_weeks(_SEPT_1_15, 0)[1] == date(2026, 9, 14)
    assert st.work_arrangement_weeks(_SEPT_16_30, 2)[0] == date(2026, 9, 14)


@pytest.mark.parametrize("weekday,day", [(2, 16), (3, 17), (4, 18)])
def test_a_monday_wfh_pushes_the_later_filing_past_the_shared_week(sept, weekday, day):
    """Monday the 14th is already work-from-home under the EARLIER period, so
    the 16th, 17th and 18th filing doesn't get refused — it just doesn't start
    until the week after, since September 1-15 keeps the week of the 14th."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    wa = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", weekday)
    row = st.get_work_arrangement(conn, wa)
    assert row["effective_from"] == "2026-09-21"          # the Monday after
    assert row["week_override_note"]
    assert "September 14-18, 2026" in row["week_override_note"]
    assert "Monday" in row["week_override_note"]
    assert "September 21-25, 2026" in row["week_override_note"]
    # Still goes on to approve normally — the resolution isn't a permanent
    # asterisk on the row, just its starting date.
    assert st.decide_work_arrangement(
        conn, wa, _u(sept["chief"], "section_manager"), "approve")["status"] == "approved"


def test_the_pushed_effective_date_shows_up_in_the_day_view(sept):
    """The one place the boundary week actually matters day to day: the 16th
    itself reads as an ordinary on-site day, and the 23rd reads as WFH."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    wa = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)   # Wed
    st.decide_work_arrangement(conn, wa, _u(sept["chief"], "section_manager"), "approve")
    day16 = st.work_arrangement_day(conn, date(2026, 9, 16))
    day23 = st.work_arrangement_day(conn, date(2026, 9, 23))
    name = "WA-Test Staff0"
    assert name in [p["name"] for p in day16["groups"]["onsite"]]
    assert name in [p["name"] for p in day23["groups"]["wfh"]]


@pytest.mark.parametrize("weekday", [0, 1, 2, 3, 4])
def test_the_week_after_the_boundary_is_free_again(sept, weekday):
    """The cap is per week, not per fortnight: September 21 onward is untouched
    by a Monday the 14th filing, whichever weekday is picked."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 4)   # Friday 11th
    wa = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", weekday)
    assert st.get_work_arrangement(conn, wa)["weekday"] == weekday


def test_a_day_that_stops_short_of_the_boundary_week_does_not_reserve_it(sept):
    """Wednesdays in September 1-15 run out on the 9th, so that filing never
    reaches the week of the 14th and a Wednesday the 16th is a first day, not a
    second. The rule counts the days actually worked, not the periods."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 2)
    assert st.work_arrangement_weeks(_SEPT_1_15, 2)[1] == date(2026, 9, 7)
    st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)


def test_a_compressed_day_off_counts_against_the_same_week(sept):
    """Both alternative arrangements share the cap, and against each other — a
    day off and a day at home in one week is still two, so the later one still
    gets pushed rather than treated as unrelated."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "CWW", 0)
    wa = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)
    assert st.get_work_arrangement(conn, wa)["effective_from"] == "2026-09-21"


def test_a_regular_week_is_exempt_at_both_ends(sept):
    """REG is not an alternative arrangement: it neither takes a week nor is
    refused one."""
    conn, emp, other = sept["conn"], sept["staff"][0], sept["staff"][1]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    st.create_work_arrangement(conn, emp, _SEPT_16_30, "REG")      # not blocked
    st.create_work_arrangement(conn, other, _SEPT_1_15, "REG")
    st.create_work_arrangement(conn, other, _SEPT_16_30, "WFH", 2)  # takes nothing


def test_the_cap_is_per_person(sept):
    conn = sept["conn"]
    st.create_work_arrangement(conn, sept["staff"][0], _SEPT_1_15, "WFH", 0)
    st.create_work_arrangement(conn, sept["staff"][1], _SEPT_16_30, "WFH", 2)


def test_a_filing_still_waiting_for_approval_holds_the_week(sept):
    """Pending occupies the week just like approved does — the later filing
    still defers to it rather than being treated as free to claim it."""
    conn, emp = sept["conn"], sept["staff"][0]
    wa = st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    assert st.get_work_arrangement(conn, wa)["status"] == "pending"
    later = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)
    assert st.get_work_arrangement(conn, later)["effective_from"] == "2026-09-21"


@pytest.mark.parametrize("release", ["withdraw", "deny"])
def test_withdrawing_or_denying_gives_the_week_back(sept, release):
    conn, emp = sept["conn"], sept["staff"][0]
    wa = st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    if release == "withdraw":
        st.cancel_work_arrangement(conn, wa, emp)
    else:
        st.decide_work_arrangement(conn, wa, _u(sept["chief"], "section_manager"), "deny")
    st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)


def test_the_override_note_names_the_week_the_day_and_the_other_period(sept):
    """A note that doesn't say what is in the way and what changed is a dead
    end — the same standard the old refusal was held to."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    wa = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)
    msg = st.get_work_arrangement(conn, wa)["week_override_note"]
    assert "September 14-18, 2026" in msg          # the week they share
    assert "Monday" in msg                         # what is already there
    assert "September 1-15, 2026" in msg           # filed under which period
    assert "September 21-25, 2026" in msg          # and when this one starts


def test_the_approver_reads_the_override_note_about_the_filer_with_a_link(sept, app_mod):
    """In the approvals queue the note is about SOMEONE ELSE's arrangement —
    "your Monday arrangement" pointed at the approver — and it links to the
    period that keeps the shared week."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    wa = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)
    row = st.get_work_arrangement(conn, wa)
    assert "your Monday" in row["week_override_note"]          # the filer's own view
    assert "your" not in row["week_override_note_approver"]
    assert f"{row['person_name']}'s Monday" in row["week_override_note_approver"]
    assert row["week_override_period"]["start"] == _SEPT_1_15

    html = _client_for_id(app_mod, sept["chief"]).get("/work-arrangement/approvals").text
    assert row["week_override_note_approver"].replace("'", "&#39;") in html
    assert f'href="/work-arrangement/schedule?period={_SEPT_1_15}"' in html


def test_the_blocked_direction_names_the_week_the_day_and_the_other_period(sept):
    """Filing the EARLIER period after the later one is already approved can't
    be resolved automatically — there's no later date within its own period to
    push it to — so this still refuses, and still says what is in the way."""
    conn, emp = sept["conn"], sept["staff"][0]
    wa = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)
    st.decide_work_arrangement(conn, wa, _u(sept["chief"], "section_manager"), "approve")
    with pytest.raises(ValueError) as e:
        st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    msg = str(e.value)
    assert "September 14-18, 2026" in msg          # the week they collide in
    assert "Wednesday WFH" in msg                  # what is already there
    assert "Wednesday, September 16, 2026" in msg  # and on which day
    assert "September 16-30, 2026" in msg          # filed under which period
    assert "already approved" in msg               # and how far along it is
    assert "can't take priority" in msg            # why this one can't win it


def test_a_change_is_judged_on_the_day_it_asks_for(sept):
    """A change can walk into a boundary week the filed day stayed clear of."""
    conn, emp = sept["conn"], sept["staff"][0]
    wa = st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 2)   # Wednesday
    st.decide_work_arrangement(conn, wa, _u(sept["chief"], "section_manager"), "approve")
    st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)       # fine
    # Mondays AND Tuesdays in September 1-15 both reach the week of the 14th —
    # and this filing IS the earlier period, so there's no later date within it
    # to push the change to; it still has to refuse.
    for weekday in (0, 1):
        with pytest.raises(ValueError, match="can't take priority"):
            st.request_work_arrangement_change(conn, wa, emp, "WFH", weekday)
    # Thursdays run out on the 10th, so that change still goes through
    assert st.request_work_arrangement_change(conn, wa, emp, "WFH", 3)


def test_a_change_may_still_reuse_the_week_its_own_filing_holds(sept):
    """The row a change replaces sits in the same period and must not be read as
    the clash — otherwise no arrangement could ever be changed."""
    conn, emp = sept["conn"], sept["staff"][0]
    wa = st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    st.decide_work_arrangement(conn, wa, _u(sept["chief"], "section_manager"), "approve")
    assert st.request_work_arrangement_change(conn, wa, emp, "CWW", 0)


def _grandfather(conn, user_id, period_start, arrangement, weekday, status):
    """A row from before the weekly rule existed: written straight past the
    trigger, the way the pre-rule code would have written it."""
    conn.execute("DROP TRIGGER trg_wa_one_per_week_ins")
    conn.execute(
        "INSERT INTO work_arrangement(user_id, period_start, arrangement, "
        "weekday, status, created_at) VALUES (?,?,?,?,?,?)",
        (user_id, period_start, arrangement, weekday, status, tz.now_iso()))
    conn.commit()
    st._wa_install_week_guard(conn)      # put the guard back


def test_arrangements_filed_before_the_rule_are_left_alone(sept):
    """Two days in one week already on the books stay exactly as they are —
    the rule judges what is written from here on, not what people already hold."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    _grandfather(conn, emp, _SEPT_16_30, "WFH", 2, "approved")
    assert st.my_work_arrangement_for(conn, emp, _SEPT_16_30)["weekday"] == 2
    assert st.my_pending_work_arrangement(conn, emp, _SEPT_1_15)["weekday"] == 0


def test_an_approval_that_would_double_up_a_week_is_refused(sept):
    """The second gate. Filing normally catches this, so it takes a row that
    predates the rule to reach an approver — and the approver is told rather
    than left to sign off a week the annex would print twice."""
    conn, emp = sept["conn"], sept["staff"][0]
    wa = st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    _grandfather(conn, emp, _SEPT_16_30, "WFH", 2, "approved")
    with pytest.raises(ValueError, match="can't take priority"):
        st.decide_work_arrangement(conn, wa, _u(sept["chief"], "section_manager"), "approve")
    # denying it is still allowed — that is how the clash gets cleared
    assert st.decide_work_arrangement(
        conn, wa, _u(sept["chief"], "section_manager"), "deny")["status"] == "denied"


def test_the_database_refuses_it_even_with_the_checks_bypassed(sept):
    """The backstop: the same rule as a trigger, so a write path that forgets to
    ask cannot leave two days in one week behind."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    with pytest.raises(sqlite3.IntegrityError, match="one alternative work"):
        conn.execute(
            "INSERT INTO work_arrangement(user_id, period_start, arrangement, "
            "weekday, status, created_at) VALUES (?,?,?,?,'pending',?)",
            (emp, _SEPT_16_30, "WFH", 2, tz.now_iso()))
    conn.rollback()


def test_the_database_refuses_a_status_flip_into_a_taken_week(sept):
    """Not just INSERT: a denied row nudged back to approved would slip in
    through UPDATE, so the trigger covers that too."""
    conn, emp = sept["conn"], sept["staff"][0]
    wa = st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)
    st.decide_work_arrangement(conn, wa, _u(sept["chief"], "section_manager"), "deny")
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    with pytest.raises(sqlite3.IntegrityError, match="one alternative work"):
        conn.execute("UPDATE work_arrangement SET status='approved' WHERE id=?",
                     (wa,))
    conn.rollback()


@pytest.mark.parametrize("ps", ["2026-01-01", "2026-02-16", "2026-09-01",
                                "2026-12-16", "2028-02-16"])
@pytest.mark.parametrize("weekday", [0, 1, 2, 3, 4])
def test_the_stored_week_bounds_agree_with_the_calendar(wa_org, ps, weekday):
    """The two GENERATED columns the trigger compares are derived in SQL; this
    holds them to the same answer the Python arithmetic gives, including over
    year ends and a leap February."""
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    conn.execute("DROP TRIGGER trg_wa_one_per_week_ins")
    conn.execute(
        "INSERT INTO work_arrangement(user_id, period_start, arrangement, "
        "weekday, status, created_at) VALUES (?,?,?,?,'pending',?)",
        (emp, ps, "WFH", weekday, tz.now_iso()))
    got = conn.execute("SELECT week_head, week_tail FROM work_arrangement "
                       "WHERE user_id=? AND period_start=?",
                       (emp, ps)).fetchone()
    conn.rollback()
    st._wa_install_week_guard(conn)
    head, tail = st.work_arrangement_weeks(ps, weekday)
    assert got == (head.isoformat(), tail.isoformat())
    assert head.weekday() == 0 and tail.weekday() == 0        # always Mondays


def test_the_week_label_reads_like_the_period_labels():
    assert st.work_arrangement_week_label(date(2026, 9, 14)) == "September 14-18, 2026"
    # a week that straddles two months names both
    assert st.work_arrangement_week_label(date(2026, 8, 31)) == \
        "August 31 - September 4, 2026"


def test_a_resolved_boundary_week_files_and_carries_the_note_to_the_redirect(
        sept, app_mod):
    """The later-period filing over HTTP: no refusal, and the explanation rides
    along on the redirect as a distinct info notice rather than being lost."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_1_15, "WFH", 0)
    c = _client_for_id(app_mod, emp)
    r = c.post("/work-arrangement/file",
               data={"period_start": _SEPT_16_30, "arrangement": "WFH",
                     "weekday": "2"}, follow_redirects=False)
    assert r.status_code == 303
    assert "info=" in r.headers["location"]
    assert st.my_pending_work_arrangement(conn, emp, _SEPT_16_30) is not None


def test_the_filing_form_says_why_a_blocked_overlap_was_refused(sept, app_mod):
    """The direction that still can't be resolved automatically (the EARLIER
    period filed after the later one) reaches the person at the form, not as a
    500."""
    conn, emp = sept["conn"], sept["staff"][0]
    st.create_work_arrangement(conn, emp, _SEPT_16_30, "WFH", 2)
    c = _client_for_id(app_mod, emp)
    r = c.post("/work-arrangement/file",
               data={"period_start": _SEPT_1_15, "arrangement": "WFH",
                     "weekday": "0"})
    assert r.status_code == 400
    assert "September 14-18, 2026" in r.text
    assert st.my_pending_work_arrangement(conn, emp, _SEPT_1_15) is None


# ------------------------------------------------ staff-initiated changes ----

def _approved(org, uid, arrangement="WFH", weekday=4, ps=None):
    conn, ps = org["conn"], ps or _open()
    wa = st.create_work_arrangement(conn, uid, ps, arrangement, weekday)
    st.decide_work_arrangement(conn, wa, _u(org["chief"], "section_manager"), "approve")
    return wa


def test_a_change_leaves_the_original_in_force_until_decided(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    orig = _approved(wa_org, emp)
    chg = st.request_work_arrangement_change(conn, orig, emp, "CWW", 0)

    assert st.get_work_arrangement(conn, chg)["supersedes_id"] == orig
    assert st.get_work_arrangement(conn, chg)["is_change"]
    # the schedule in force has not moved
    assert st.my_work_arrangement_for(conn, emp, ps)["id"] == orig
    assert st.get_work_arrangement(conn, orig)["status"] == "approved"


def test_approving_a_change_supersedes_the_original(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    orig = _approved(wa_org, emp)
    chg = st.request_work_arrangement_change(conn, orig, emp, "CWW", 0)
    st.decide_work_arrangement(conn, chg, _u(wa_org["chief"], "section_manager"), "approve")

    assert st.get_work_arrangement(conn, orig)["status"] == "superseded"
    live = st.my_work_arrangement_for(conn, emp, ps)
    assert live["id"] == chg and live["remarks"] == "Monday Day off"
    # exactly one approved row survives for the period
    assert conn.execute(
        "SELECT COUNT(*) FROM work_arrangement WHERE user_id=? AND period_start=? "
        "AND status='approved'", (emp, ps)).fetchone()[0] == 1


@pytest.mark.parametrize("action", ["deny", "cancel"])
def test_a_rejected_change_leaves_the_original_standing(wa_org, action):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    orig = _approved(wa_org, emp)
    chg = st.request_work_arrangement_change(conn, orig, emp, "CWW", 0)
    if action == "deny":
        st.decide_work_arrangement(conn, chg, _u(wa_org["chief"], "section_manager"), "deny")
    else:
        st.cancel_work_arrangement(conn, chg, emp)

    assert st.my_work_arrangement_for(conn, emp, ps)["id"] == orig
    assert st.my_work_arrangement_for(conn, emp, ps)["remarks"] == "Friday WFH"


def test_only_the_owner_may_start_a_change(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    orig = _approved(wa_org, emp)
    # not the section chief, and not an admin — a schedule is the person's own
    with pytest.raises(PermissionError):
        st.request_work_arrangement_change(conn, orig, wa_org["chief"], "REG")


def test_a_pending_request_is_withdrawn_not_changed(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    wa = st.create_work_arrangement(conn, emp, ps, "WFH", 4)
    with pytest.raises(ValueError, match="withdraw it and file again"):
        st.request_work_arrangement_change(conn, wa, emp, "REG")


def test_one_change_in_flight_at_a_time_and_no_no_ops(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    orig = _approved(wa_org, emp)
    with pytest.raises(ValueError, match="already have"):
        st.request_work_arrangement_change(conn, orig, emp, "WFH", 4)
    st.request_work_arrangement_change(conn, orig, emp, "CWW", 0)
    with pytest.raises(ValueError, match="already have a change waiting"):
        st.request_work_arrangement_change(conn, orig, emp, "REG")


def test_a_change_after_the_deadline_is_refused(wa_org, monkeypatch):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    orig = _approved(wa_org, emp)
    past = st.work_arrangement_deadline(ps) + timedelta(days=1)
    monkeypatch.setattr(tz, "today_local", lambda: past)
    with pytest.raises(ValueError, match="deadline"):
        st.request_work_arrangement_change(conn, orig, emp, "CWW", 0)


# ------------------------------------------------- admin reopen override ----

def test_an_admin_reopen_lets_filing_past_the_deadline(wa_org, monkeypatch):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    past = st.work_arrangement_deadline(ps) + timedelta(days=1)
    monkeypatch.setattr(tz, "today_local", lambda: past)
    with pytest.raises(ValueError, match="deadline"):
        st.create_work_arrangement(conn, emp, ps, "REG")

    st.set_work_arrangement_reopen(conn, ps, wa_org["admin"]["id"], "late request")
    assert st.get_work_arrangement_reopen(conn, ps)["note"] == "late request"
    assert [r["period_start"] for r in st.list_work_arrangement_reopens(conn)] == [ps]

    wa_id = st.create_work_arrangement(conn, emp, ps, "REG")
    assert st.get_work_arrangement(conn, wa_id)["period_start"] == ps

    st.clear_work_arrangement_reopen(conn, ps)
    assert st.get_work_arrangement_reopen(conn, ps) is None
    with pytest.raises(ValueError, match="deadline"):
        st.create_work_arrangement(conn, wa_org["staff"][1], ps, "REG")


# ------------------------------------------------------ approver routing ----

def test_escalation_uses_the_shared_org_chain(wa_org):
    conn = wa_org["conn"]
    assert st.get_approvers_for(conn, wa_org["staff"][0]) == [wa_org["chief"]]
    # a section chief escalates to the division chief, never to themselves
    assert st.get_approvers_for(conn, wa_org["chief"]) == [wa_org["dchief"]]
    # nobody off the chain resolves; admins are the fallback
    assert st.get_approvers_for(conn, wa_org["loner"]) == []
    assert st.is_approver_for(conn, wa_org["admin"], wa_org["loner"])


def test_the_wellness_names_are_the_same_functions():
    """Phase 0 bound the old names to the shared implementations rather than
    copying them. If someone un-aliases these, the two modules drift apart."""
    for old, new in [
        (st._wellness_rank, st._org_rank),
        (st._wellness_parent_path, st._org_parent_path),
        (st._wellness_accounts_at_rank, st._org_accounts_at_rank),
        (st._wellness_role_filled, st._org_role_filled),
        (st._wellness_seat_name, st._org_seat_name),
        (st._wellness_dashboard_scope, st._org_scope_for_user),
        (st._wellness_in_scope, st._path_in_scope),
        (st.wellness_approver_name, st.approver_name_for),
        (st.get_wellness_approvers, st.get_approvers_for),
        (st.is_wellness_approver, st.is_approver_for),
        (st.user_is_wellness_chief, st.user_is_approving_chief),
        (st.wellness_approver_title, st.approver_title_for),
    ]:
        assert old is new


def test_queue_is_scoped_to_what_the_chain_routes(wa_org):
    conn, ps = wa_org["conn"], _open()
    mine = st.create_work_arrangement(conn, wa_org["staff"][0], ps, "REG")
    chiefs_own = st.create_work_arrangement(conn, wa_org["chief"], ps, "REG")
    st.create_work_arrangement(conn, wa_org["loner"], ps, "REG")

    queue = st.list_work_arrangement_approvals(
        conn, _u(wa_org["chief"], "section_manager"))
    ids = [w["id"] for w in queue]
    assert mine in ids
    assert chiefs_own not in ids                     # escalates one level further
    # ...and it lands with the division chief instead
    dq = [w["id"] for w in st.list_work_arrangement_approvals(
        conn, _u(wa_org["dchief"], "section_manager"))]
    assert chiefs_own in dq and mine not in dq
    # a plain viewer sees no queue at all
    assert st.list_work_arrangement_approvals(conn, _u(wa_org["staff"][1])) == []


def test_decide_guards_state_and_authority(wa_org):
    conn, ps = wa_org["conn"], _open()
    chief = _u(wa_org["chief"], "section_manager")
    wa = st.create_work_arrangement(conn, wa_org["staff"][0], ps, "REG")
    with pytest.raises(PermissionError):
        st.decide_work_arrangement(conn, wa, _u(wa_org["staff"][1]), "approve")
    with pytest.raises(ValueError, match="Unknown decision"):
        st.decide_work_arrangement(conn, wa, chief, "maybe")

    row = st.decide_work_arrangement(conn, wa, chief, "approve",
                                     "keeps Tuesday covered")
    assert row["status"] == "approved"
    assert row["decided_at"] and row["decide_note"] == "keeps Tuesday covered"
    with pytest.raises(ValueError, match="already been decided"):
        st.decide_work_arrangement(conn, wa, chief, "deny")


def test_admin_cannot_decide_a_request_that_has_a_real_chief(wa_org):
    """An admin is the fallback ABOVE the chain, for requests nobody in it can
    reach — never a bypass around a section that already has an approver.
    wa_org["staff"][0] has a real section chief, so an admin with no seat of
    their own in that section may not decide it, even though the same admin
    correctly decides the loner's request (an orphan, per
    test_escalation_uses_the_shared_org_chain)."""
    conn, ps = wa_org["conn"], _open()
    wa = st.create_work_arrangement(conn, wa_org["staff"][0], ps, "REG")
    assert not st.is_approver_for(conn, wa_org["admin"], wa_org["staff"][0])
    with pytest.raises(PermissionError):
        st.decide_work_arrangement(conn, wa, wa_org["admin"], "approve")
    assert st.get_work_arrangement(conn, wa)["status"] == "pending"

    orphan_wa = st.create_work_arrangement(conn, wa_org["loner"], ps, "REG")
    row = st.decide_work_arrangement(conn, orphan_wa, wa_org["admin"], "approve")
    assert row["status"] == "approved"


# ---------------------------------------------------------------- coverage ---

def test_coverage_counts_cww_and_wfh_apart(wa_org):
    conn, ps, path = wa_org["conn"], _open(), wa_org["path"]
    for i in range(3):
        _approved(wa_org, wa_org["staff"][i], "CWW", 4)
    _approved(wa_org, wa_org["staff"][3], "WFH", 4)

    cov = st.work_arrangement_coverage(conn, path, ps)
    fri = cov["weekdays"][4]
    assert (fri["cww"], fri["wfh"], fri["offsite"]) == (3, 1, 4)
    assert cov["headcount"] == 7                     # 6 staff + the section chief
    assert fri["onsite"] == 3
    assert cov["weekdays"][0]["offsite"] == 0        # Monday untouched


def test_coverage_ignores_denied_and_withdrawn_rows(wa_org):
    conn, ps, path = wa_org["conn"], _open(), wa_org["path"]
    _approved(wa_org, wa_org["staff"][0], "CWW", 4)
    denied = st.create_work_arrangement(conn, wa_org["staff"][1], ps, "CWW", 4)
    st.decide_work_arrangement(conn, denied, _u(wa_org["chief"], "section_manager"), "deny")
    withdrawn = st.create_work_arrangement(conn, wa_org["staff"][2], ps, "CWW", 4)
    st.cancel_work_arrangement(conn, withdrawn, wa_org["staff"][2])

    assert st.work_arrangement_coverage(conn, path, ps)["weekdays"][4]["offsite"] == 1


def test_a_pending_change_outranks_the_row_it_replaces(wa_org):
    """Coverage answers "what is being asked for", so an in-flight change is the
    state the chief is deciding on."""
    conn, ps, path = wa_org["conn"], _open(), wa_org["path"]
    emp = wa_org["staff"][0]
    orig = _approved(wa_org, emp, "CWW", 4)
    st.request_work_arrangement_change(conn, orig, emp, "CWW", 0)

    cov = st.work_arrangement_coverage(conn, path, ps)
    assert cov["weekdays"][4]["offsite"] == 0        # no longer off on Friday
    assert cov["weekdays"][0]["offsite"] == 1        # off on Monday instead


def test_coverage_warns_but_never_blocks(wa_org):
    """The whole point of the rule: a thin day is advice, not a refusal."""
    conn, ps, path = wa_org["conn"], _open(), wa_org["path"]
    for i in range(3):
        _approved(wa_org, wa_org["staff"][i], "CWW", 4)
    _approved(wa_org, wa_org["staff"][3], "WFH", 4)

    emp = wa_org["staff"][4]
    note = st.work_arrangement_coverage_note(conn, emp, ps, "CWW", 4)
    assert "Friday" in note and "4 of 6" in note

    # ...and it still goes through, and can still be approved.
    wa = st.create_work_arrangement(conn, emp, ps, "CWW", 4)
    assert st.decide_work_arrangement(
        conn, wa, _u(wa_org["chief"], "section_manager"), "approve")["status"] == "approved"
    assert st.work_arrangement_coverage(conn, path, ps)["weekdays"][4]["onsite"] == 2


def test_coverage_flag_escalates_to_nobody_on_site(wa_org):
    """The section is 7 (six staff + the chief). With five already off on
    Friday, the sixth filing leaves only the chief, and the chief filing leaves
    nobody at all."""
    conn, ps = wa_org["conn"], _open()
    for i in range(5):
        _approved(wa_org, wa_org["staff"][i], "CWW", 4)

    def flag_for(uid):
        wa = st.create_work_arrangement(conn, uid, ps, "CWW", 4)
        return st.work_arrangement_coverage_flag(
            conn, st.get_work_arrangement(conn, wa))

    assert flag_for(wa_org["staff"][5]) == "Only 1 on-site on Friday after this"
    assert flag_for(wa_org["chief"]) == "Approving leaves nobody on-site on Friday"


def test_coverage_is_silent_for_reg_and_the_undirectoried(wa_org):
    conn, ps = wa_org["conn"], _open()
    assert st.work_arrangement_coverage_note(
        conn, wa_org["staff"][0], ps, "REG", None) == ""
    assert st.work_arrangement_coverage_note(
        conn, wa_org["loner"], ps, "CWW", 4) == ""


# ------------------------------------------------------ roster and exports ---

def test_the_roster_shows_people_who_never_filed(wa_org):
    """The blanks are the information — the whole reason the export reads from
    admin_user outward rather than from work_arrangement inward."""
    conn, ps = wa_org["conn"], _open()
    _approved(wa_org, wa_org["staff"][0], "WFH", 4)

    d = st.work_arrangement_period_summary(conn, ps,
                                           _u(wa_org["chief"], "section_manager"))
    assert d["filed"] == 1 and d["headcount"] == 7
    assert "WA-Test Staff5" in d["unfiled"]
    assert d["thinnest"]["name"] == "Friday"


def test_the_system_admin_account_is_not_on_the_staff_roster(wa_org):
    """ray@gmail.com is the system administrator's check-in account, not staff:
    never a roster row (screen or ANNEX C export), never nudged to file."""
    conn, ps = wa_org["conn"], _open()
    conn.execute("DELETE FROM admin_user WHERE lower(username)='ray@gmail.com'")
    ray = st.create_admin_user(conn, "ray@gmail.com", "x", "y", "admin")
    conn.commit()
    try:
        assert ray not in {r["user_id"] for r in st.work_arrangement_roster(conn, ps)}
        assert "ray@gmail.com" not in {r["name"] for r in
                                       st.work_arrangement_export_rows(conn, ps)}
        assert ray not in {r["id"] for r in st.staff_without_work_arrangement(conn, ps)}
        week = st.wellness_week_start()
        assert ray not in {r["id"] for r in
                           st.staff_without_wellness_claim_this_week(conn, week)}
        # a real staff account is still there
        assert wa_org["staff"][0] in {r["user_id"] for r in
                                      st.work_arrangement_roster(conn, ps)}
    finally:
        conn.execute("DELETE FROM admin_user WHERE id=?", (ray,))
        conn.commit()


def test_export_prints_only_approved_rows(wa_org):
    conn, ps = wa_org["conn"], _open()
    _approved(wa_org, wa_org["staff"][0], "CWW", 4)
    st.create_work_arrangement(conn, wa_org["staff"][1], ps, "WFH", 2)   # pending

    rows = {r["name"]: r for r in st.work_arrangement_export_rows(conn, ps)}
    assert rows["WA-Test Staff0"]["arrangement"] == "CWW"
    assert rows["WA-Test Staff0"]["remarks"] == "Friday Day off"
    # pending and unfiled people are PRESENT, with both cells blank
    assert rows["WA-Test Staff1"]["arrangement"] == ""
    assert rows["WA-Test Staff1"]["remarks"] == ""
    assert rows["WA-Test Staff5"]["arrangement"] == ""
    assert [r["no"] for r in st.work_arrangement_export_rows(conn, ps)][:3] == [1, 2, 3]


def test_only_accounts_appear_on_the_annex(wa_org):
    """Product decision: staff without an app account are out of scope."""
    conn, ps = wa_org["conn"], _open()
    conn.execute(
        "INSERT INTO resource_person(name, is_internal, org_unit_id, position) "
        "VALUES ('WA-Test NoAccount', 1, ?, 'GIP')", (wa_org["unit"],))
    conn.commit()
    names = [r["name"] for r in st.work_arrangement_export_rows(conn, ps)]
    assert "WA-Test NoAccount" not in names


def test_prepared_by_override_still_reads_the_directory_title(wa_org):
    """Pinning who prepares the annex must not require re-typing their title:
    the printed title comes from their live directory record when they have
    one there, so it stays accurate even if their position changes later."""
    conn = wa_org["conn"]
    st.set_setting(conn, "wa_prepared_by_name", "WA-Test Chief")
    st.set_setting(conn, "wa_prepared_by_title", "Stale Title")
    try:
        sig = st.work_arrangement_signatories(conn)
        assert sig["prepared_by"]["name"] == "WA-Test Chief"
        assert sig["prepared_by"]["title"] == "CHPO"     # from resource_person
    finally:
        st.set_setting(conn, "wa_prepared_by_name", None)
        st.set_setting(conn, "wa_prepared_by_title", None)


def test_seed_wa_prepared_by_default_fires_once_and_never_overwrites(wa_org):
    conn = wa_org["conn"]
    conn.execute("DELETE FROM app_setting WHERE key IN (?, ?)",
                 (st._WA_PREPARED_BY_SEEDED, "wa_prepared_by_name"))
    conn.commit()
    try:
        assert st.seed_wa_prepared_by_default(conn) is True
        assert st.get_setting(conn, "wa_prepared_by_name") == "Alexander R. Medel"

        # A later admin override survives the next boot's re-run.
        st.set_setting(conn, "wa_prepared_by_name", "Someone Else")
        assert st.seed_wa_prepared_by_default(conn) is False
        assert st.get_setting(conn, "wa_prepared_by_name") == "Someone Else"
    finally:
        conn.execute("DELETE FROM app_setting WHERE key IN (?, ?)",
                     (st._WA_PREPARED_BY_SEEDED, "wa_prepared_by_name"))
        conn.commit()


def test_annex_xlsx_reproduces_the_form(wa_org, app_mod):
    from openpyxl import load_workbook
    conn, ps = wa_org["conn"], _open()
    _approved(wa_org, wa_org["staff"][0], "CWW", 4)

    r = _client_for_id(app_mod, wa_org["chief"]).get(
        f"/work-arrangement/annex-c.xlsx?period={ps}")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "ANNEX-C-Work-Arrangement" in r.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [st.work_arrangement_period_label(ps)]
    flat = "\n".join(
        " | ".join("" if c is None else str(c) for c in row)
        for row in wb.active.iter_rows(values_only=True))
    for probe in ["ANNEX C", "INTERNAL WORK ARRANGEMENT SCHEDULE",
                  st.BUREAU_NAME, st.work_arrangement_period_label(ps),
                  st.WORK_ARRANGEMENT_LEGEND,          # verbatim, not reflowed
                  "Prepared by:", "Approved by:",
                  "WA-Test Staff0 | HPO II | WASec | CWW | Friday Day off"]:
        assert probe in flat, probe
    for header in wanx.HEADERS:
        assert header in flat


def test_annex_xlsx_month_gives_both_halves(wa_org, app_mod):
    from openpyxl import load_workbook
    ps = _open()
    r = _client_for_id(app_mod, wa_org["chief"]).get(
        f"/work-arrangement/annex-c.xlsx?month={ps[:7]}")
    assert r.status_code == 200
    assert load_workbook(io.BytesIO(r.content)).sheetnames == [
        st.work_arrangement_period_label(f"{ps[:7]}-01"),
        st.work_arrangement_period_label(f"{ps[:7]}-16")]


def test_annex_pdf_downloads(wa_org, app_mod):
    conn, ps = wa_org["conn"], _open()
    _approved(wa_org, wa_org["staff"][0], "WFH", 4)
    r = _client_for_id(app_mod, wa_org["chief"]).get(
        f"/work-arrangement/annex-c.pdf?period={ps}")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:4] == b"%PDF"


def test_a_malformed_month_redirects_rather_than_500s(wa_org, app_mod):
    r = _client_for_id(app_mod, wa_org["chief"]).get(
        "/work-arrangement/annex-c.xlsx?month=nonsense", follow_redirects=False)
    assert r.status_code == 303


# ------------------------------------------------------------------- RBAC ---

@pytest.mark.parametrize("path", [
    "/work-arrangement/approvals", "/work-arrangement/schedule",
    "/work-arrangement/annex-c.xlsx", "/work-arrangement/annex-c.pdf",
])
def test_chief_only_pages_are_gated(wa_org, app_mod, path):
    assert _client_for_id(app_mod, wa_org["staff"][0]).get(path).status_code == 403
    assert _client_for_id(app_mod, wa_org["chief"]).get(path).status_code == 200


def test_everyone_signed_in_gets_the_hub_and_the_form(wa_org, app_mod):
    c = _client_for_id(app_mod, wa_org["staff"][0])
    assert c.get("/work-arrangement").status_code == 200
    assert c.get("/work-arrangement/file").status_code == 200


def test_a_non_approver_cannot_decide_over_http(wa_org, app_mod):
    conn, ps = wa_org["conn"], _open()
    wa = st.create_work_arrangement(conn, wa_org["staff"][0], ps, "REG")
    r = _client_for_id(app_mod, wa_org["staff"][1]).post(
        f"/work-arrangement/{wa}/decide", data={"action": "approve"})
    assert r.status_code == 403
    assert st.get_work_arrangement(conn, wa)["status"] == "pending"


def test_only_an_admin_reaches_the_reopen_console(wa_org, app_mod):
    staff = _client_for_id(app_mod, wa_org["staff"][0])
    assert staff.get("/admin/work-arrangement").status_code == 403
    assert staff.post("/admin/work-arrangement/reopen",
                      data={"period_start": _open(), "action": "reopen"}
                      ).status_code == 403
    admin = _client_for_id(app_mod, wa_org["admin"]["id"])
    assert admin.get("/admin/work-arrangement").status_code == 200


def test_reopening_over_http_lets_the_employee_file_it_themselves(
        wa_org, app_mod, monkeypatch):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    past = st.work_arrangement_deadline(ps) + timedelta(days=1)
    monkeypatch.setattr(tz, "today_local", lambda: past)

    staff = _client_for_id(app_mod, emp)
    admin = _client_for_id(app_mod, wa_org["admin"]["id"])

    r = admin.post("/admin/work-arrangement/reopen",
                   data={"period_start": ps, "action": "reopen",
                         "note": "late request"}, follow_redirects=False)
    assert r.status_code == 303
    assert st.get_work_arrangement_reopen(conn, ps) is not None

    hub = staff.get("/work-arrangement")
    assert "Reopened by an admin" in hub.text

    # The filing form's own "Which period" picker has to offer the reopened
    # period too — the hub's "File it now" link is worthless if the form it
    # lands on doesn't actually list the period as choosable.
    file_form = staff.get("/work-arrangement/file")
    assert f'value="{ps}"' in file_form.text
    assert "reopened by an admin" in file_form.text

    r = staff.post("/work-arrangement/file",
                   data={"period_start": ps, "arrangement": "REG"},
                   follow_redirects=False)
    assert r.status_code == 303
    assert st.my_pending_work_arrangement(conn, emp, ps) is not None

    admin.post("/admin/work-arrangement/reopen",
              data={"period_start": ps, "action": "close"},
              follow_redirects=False)
    assert st.get_work_arrangement_reopen(conn, ps) is None


def test_staff_without_work_arrangement_excludes_filers(wa_org):
    conn, ps = wa_org["conn"], _open()
    st.create_work_arrangement(conn, wa_org["staff"][0], ps, "REG")
    non_filers = {r["id"] for r in st.staff_without_work_arrangement(conn, ps)}
    assert wa_org["staff"][0] not in non_filers
    assert wa_org["staff"][1] in non_filers
    assert wa_org["loner"] in non_filers


def test_staff_without_work_arrangement_skips_the_noawa_tagged(wa_org):
    """create_work_arrangement REFUSES a filing from someone tagged "No AWA",
    so nudging them would point at a door the app holds shut."""
    conn, ps = wa_org["conn"], _open()
    tagged = wa_org["staff"][2]
    st.set_work_arrangement_noawa(conn, tagged, "on secondment", _u(tagged))
    non_filers = {r["id"] for r in st.staff_without_work_arrangement(conn, ps)}
    assert tagged not in non_filers
    assert wa_org["staff"][1] in non_filers          # untagged, still owes one


def test_staff_without_work_arrangement_counts_a_denied_filing_as_unfiled(wa_org):
    """A denied (or withdrawn) row frees the period again — _WA_LIVE is what
    settles it — so that person still has to file and still needs the nudge."""
    conn, ps = wa_org["conn"], _open()
    emp = wa_org["staff"][3]
    wa = st.create_work_arrangement(conn, emp, ps, "REG")
    st.decide_work_arrangement(conn, wa, _u(wa_org["chief"], "section_manager"), "deny")
    assert emp in {r["id"] for r in st.staff_without_work_arrangement(conn, ps)}


def test_staff_without_work_arrangement_drops_an_archived_person(wa_org):
    """Resigned/transferred out: off work_arrangement_roster, so off this list
    too, even when the account itself was never deactivated."""
    conn, ps = wa_org["conn"], _open()
    gone = wa_org["staff"][4]
    rp_id = st.get_admin_user_by_id(conn, gone)["rp_id"]
    conn.execute("UPDATE resource_person SET archived_at=? WHERE id=?",
                 (tz.now_iso(), rp_id))
    conn.commit()
    assert gone not in {r["id"] for r in st.staff_without_work_arrangement(conn, ps)}


def test_reopen_emails_only_staff_who_did_not_file(wa_org, app_mod, monkeypatch):
    conn, ps = wa_org["conn"], _open()
    # staff[0] filed before the deadline; everyone else missed it.
    st.create_work_arrangement(conn, wa_org["staff"][0], ps, "REG")
    past = st.work_arrangement_deadline(ps) + timedelta(days=1)
    monkeypatch.setattr(tz, "today_local", lambda: past)

    calls = []
    monkeypatch.setattr(mailer, "send", lambda to, subject, body, html=None:
                        calls.append({"to": to, "subject": subject, "body": body})
                        or {"sent": True})
    admin = _client_for_id(app_mod, wa_org["admin"]["id"])
    r = admin.post("/admin/work-arrangement/reopen",
                   data={"period_start": ps, "action": "reopen", "note": "late request"},
                   follow_redirects=False)
    assert r.status_code == 303
    recipients = {c["to"] for c in calls}
    assert "wa-staff0@test.local" not in recipients     # already filed -> not nudged
    assert "wa-staff1@test.local" in recipients
    assert "wa-loner@test.local" in recipients
    mail = next(c for c in calls if c["to"] == "wa-loner@test.local")
    assert "reopened" in mail["subject"].lower()
    assert "/work-arrangement/file" in mail["body"]
    assert "late request" in mail["body"]


# ------------------------------------------------------- pages and wiring ---

def test_file_page_carries_live_coverage_data(wa_org, app_mod):
    conn, ps = wa_org["conn"], _open()
    _approved(wa_org, wa_org["staff"][0], "CWW", 4)
    r = _client_for_id(app_mod, wa_org["staff"][1]).get("/work-arrangement/file")
    assert r.status_code == 200
    assert "data-coverage=" in r.text


def test_approvals_page_flags_thin_coverage_without_disabling_the_buttons(
        wa_org, app_mod):
    conn, ps = wa_org["conn"], _open()
    for i in range(5):
        _approved(wa_org, wa_org["staff"][i], "CWW", 4)
    st.create_work_arrangement(conn, wa_org["staff"][5], ps, "CWW", 4)

    r = _client_for_id(app_mod, wa_org["chief"]).get("/work-arrangement/approvals")
    assert r.status_code == 200
    assert "on-site on Friday" in r.text
    assert 'value="approve"' in r.text and 'value="deny"' in r.text
    assert "disabled" not in r.text


def test_schedule_page_lists_the_unfiled(wa_org, app_mod):
    r = _client_for_id(app_mod, wa_org["chief"]).get("/work-arrangement/schedule")
    assert r.status_code == 200
    assert "not yet filed" in r.text


def test_landing_page_offers_the_third_door(wa_org, app_mod):
    r = _client_for_id(app_mod, wa_org["staff"][0]).get("/")
    assert r.status_code == 200
    assert 'href="/work-arrangement"' in r.text


def test_hub_leads_with_the_current_in_force_schedule(wa_org, app_mod, monkeypatch):
    """The period actually happening now must appear on the hub even once
    it's closed for filing — and only there, not duplicated into history."""
    emp, ps = wa_org["staff"][0], "2026-09-01"
    _at(monkeypatch, date(2026, 8, 26))            # Sept 1-15 still open then
    _approved(wa_org, emp, "CWW", 4, ps=ps)
    _at(monkeypatch, date(2026, 9, 3))             # now inside Sept 1-15

    r = _client_for_id(app_mod, emp).get("/work-arrangement")
    assert r.status_code == 200
    assert "Current schedule" in r.text
    label = st.work_arrangement_period_label(ps)
    assert label in r.text
    assert "in force" in r.text
    assert r.text.count(label) == 1


def test_hub_shows_the_next_period_on_a_worked_out_weekend(wa_org, app_mod, monkeypatch):
    """Saturday August 15: the 1st-15th has no workdays left, so the hub must
    already lead with the 16th-31st — the schedule about to be worked."""
    emp, ps = wa_org["staff"][0], "2026-08-16"
    _at(monkeypatch, date(2026, 8, 14))            # Aug 16-31 still open then
    _approved(wa_org, emp, "WFH", 2, ps=ps)
    _at(monkeypatch, date(2026, 8, 15))            # Saturday

    r = _client_for_id(app_mod, emp).get("/work-arrangement")
    assert r.status_code == 200
    assert "August 16-31, 2026" in r.text
    assert "August 1-15, 2026" not in r.text


def test_landing_team_card_leads_with_the_fortnight_being_worked(
        wa_org, app_mod, monkeypatch):
    """The team panel used to be built entirely from the period open for
    filing, so on August 15 a chief saw September's empty roster and nothing
    at all about the fortnight their team was about to work."""
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], "2026-08-16"
    chief = wa_org["chief"]
    st.mark_onboarded(conn, chief)
    _at(monkeypatch, date(2026, 8, 14))            # Aug 16-31 still open then
    _approved(wa_org, emp, "WFH", 2, ps=ps)
    _at(monkeypatch, date(2026, 8, 15))            # Saturday

    html = _client_for_id(app_mod, chief).get("/").text
    assert "Now: <strong>August 16-31, 2026</strong>" in html
    assert "Who is on-site now, by weekday" in html
    # ...and the filing period is still there, framed as the deadline it is.
    assert "Filing for <strong>September 1-15, 2026</strong>" in html
    # Nothing filed for September yet, so its three preview charts stay out of
    # the way rather than sitting empty in front of the fortnight being worked.
    assert "Who would be on-site once filed" not in html

    # File for September and the previews earn their place back.
    _approved(wa_org, emp, "CWW", 0, ps="2026-09-01")
    html = _client_for_id(app_mod, chief).get("/").text
    assert "Who would be on-site once filed" in html
    assert "What people filed" in html


def test_schedule_page_opens_on_the_worked_period_not_the_filing_one(
        wa_org, app_mod, monkeypatch):
    _at(monkeypatch, date(2026, 8, 15))            # Saturday
    r = _client_for_id(app_mod, wa_org["chief"]).get("/work-arrangement/schedule")
    assert r.status_code == 200
    assert "August 16-31, 2026" in r.text


def test_hub_current_schedule_says_so_when_nothing_was_filed(wa_org, app_mod, monkeypatch):
    _at(monkeypatch, date(2026, 9, 3))             # Sept 1-15, nothing filed
    r = _client_for_id(app_mod, wa_org["staff"][0]).get("/work-arrangement")
    assert r.status_code == 200
    assert "Nothing was filed for this period" in r.text


def test_home_door_shows_the_in_force_arrangement_not_just_the_filing_window(
        wa_org, app_mod, monkeypatch):
    """Regression: the door used to report only the next open-for-filing
    period, so a schedule already in force for TODAY never showed there."""
    emp, ps = wa_org["staff"][0], "2026-09-01"
    st.mark_onboarded(wa_org["conn"], emp)     # skip /welcome, land on the doors
    _at(monkeypatch, date(2026, 8, 26))
    _approved(wa_org, emp, "CWW", 4, ps=ps)
    _at(monkeypatch, date(2026, 9, 3))

    r = _client_for_id(app_mod, emp).get("/")
    assert r.status_code == 200
    assert "Now: " + st.work_arrangement_period_label(ps) in r.text
    assert "CWW" in r.text


def test_nav_and_search_expose_the_module(app_mod):
    class U:
        role = "viewer"
    urls = {it["url"] for g in app_mod.nav_groups_for(U()) for it in g["items"]}
    assert "/work-arrangement" in urls
    assert "/work-arrangement" in {it["url"] for it in app_mod.search_index_for(U())}


# ------------------------------------------------- one day, by name ---------
# Aug 16-31 2026 is a fortnight with TWO Fridays (the 21st and the 28th), which
# is what makes it the right period to prove the recurrence on.

_DAY_PS = "2026-08-16"
_FILE_BY = date(2026, 8, 14)          # on or before the Aug 16-31 deadline


def _where(d, uid):
    """Which group a person landed in — exactly one, or None if absent."""
    found = [k for k, people in d["groups"].items()
             if any(p["user_id"] == uid for p in people)]
    assert len(found) <= 1, f"{uid} appears in {found}"
    return found[0] if found else None


def test_a_nominated_day_off_recurs_on_every_such_weekday(wa_org, monkeypatch):
    """The whole basis of the view: a fortnightly filing resolves to specific
    dates because "Friday Day off" means EVERY Friday, not one of them."""
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    _at(monkeypatch, _FILE_BY)
    _approved(wa_org, emp, "CWW", 4, ps=_DAY_PS)          # Fridays off

    for friday in (date(2026, 8, 21), date(2026, 8, 28)):
        assert _where(st.work_arrangement_day(conn, friday), emp) == "dayoff"
    for other in (date(2026, 8, 20), date(2026, 8, 27), date(2026, 8, 17)):
        assert _where(st.work_arrangement_day(conn, other), emp) == "onsite"


def test_a_day_off_is_told_apart_from_a_day_worked_from_home(wa_org, monkeypatch):
    """Conflating the two would report someone on leave as reachable."""
    conn = wa_org["conn"]
    off, home = wa_org["staff"][0], wa_org["staff"][1]
    _at(monkeypatch, _FILE_BY)
    _approved(wa_org, off, "CWW", 2, ps=_DAY_PS)          # Wednesday off
    _approved(wa_org, home, "WFH", 2, ps=_DAY_PS)         # Wednesday from home

    wed = st.work_arrangement_day(conn, date(2026, 8, 19))
    assert _where(wed, off) == "dayoff"
    assert _where(wed, home) == "wfh"
    assert wed["counts"]["dayoff"] == 1 and wed["counts"]["wfh"] == 1
    # Everyone lands in exactly one group, so the groups partition the roster.
    assert sum(wed["counts"].values()) == wed["headcount"]

    tue = st.work_arrangement_day(conn, date(2026, 8, 18))
    assert _where(tue, off) == "onsite" and _where(tue, home) == "onsite"


def test_nothing_filed_is_never_counted_as_in_the_office(wa_org, monkeypatch):
    """A blank row is not a person at their desk. Folding the unfiled into
    "in the office" would overstate the count with people nobody has heard
    from — the exact gap the chase list exists to close."""
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    _at(monkeypatch, _FILE_BY)
    st.set_work_arrangement_noawa(conn, wa_org["staff"][1], "on leave",
                                  _u(wa_org["staff"][1]))

    d = st.work_arrangement_day(conn, date(2026, 8, 18))
    assert _where(d, emp) == "unfiled"                    # never filed
    assert _where(d, wa_org["staff"][1]) == "noawa"
    assert d["counts"]["onsite"] == 0
    tagged = d["groups"]["noawa"][0]
    assert tagged["remarks"] == "on leave"                # the reason travels


def test_nobody_is_scheduled_on_a_weekend(wa_org, monkeypatch):
    conn = wa_org["conn"]
    _at(monkeypatch, _FILE_BY)
    _approved(wa_org, wa_org["staff"][0], "REG", ps=_DAY_PS)

    sat = st.work_arrangement_day(conn, date(2026, 8, 22))
    assert sat["is_weekend"] and sat["headcount"] == 0
    assert all(not p for p in sat["groups"].values())
    # ...and the arrows step over it rather than landing on it.
    assert st.work_arrangement_day_step(date(2026, 8, 21), 1) == date(2026, 8, 24)
    assert st.work_arrangement_day_step(date(2026, 8, 24), -1) == date(2026, 8, 21)


def test_an_unlinked_account_is_shown_nobody_not_everybody(wa_org, monkeypatch):
    """The guard that matters. This page has no chief route guard, and
    work_arrangement_roster scopes through _org_scope_for_user, which folds "no
    team" onto the same None that means bureau-wide. If work_arrangement_day
    ever stops resolving scope itself, an account with no directory record sees
    every section in the bureau."""
    conn = wa_org["conn"]
    _at(monkeypatch, _FILE_BY)
    _approved(wa_org, wa_org["staff"][0], "CWW", 4, ps=_DAY_PS)

    d = st.work_arrangement_day(conn, date(2026, 8, 18), _u(wa_org["loner"]))
    assert d["kind"] == "none"
    assert d["headcount"] == 0
    assert "WA-Test" not in repr(d["groups"])             # not one colleague


def test_the_day_view_is_open_to_everyone_unlike_the_team_schedule(
        wa_org, app_mod, monkeypatch):
    """Asserted together so the difference stays deliberate: knowing whether a
    colleague is in is not the same power as reading the filing progress and
    printing ANNEX C."""
    _at(monkeypatch, date(2026, 8, 18))
    c = _client_for_id(app_mod, wa_org["staff"][0])
    assert c.get("/work-arrangement/day").status_code == 200
    assert c.get("/work-arrangement/schedule").status_code == 403


def test_the_day_view_names_people_and_bounces_a_bad_date(
        wa_org, app_mod, monkeypatch):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    _at(monkeypatch, _FILE_BY)
    _approved(wa_org, emp, "WFH", 4, ps=_DAY_PS)
    _at(monkeypatch, date(2026, 8, 18))
    c = _client_for_id(app_mod, wa_org["chief"])

    html = c.get("/work-arrangement/day?day=2026-08-21").text
    assert "Friday, August 21, 2026" in html
    assert "Working from home" in html
    assert "WA-Test Staff0" in html                       # by name, not a count

    r = c.get("/work-arrangement/day?day=not-a-date", follow_redirects=False)
    assert r.status_code == 303                           # bounces, never 500s


def test_landing_card_leads_with_who_is_in_today(wa_org, app_mod, monkeypatch):
    conn, chief = wa_org["conn"], wa_org["chief"]
    st.mark_onboarded(conn, chief)
    _at(monkeypatch, _FILE_BY)
    _approved(wa_org, wa_org["staff"][0], "CWW", 1, ps=_DAY_PS)   # Tuesday off
    _at(monkeypatch, date(2026, 8, 18))                            # a Tuesday

    html = _client_for_id(app_mod, chief).get("/").text
    assert "in the office today" in html
    assert "Tuesday, August 18, 2026" in html
    assert 'href="/work-arrangement/day"' in html
    # the headline count itself is the link — people clicked it for the names
    assert 'class="oc-big oc-big-link" href="/work-arrangement/day"' in html


# ------------------------------------------------------------- "No AWA" ------

def _at(monkeypatch, d):
    monkeypatch.setattr(tz, "today_local", lambda: d)


def test_noawa_needs_a_reason(wa_org):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    for blank in ("", "   ", None):
        with pytest.raises(ValueError, match="Say why"):
            st.set_work_arrangement_noawa(conn, emp, blank, _u(emp))


def test_noawa_carries_into_later_periods_with_nothing_run(wa_org):
    """The whole point of storing it as an interval: no rollover step exists, so
    if this ever fails the tag has quietly become per-period."""
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    tag = st.set_work_arrangement_noawa(conn, emp, "on leave", _u(emp))

    # It starts in the period we are IN, not the next one open for filing —
    # that is what lets it cover something that comes up mid-period.
    assert tag["effective_from"] == st.work_arrangement_period()
    p = tag["effective_from"]
    for _ in range(5):
        assert st.get_work_arrangement_noawa(conn, emp, p)["remark"] == "on leave"
        p = st.work_arrangement_next_period(p)
    # ...but it does not reach backwards into a period already submitted.
    assert st.get_work_arrangement_noawa(
        conn, emp, st.work_arrangement_prev_period(tag["effective_from"])) is None


def test_rewording_later_keeps_what_the_old_period_said(wa_org, monkeypatch):
    """An annex already submitted has to keep the wording it went out with."""
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    st.set_work_arrangement_noawa(conn, emp, "awaiting OSEC approval", _u(emp))

    later = st.work_arrangement_next_period(st.work_arrangement_next_period(ps))
    _at(monkeypatch, date.fromisoformat(later))
    st.set_work_arrangement_noawa(conn, emp, "OSEC approved, filing next period",
                                  _u(emp))

    assert st.get_work_arrangement_noawa(
        conn, emp, ps)["remark"] == "awaiting OSEC approval"
    assert st.get_work_arrangement_noawa(
        conn, emp, later)["remark"] == "OSEC approved, filing next period"
    # two intervals, exactly one of them still open
    assert conn.execute(
        "SELECT COUNT(*) FROM work_arrangement_noawa WHERE user_id=?",
        (emp,)).fetchone()[0] == 2
    assert conn.execute(
        "SELECT COUNT(*) FROM work_arrangement_noawa "
        "WHERE user_id=? AND effective_to IS NULL", (emp,)).fetchone()[0] == 1


def test_rewording_within_the_same_period_corrects_in_place(wa_org):
    """No history to preserve yet, so a typo fix shouldn't leave a stub behind."""
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    st.set_work_arrangement_noawa(conn, emp, "on leave", _u(emp))
    st.set_work_arrangement_noawa(conn, emp, "on maternity leave", _u(emp))
    assert conn.execute(
        "SELECT COUNT(*) FROM work_arrangement_noawa WHERE user_id=?",
        (emp,)).fetchone()[0] == 1
    assert st.get_work_arrangement_noawa(conn, emp)["remark"] == "on maternity leave"


def test_removal_lifts_it_from_here_on_and_leaves_history(wa_org, monkeypatch):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    st.set_work_arrangement_noawa(conn, emp, "on leave", _u(emp))

    later = st.work_arrangement_next_period(st.work_arrangement_next_period(ps))
    _at(monkeypatch, date.fromisoformat(later))
    st.end_work_arrangement_noawa(conn, emp, _u(emp))

    assert st.get_work_arrangement_noawa(conn, emp, later) is None
    assert st.get_work_arrangement_noawa(conn, emp, ps)["remark"] == "on leave"
    with pytest.raises(ValueError, match="no .No AWA. to remove"):
        st.end_work_arrangement_noawa(conn, emp, _u(emp))


def test_only_one_tag_can_be_live_at_a_time(wa_org):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    st.set_work_arrangement_noawa(conn, emp, "on leave", _u(emp))
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute(
            "INSERT INTO work_arrangement_noawa(user_id, remark, effective_from) "
            "VALUES (?,'second','2099-01-01')", (emp,))
    conn.rollback()


def test_the_person_or_an_admin_may_act_and_nobody_else(wa_org):
    conn, emp, peer = wa_org["conn"], wa_org["staff"][0], wa_org["staff"][1]
    # a peer cannot speak for someone — nor can their section chief
    for actor in (_u(peer), _u(wa_org["chief"], "section_manager")):
        with pytest.raises(PermissionError):
            st.set_work_arrangement_noawa(conn, emp, "on leave", actor)

    # an admin can record it for someone who isn't logging in
    st.set_work_arrangement_noawa(conn, emp, "on leave", wa_org["admin"])
    assert st.get_work_arrangement_noawa(conn, emp)["remark"] == "on leave"
    with pytest.raises(PermissionError):
        st.end_work_arrangement_noawa(conn, emp, _u(peer))
    st.end_work_arrangement_noawa(conn, emp, _u(emp))       # ...and they can lift it
    assert st.get_work_arrangement_noawa(conn, emp) is None


def test_noawa_needs_no_approval_and_ignores_the_deadline(wa_org, monkeypatch):
    """The case this exists for: something comes up mid-period, long after
    filing closed. If this ever starts raising, Sales' OSEC situation has no
    path through the system again."""
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    inside = st.work_arrangement_period_bounds(ps)[0] + timedelta(days=3)
    _at(monkeypatch, inside)
    assert st.work_arrangement_deadline(ps) < inside          # filing has closed
    with pytest.raises(ValueError, match="deadline"):
        st.create_work_arrangement(conn, emp, ps, "WFH", 4)

    tag = st.set_work_arrangement_noawa(conn, emp, "hospitalised", _u(emp))
    assert tag["remark"] == "hospitalised"
    assert st.get_work_arrangement_noawa(conn, emp, ps)         # in force immediately
    # ...and it never reaches an approver
    assert st.list_work_arrangement_approvals(conn, wa_org["admin"]) == []


def test_filing_while_tagged_is_refused_and_says_how(wa_org):
    conn, emp, ps = wa_org["conn"], wa_org["staff"][0], _open()
    st.set_work_arrangement_noawa(conn, emp, "on leave", _u(emp))
    with pytest.raises(ValueError, match="Remove that first"):
        st.create_work_arrangement(conn, emp, ps, "WFH", 4)
    # lifting it hands the period back
    st.end_work_arrangement_noawa(conn, emp, _u(emp))
    assert st.create_work_arrangement(conn, emp, ps, "WFH", 4)


def test_tagged_people_leave_the_coverage_denominator(wa_org):
    """A section of 7 with one person on leave is a section of 6 — they were
    never going to be at a desk on Friday, and counting them either way
    misstates what this warning is for."""
    conn, ps, path = wa_org["conn"], _open(), wa_org["path"]
    before = st.work_arrangement_coverage(conn, path, ps)
    assert before["headcount"] == 7 and before["excluded"] == 0
    assert before["weekdays"][4]["onsite"] == 7

    st.set_work_arrangement_noawa(conn, wa_org["staff"][0], "on leave", _u(wa_org["staff"][0]))
    after = st.work_arrangement_coverage(conn, path, ps)
    assert after["headcount"] == 6 and after["excluded"] == 1
    assert after["weekdays"][4]["onsite"] == 6      # not on-site AND not off-site
    assert after["weekdays"][4]["offsite"] == 0


def test_the_flag_thresholds_move_with_the_smaller_section(wa_org):
    """Six people once one is tagged, so it takes one fewer Friday absence to
    reach 'nobody on-site' than the seven-person case does."""
    conn, ps = wa_org["conn"], _open()
    st.set_work_arrangement_noawa(conn, wa_org["staff"][0], "on leave", _u(wa_org["staff"][0]))
    for i in range(1, 5):
        _approved(wa_org, wa_org["staff"][i], "CWW", 4)

    def flag_for(uid):
        wa = st.create_work_arrangement(conn, uid, ps, "CWW", 4)
        return st.work_arrangement_coverage_flag(
            conn, st.get_work_arrangement(conn, wa))

    assert flag_for(wa_org["staff"][5]) == "Only 1 on-site on Friday after this"
    assert flag_for(wa_org["chief"]) == "Approving leaves nobody on-site on Friday"


def test_a_tagged_person_is_not_a_chase_list_gap(wa_org):
    conn, ps = wa_org["conn"], _open()
    emp = wa_org["staff"][0]
    st.set_work_arrangement_noawa(conn, emp, "on leave", _u(emp))

    d = st.work_arrangement_period_summary(conn, ps,
                                           _u(wa_org["chief"], "section_manager"))
    assert "WA-Test Staff0" in d["noawa"]
    assert "WA-Test Staff0" not in d["unfiled"]
    assert d["headcount"] == 6                       # out of "filed n of m" too
    row = next(r for r in d["roster"] if r["name"] == "WA-Test Staff0")
    assert row["filed"] and row["noawa"]
    assert row["arrangement"] == st.WORK_ARRANGEMENT_NOAWA


def test_a_tag_outranks_a_stale_filing(wa_org):
    """Someone files, then goes on leave. The stated reason is the truth."""
    conn, ps, emp = wa_org["conn"], _open(), wa_org["staff"][0]
    _approved(wa_org, emp, "WFH", 4)
    st.set_work_arrangement_noawa(conn, emp, "on leave", _u(emp))

    row = next(r for r in st.work_arrangement_export_rows(conn, ps)
               if r["name"] == "WA-Test Staff0")
    assert row["arrangement"] == st.WORK_ARRANGEMENT_NOAWA
    assert row["remarks"] == "on leave"


def test_the_annex_distinguishes_no_awa_from_never_filed(wa_org, app_mod):
    """The Bagos/Sales shape: a stated reason prints, a genuine gap stays blank.
    On the source spreadsheet these two looked identical."""
    from openpyxl import load_workbook
    conn, ps = wa_org["conn"], _open()
    st.set_work_arrangement_noawa(conn, wa_org["staff"][0], "Leave", _u(wa_org["staff"][0]))
    st.set_work_arrangement_noawa(conn, wa_org["staff"][1],
                                  "with request WFH due to medical condition, "
                                  "for approval from OSEC", _u(wa_org["staff"][1]))
    # staff[2] files nothing at all

    rows = {r["name"]: r for r in st.work_arrangement_export_rows(conn, ps)}
    assert rows["WA-Test Staff0"]["arrangement"] == "No AWA"
    assert rows["WA-Test Staff0"]["remarks"] == "Leave"
    assert "OSEC" in rows["WA-Test Staff1"]["remarks"]
    assert rows["WA-Test Staff2"]["arrangement"] == ""       # genuinely unfilled
    assert rows["WA-Test Staff2"]["remarks"] == ""

    c = _client_for_id(app_mod, wa_org["chief"])
    wb = load_workbook(io.BytesIO(
        c.get(f"/work-arrangement/annex-c.xlsx?period={ps}").content))
    flat = "\n".join(" | ".join("" if v is None else str(v) for v in row)
                     for row in wb.active.iter_rows(values_only=True))
    assert "WA-Test Staff0 | HPO II | WASec | No AWA | Leave" in flat
    assert "OSEC" in flat

    pdf = c.get(f"/work-arrangement/annex-c.pdf?period={ps}")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"


def test_noawa_over_http(wa_org, app_mod):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    c = _client_for_id(app_mod, emp)

    assert c.post("/work-arrangement/no-awa", data={"remark": "on leave"},
                  follow_redirects=False).status_code == 303
    hub = c.get("/work-arrangement").text
    assert "on leave" in hub and "no alternative work arrangement" in hub
    assert "File it now" not in hub          # the period cards step aside

    # a peer cannot act for them
    assert _client_for_id(app_mod, wa_org["staff"][1]).post(
        "/work-arrangement/no-awa",
        data={"user_id": emp, "remark": "x"}).status_code == 403

    assert c.post("/work-arrangement/no-awa/remove",
                  follow_redirects=False).status_code == 303
    assert st.get_work_arrangement_noawa(conn, emp) is None


def test_admin_can_tag_someone_from_the_schedule_page(wa_org, app_mod):
    conn, emp = wa_org["conn"], wa_org["staff"][0]
    admin_c = _client_for_id(app_mod, wa_org["admin"]["id"])
    page = admin_c.get("/work-arrangement/schedule")
    assert page.status_code == 200 and 'name="user_id"' in page.text

    r = admin_c.post("/work-arrangement/no-awa",
                     data={"user_id": emp, "remark": "on leave",
                           "back": "/work-arrangement/schedule"},
                     follow_redirects=False)
    assert r.status_code == 303 and "/work-arrangement/schedule" in r.headers["location"]
    assert st.get_work_arrangement_noawa(conn, emp)["remark"] == "on leave"
    # attributable: an admin acting for someone else names them in the trail
    summary = conn.execute(
        "SELECT summary FROM audit_log WHERE action='work_arrangement_noawa_set' "
        "ORDER BY id DESC LIMIT 1").fetchone()[0]
    assert "WA-Test Staff0" in summary and "on leave" in summary


# ---- finding the right period's annex after an approval --------------------

def test_the_schedule_points_at_the_upcoming_period_once_it_has_filings(wa_org, app_mod):
    """Approvals are made for the NEXT period, but the schedule opens on the one
    being worked — so the page must surface the upcoming period's own downloads,
    or the approver prints a sheet where the person they just approved is blank."""
    o = wa_org
    now_ps = st.work_arrangement_schedule_period()
    nxt = st.work_arrangement_next_period(now_ps)
    if nxt not in st.work_arrangement_open_periods():
        pytest.skip("next period not open for filing on this date")
    c = _client_for_id(app_mod, o["chief"])
    assert f"annex-c.pdf?period={nxt}" not in c.get("/work-arrangement/schedule").text
    wid = st.create_work_arrangement(o["conn"], o["staff"][0], nxt, "WFH", 4)
    r = c.post(f"/work-arrangement/{wid}/decide", data={"action": "approve"},
               follow_redirects=False)
    assert f"period={nxt}" in r.headers["location"]
    assert f"annex-c.pdf?period={nxt}" in c.get(r.headers["location"]).text
    page = c.get("/work-arrangement/schedule").text
    assert f"annex-c.pdf?period={nxt}" in page
    assert f"annex-c.pdf?period={now_ps}" in page
    # ...and that period's annex really does carry the approval.
    row = next(x for x in st.work_arrangement_export_rows(o["conn"], nxt)
               if x["name"] == "WA-Test Staff0")
    assert row["arrangement"] == "WFH"
