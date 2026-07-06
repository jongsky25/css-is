"""Spot Feedback #56 — raw / per-submission data behind the SERVQUAL numbers.

The Type-2/3 raw-data Excel gains a "Responses" sheet: one row per submission,
carrying that respondent's OWN per-dimension means, their overall, expectation
(item A) and overall service quality (item C). These lock:
  * the pure per-response maths (each answered item counted once per person);
  * `servqual_xlsx` renders a populated "Responses" sheet from the payload;
  * end-to-end, downloading the RP report's Excel yields one distinct row per
    submission, narrowed to the chosen activity.
"""
import io

import pandas as pd
from openpyxl import load_workbook

import app
import css_is_engine as eng
import report_docs
import storage as st


def _atoms(rows):
    """Build a minimal atoms frame (only the columns the row builder reads)."""
    df = pd.DataFrame(rows, columns=["response_id", "item_id", "value"])
    df["is_na"] = False
    return df


# ---- pure per-submission maths ---------------------------------------------

def test_response_rows_disaggregate_per_submitter():
    rows = []
    # Respondent 101: Reliability 6, Responsiveness 5, Assurance (6,4)->5,
    # Empathy 6, Tangible 6  => overall 5.6 ; expectation 6 ; OSQ 4
    for it, val in [("A", 6), ("B1", 6), ("B2", 6), ("B3", 6), ("B4", 6),
                    ("B5", 5), ("B6", 6), ("B7", 4), ("B8", 6), ("B9", 6), ("C", 4)]:
        rows.append((101, it, val))
    # Respondent 102: every rating 4 => all dims 4, overall 4 ; expectation 5 ; OSQ 3
    for it, val in [("A", 5), ("B1", 4), ("B2", 4), ("B3", 4), ("B4", 4),
                    ("B5", 4), ("B6", 4), ("B7", 4), ("B8", 4), ("B9", 4), ("C", 3)]:
        rows.append((102, it, val))

    out = app._servqual_response_rows(_atoms(rows), eng.DIM_MAP_FORM3)
    assert set(out) == {101, 102}

    a = out[101]
    assert a["dims"] == {"Reliability": 6.0, "Responsiveness": 5.0,
                         "Assurance": 5.0, "Empathy": 6.0, "Tangible": 6.0}
    assert a["overall"] == 5.6 and a["expectation"] == 6.0 and a["osq"] == 4.0

    b = out[102]
    assert b["dims"]["Reliability"] == 4.0 and b["overall"] == 4.0
    assert b["expectation"] == 5.0 and b["osq"] == 3.0


def test_response_rows_empty_frame_is_safe():
    assert app._servqual_response_rows(_atoms([]), eng.DIM_MAP_FORM3) == {}


# ---- the XLSX "Responses" sheet --------------------------------------------

def test_servqual_xlsx_has_responses_sheet():
    sq = {
        "has_data": True, "respondents": 2,
        "dimensions": [{"dim": d, "members": "", "value": 5.0, "n": 2,
                        "gap_shown": "—", "verdict": None}
                       for d in eng.DIM_MAP_FORM3],
        "summary": {"total": "25.00", "average": "5.00", "expectation": "5.50",
                    "servqual": "-0.50", "servqual_label": "met",
                    "osq": "3.50", "osq_band": "Excellent"},
        "responses": [
            {"response_id": 101, "when": "2026-06-10 09:00", "activity": "Alpha Talk",
             "dims": {"Reliability": 6.0, "Responsiveness": 5.0, "Assurance": 5.0,
                      "Empathy": 6.0, "Tangible": 6.0},
             "overall": 5.6, "expectation": 6.0, "osq": 4.0, "comment": "great session"},
            {"response_id": 102, "when": "2026-06-10 10:00", "activity": "Alpha Talk",
             "dims": {"Reliability": 4.0, "Responsiveness": 4.0, "Assurance": 4.0,
                      "Empathy": 4.0, "Tangible": 4.0},
             "overall": 4.0, "expectation": 5.0, "osq": 3.0, "comment": ""},
        ],
    }
    wb = load_workbook(io.BytesIO(report_docs.servqual_xlsx(sq, {"title": "RP"})))
    assert "Responses" in wb.sheetnames
    ws = wb["Responses"]
    headers = [c.value for c in ws[3]]
    assert headers[:3] == ["#", "Submitted", "Activity"]
    assert "Reliability" in headers and "Comment" in headers
    # two data rows (title row 1, headers row 3, data from row 4)
    body = [[c.value for c in r] for r in ws.iter_rows(min_row=4, max_row=5)]
    assert body[0][0] == 1 and body[1][0] == 2
    rel = headers.index("Reliability")
    assert body[0][rel] == 6.0 and body[1][rel] == 4.0          # disaggregated
    assert body[0][headers.index("Comment")] == "great session"


# ---- end-to-end: the real download, narrowed to one activity ---------------

_F3_A = [("A", 6), ("B1", 6), ("B2", 6), ("B3", 6), ("B4", 6),
         ("B5", 6), ("B6", 6), ("B7", 6), ("B8", 6), ("B9", 6), ("C", 4)]
_F3_B = [("A", 5), ("B1", 4), ("B2", 4), ("B3", 4), ("B4", 4),
         ("B5", 4), ("B6", 4), ("B7", 4), ("B8", 4), ("B9", 4), ("C", 3)]


def _one_internal_rp(db_path):
    conn = st.connect(db_path)
    row = conn.execute(
        "SELECT id FROM resource_person WHERE is_internal=1 ORDER BY id LIMIT 1").fetchone()
    conn.close()
    return row[0]


def _make_bureau(manager, db_path, name, date, venue):
    rp = _one_internal_rp(db_path)
    r = manager.post("/create", data={
        "event_type": "bureau", "activity_name": name,
        "event_date": date, "venue": venue, "qr_mode": "combined",
        "rp_id": [str(rp)], "topic": ["Session topic"],
    }, follow_redirects=True)
    assert r.status_code == 200, r.text[:400]
    conn = st.connect(db_path)
    ev = conn.execute("SELECT id FROM event WHERE title=? ORDER BY id DESC LIMIT 1",
                      (name,)).fetchone()[0]
    f3 = conn.execute(
        "SELECT fi.id FROM event_form ef JOIN form_instance fi ON fi.id = ef.form_instance_id "
        "WHERE ef.event_id=? AND fi.survey_type=3 ORDER BY fi.id LIMIT 1", (ev,)).fetchone()[0]
    conn.close()
    return ev, f3


def _option_value(body, name):
    import re
    m = re.search(rf'<option value="(\d+)"[^>]*>{re.escape(name)} —', body)
    return m.group(1) if m else None


def test_rp_xlsx_has_one_row_per_submission_for_activity(manager, app_mod):
    db_path = app_mod.DB
    _ev, f3 = _make_bureau(manager, db_path, "Raw Rows Activity", "2026-06-12", "Hall R")
    conn = st.connect(db_path)
    st.record_response(conn, f3, _F3_A)
    st.record_response(conn, f3, _F3_B)
    conn.close()
    app_mod.invalidate_atoms()

    body = manager.get("/reports", params={"type": 3}).text
    act = _option_value(body, "Raw Rows Activity")
    assert act, "the seeded activity should be an Activity filter option"

    r = manager.get("/reports/export",
                    params={"type": 3, "fmt": "xlsx", "activity": act})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert "Responses" in wb.sheetnames
    ws = wb["Responses"]
    headers = [c.value for c in ws[3]]
    rel = headers.index("Reliability")
    body_rows = [[c.value for c in row] for row in ws.iter_rows(min_row=4)
                 if row[0].value is not None]
    assert len(body_rows) == 2                                 # one row per submission
    rel_vals = sorted(v[rel] for v in body_rows)
    assert rel_vals == [4.0, 6.0]                              # the two differ (disaggregated)
    # Every row is labelled — even the comment-less one (guards the metadata source:
    # it must cover ALL submissions, not only those that left free text).
    sub, act = headers.index("Submitted"), headers.index("Activity")
    assert all(v[sub] and v[act] == "Raw Rows Activity" for v in body_rows)
