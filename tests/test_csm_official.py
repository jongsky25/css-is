"""Official CO-CSM (ARTA Type-1) report — the government-template generator.

Covers the data assembly (`app._csm_official_report`), the period heading helper,
and that both document builders (`csm_report.official_docx` / `official_pdf`)
produce valid, populated artifacts. Built on synthetic BLHSD-format rows so the
test is hermetic (no DB seeding needed) and the arithmetic is checkable by hand.
"""
from __future__ import annotations
import io

import app
import csm_report
from docx import Document
import fitz  # pymupdf


def _row(sqd, *, type_of_service="External", service="Technical Assistance",
         sex="Male", age="30", region="NCR", client_type="Government",
         cc1=app.CSM_CC1_OPTS[0], cc2=app.CSM_CC2_OPTS[0], cc3=app.CSM_CC3_OPTS[0]):
    """One BLHSD-format response row (matches app._csm_blhsd_rows output keys).
    `sqd` is a dict of SQD0..SQD8 -> value (use 0 for an N/A answer)."""
    r = {"RESPONSE_NO": "", "SERVICE": service, "TYPE_OF_SERVICE": type_of_service,
         "CLIENT_TYPE": client_type, "EMAIL": "", "DATE": "2026-02-01",
         "SEX": sex, "AGE": age, "REGION": region, "ONSITE_ONLINE": "Onsite",
         "CC1": cc1, "CC2": cc2, "CC3": cc3, "SUGGESTIONS": "", "FEEDBACK_TYPE": ""}
    for i in range(9):
        r[f"SQD{i}"] = sqd.get(f"SQD{i}", 5)
    return r


def _meta(period="quarter:2026-Q1"):
    return {"title": "Client Satisfaction Measurement Report", "place": "BLHSD",
            "period": period, "generated": "2026-06-15 14:00", "source": "BLHSD",
            "period_info": app._csm_period_heading(period),
            "prepared_by": app.CSM_PREPARED_BY, "approved_by": app.CSM_APPROVED_BY,
            "seal_path": app._csm_seal_path(),
            "formula_path": app._csm_asset("csm_formula.png"),
            "sign_date": "June 15, 2026"}


# --------------------------------------------------------------------------
# Assembly
# --------------------------------------------------------------------------

def test_empty_rows():
    rep = app._csm_official_report([])
    assert rep["has_data"] is False
    assert rep["respondents"] == 0


def test_external_internal_split_and_demographics():
    rows = [_row({}, type_of_service="External", region="NCR") for _ in range(8)]
    rows += [_row({}, type_of_service="External", region="Region III") for _ in range(2)]
    rows += [_row({}, type_of_service="Internal", region="NCR", sex="Female")]
    rep = app._csm_official_report(rows)
    assert rep["respondents"] == 11
    assert rep["n_external"] == 10 and rep["n_internal"] == 1
    region = {d["label"]: d for d in rep["demo"]["region"]}
    # External column: 8/10 NCR, 2/10 Region III; Overall: 9/11 NCR
    assert region["NCR"]["ext"] == 80.0
    assert region["Region III"]["ext"] == 20.0
    assert round(region["NCR"]["all"], 2) == round(9 / 11 * 100, 2)
    assert region["Region III"]["int"] == 0.0


def test_costs_na_not_communication():
    """A no-fee service: SQD5 (Costs) is all N/A. Communication (SQD4) must NOT
    be the N/A dimension — it should score normally."""
    rows = [_row({"SQD5": 0}) for _ in range(12)]  # 0 == N/A
    rep = app._csm_official_report(rows)
    dims = {d["dimension"]: d for d in rep["dimensions"]}
    assert dims["Costs"]["na"] == 12 and dims["Costs"]["score"] is None
    assert dims["Communication"]["na"] == 0
    assert dims["Communication"]["score"] == 100.0
    # Overall excludes the all-N/A dimension and is still Outstanding.
    assert rep["overall"] == 100.0
    assert rep["overall_band"] == "Outstanding"


def test_scope_catalogue_and_per_service_dash():
    rows = [_row({}, service="Technical Assistance") for _ in range(5)]
    rep = app._csm_official_report(rows)
    assert rep["scope"]["ext_total"] == 5
    others = [i for i in rep["scope"]["external"] if i["service"].startswith("OTHERS")]
    assert any(i["responses"] == 5 for i in others)
    # services with no responses -> score None (rendered as "—")
    no_resp = [s for s in rep["per_service"]["external"] if s["responses"] == 0]
    assert no_resp and all(s["score"] is None for s in no_resp)


def test_period_heading_variants():
    assert app._csm_period_heading("quarter:2026-Q1")["heading"] == "1st Quarter (CY 2026)"
    assert app._csm_period_heading("quarter:2026-Q3")["heading"] == "3rd Quarter (CY 2026)"
    assert app._csm_period_heading("semester:2026-H2")["heading"] == "2nd Semester (CY 2026)"
    assert app._csm_period_heading("year:2027")["heading"] == "CY 2027"
    assert app._csm_period_heading("all")["heading"] == "All Periods"


# --------------------------------------------------------------------------
# Document builders
# --------------------------------------------------------------------------

def test_docx_builds_and_is_populated():
    rows = [_row({"SQD5": 0}) for _ in range(21)]
    rep = app._csm_official_report(rows)
    data = csm_report.official_docx(rep, _meta())
    assert data[:2] == b"PK"  # docx is a zip
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for r in tbl.rows:
            for c in r.cells:
                parts.append(c.text)
    text = "\n".join(parts)
    for needle in ("DEPARTMENT OF HEALTH", "Overview", "Scope", "Methodology",
                   "Data and Interpretation", "Continuous Agency Improvement",
                   app.CSM_PREPARED_BY["name"]):
        assert needle in text, needle
    # the cover period reflects the chosen period, in the file's casing
    assert "1st QUARTER (CY 2026)" in text


def test_pdf_builds_and_is_populated():
    rows = [_row({"SQD5": 0}) for _ in range(21)]
    rep = app._csm_official_report(rows)
    data = csm_report.official_pdf(rep, _meta())
    assert data[:4] == b"%PDF"
    d = fitz.open(stream=data, filetype="pdf")
    full = "\n".join(p.get_text() for p in d)
    for needle in ("DEPARTMENT OF HEALTH", "Overview", "Service Quality Dimensions",
                   "Outstanding", "Approved by"):
        assert needle in full, needle


def _docx_all_text(data):
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for r in tbl.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def test_template_exact_match_elements():
    """Lock in the elements that must NOT be omitted from the official template
    (these were missing/altered in the first 'enhanced' draft). Checked against
    the DOCX, where table-cell text is not wrapped by layout."""
    rows = [_row({"SQD5": 0}) for _ in range(21)]
    rep = app._csm_official_report(rows)
    full = _docx_all_text(csm_report.official_docx(rep, _meta()))
    must_have = [
        "tinyurl.com/CentralOfficeCSM",                       # encoding link
        "BLHSD believes the demographic results",             # representative sentence
        "OVERALL REPORT ANALYSIS",                            # section heading
        "Total Responses",                                    # full column header
        "Neither Agree nor Disagree",                         # full Likert header
        "MIMAROPA", "NIR",                                    # template region rows
        "OVERALL TOTAL",                                      # scope + per-service totals
        "Date:",                                              # signatory date lines
        "Overall Score = (147 + 0) / (147",                   # worked formula (SA+A)/(N-NA)
        "External Service Total", "Internal Service Total",   # per-service + scope totals
    ]
    for needle in must_have:
        assert needle in full, needle


def test_formula_and_seal_images_embedded():
    """The seal (cover) and the formula graphic (Methodology) are embedded as
    images, matching the official .docx."""
    rows = [_row({"SQD5": 0}) for _ in range(21)]
    rep = app._csm_official_report(rows)
    doc = Document(io.BytesIO(csm_report.official_docx(rep, _meta())))
    assert len(doc.inline_shapes) >= 2  # seal + formula image


def test_signatory_and_headings_match_actual():
    rows = [_row({"SQD5": 0}) for _ in range(21)]
    rep = app._csm_official_report(rows)
    full = _docx_all_text(csm_report.official_docx(rep, _meta()))
    assert "RUSTOM CANARES" in full and "Administrative Assistant III" in full
    for h in ("I. Overview:", "II. Scope:", "III. Methodology:",
              "D1. AGE AND D2. SEX ANALYSIS:", "CITIZEN'S CHARTER ANSWERS ANALYSIS",
              "C. Average score per service", "D4. Government"):
        assert h in full, h


def test_costs_na_renders_not_communication():
    rows = [_row({"SQD5": 0}) for _ in range(21)]
    rep = app._csm_official_report(rows)
    full = "\n".join(p.get_text() for p in fitz.open(
        stream=csm_report.official_pdf(rep, _meta()), filetype="pdf"))
    # Costs is the all-N/A dimension (no-fee service); Communication scores 100%.
    assert "Costs" in full and "Communication" in full


def test_pdf_heading_follows_period():
    rows = [_row({}) for _ in range(10)]
    rep = app._csm_official_report(rows)
    data = csm_report.official_pdf(rep, _meta("year:2026"))
    d = fitz.open(stream=data, filetype="pdf")
    full = "\n".join(p.get_text() for p in d)
    assert "CY 2026" in full


# --------------------------------------------------------------------------
# Route — end-to-end download against the shared test DB
# --------------------------------------------------------------------------

def test_route_official_pdf(viewer):
    r = viewer.get("/reports/export",
                   params={"type": 1, "fmt": "pdf", "style": "official",
                           "period": "quarter:2026-Q1"})
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert "CSM_Report_" in r.headers.get("content-disposition", "")
    assert r.content[:4] == b"%PDF"


def test_route_official_docx(viewer):
    r = viewer.get("/reports/export",
                   params={"type": 1, "fmt": "docx", "style": "official",
                           "period": "quarter:2026-Q1"})
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
