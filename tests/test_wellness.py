"""
Wellness module (DOH DM 2025-0135) — the 2-hour weekly break: entitlement
validation, section-chief approver resolution, the apply→approve→reflect flow
with its gamification side-effects, in-handler RBAC on the chief/admin pages,
and the monitoring register export.
"""
import base64
import io
import re
from datetime import date, datetime, time, timedelta

import pytest

import auth
import storage as st
import tzutil as tz
from tests.conftest import _auth_client


def _next_weekday(d):
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def _future_weekday():
    """A weekday comfortably in the future. Claim creation requires a minimum
    lead time before the start (Spot Feedback #73), so tests that file at fixed
    clock times must not use *today* — their time would read as 'in the past'
    whenever the suite runs late in the day."""
    return _next_weekday(tz.today_local() + timedelta(days=3))


def _monday_this_week():
    t = tz.today_local()
    return t - timedelta(days=t.weekday())


@pytest.fixture()
def wellness_org(app_mod):
    """A section with a chief (section_manager account), an employee (viewer
    account, directory-linked) and a loner (viewer account, no directory link).
    Fresh rows per test; wellness tables are wiped so entitlement math and
    dashboards start from zero."""
    conn = st.connect(app_mod.DB)
    st.init_db(conn)
    conn.execute("DELETE FROM wellness_reflection")
    conn.execute("DELETE FROM wellness_badge")
    conn.execute("DELETE FROM wellness_claim")
    conn.execute("DELETE FROM admin_user WHERE username LIKE 'w-%@test.local'")
    conn.execute("DELETE FROM resource_person WHERE name LIKE 'W-Test %'")
    conn.execute("DELETE FROM org_unit WHERE path = 'TEST/WELL/WSec'")
    cur = conn.execute(
        "INSERT INTO org_unit(name, short_name, level, path) "
        "VALUES ('Wellness Test Section','WSec','section','TEST/WELL/WSec')")
    unit = cur.lastrowid
    chief_rp = conn.execute(
        "INSERT INTO resource_person(name, is_internal, org_unit_id, designation) "
        "VALUES ('W-Test Chief', 1, ?, 'Section Head')", (unit,)).lastrowid
    emp_rp = conn.execute(
        "INSERT INTO resource_person(name, is_internal, org_unit_id) "
        "VALUES ('W-Test Employee', 1, ?)", (unit,)).lastrowid
    h, salt = "x", "y"          # never logged in via password — cookies are minted
    chief = st.create_admin_user(conn, "w-chief@test.local", h, salt,
                                 "section_manager", rp_id=chief_rp)
    emp = st.create_admin_user(conn, "w-emp@test.local", h, salt,
                               "viewer", rp_id=emp_rp)
    loner = st.create_admin_user(conn, "w-loner@test.local", h, salt, "viewer")
    conn.commit()
    yield {"conn": conn, "unit": unit, "chief": chief, "emp": emp, "loner": loner}
    conn.close()


def _client_for_id(app_mod, uid):
    from starlette.testclient import TestClient
    conn = st.connect(app_mod.DB)
    u = st.get_admin_user_by_id(conn, uid)
    conn.close()
    c = TestClient(app_mod.app, raise_server_exceptions=True)
    c.cookies.set(auth.COOKIE_NAME, auth.make_token(
        {"id": u["id"], "username": u["username"], "role": u["role"]}))
    return c


# ---------------------------------------------------------------- entitlement

def test_third_hour_in_a_week_is_rejected(wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    st.create_wellness_claim(conn, emp, d, "09:00", "exercise")
    st.create_wellness_claim(conn, emp, d, "15:00", "meditation")
    with pytest.raises(ValueError, match="entitlement"):
        st.create_wellness_claim(conn, emp, d, "16:00", "hobby")


def test_two_hour_break_in_one_claim(wellness_org):
    """The DM allows the whole 2-hour weekly entitlement in a single sitting."""
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    cid = st.create_wellness_claim(conn, emp, d, "09:00", "exercise", hours=2)
    ws = st.wellness_week_start(date.fromisoformat(d))
    assert st.wellness_hours_used(conn, emp, ws) == 2
    assert st.get_wellness_claim(conn, cid)["hours"] == 2
    # The full entitlement is spent — no further hour fits that week.
    with pytest.raises(ValueError, match="entitlement"):
        st.create_wellness_claim(conn, emp, d, "15:00", "meditation")


def test_two_hour_claim_rejected_when_only_one_hour_left(wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    st.create_wellness_claim(conn, emp, d, "09:00", "exercise")     # 1 hour used
    with pytest.raises(ValueError, match="Only 1 hour left"):
        st.create_wellness_claim(conn, emp, d, "15:00", "meditation", hours=2)


def test_break_length_bounds(wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    with pytest.raises(ValueError, match="1 to 2 hours"):
        st.create_wellness_claim(conn, emp, d, "09:00", "exercise", hours=3)
    with pytest.raises(ValueError, match="1 to 2 hours"):
        st.create_wellness_claim(conn, emp, d, "09:00", "exercise", hours=0)


def test_denied_and_cancelled_hours_return_to_the_week(wellness_org):
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    u_chief = {"id": chief, "username": "w-chief@test.local", "role": "section_manager"}
    d = _future_weekday().isoformat()
    c1 = st.create_wellness_claim(conn, emp, d, "09:00", "exercise")
    c2 = st.create_wellness_claim(conn, emp, d, "15:00", "meditation")
    st.decide_wellness_claim(conn, c1, u_chief, "deny", "busy morning")
    st.cancel_wellness_claim(conn, c2, emp)
    ws = st.wellness_week_start(date.fromisoformat(d))
    assert st.wellness_hours_used(conn, emp, ws) == 0
    st.create_wellness_claim(conn, emp, d, "16:00", "hobby")       # fits again


def test_cancel_reason_is_recorded_and_optional(wellness_org):
    """Spot Feedback #77: a documented reason on withdrawal, same shape as
    decide_note (the approver's side) — present when given, blank when not."""
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    c1 = st.create_wellness_claim(conn, emp, d, "09:00", "exercise")
    c2 = st.create_wellness_claim(conn, emp, d, "15:00", "meditation")
    st.cancel_wellness_claim(conn, c1, emp, reason="  urgent office matter  ")
    st.cancel_wellness_claim(conn, c2, emp)                        # no reason given
    claim1 = st.get_wellness_claim(conn, c1)
    claim2 = st.get_wellness_claim(conn, c2)
    assert claim1["status"] == "cancelled"
    assert claim1["cancel_reason"] == "urgent office matter"       # trimmed
    assert claim2["cancel_reason"] is None


def test_http_withdraw_records_reason(app_mod, wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    claim_id = st.create_wellness_claim(conn, emp, d, "09:00", "exercise")
    c = _client_for_id(app_mod, emp)
    r = c.post(f"/wellness/claims/{claim_id}/cancel",
               data={"reason": "urgent work came up"}, follow_redirects=False)
    assert r.status_code == 303
    claim = st.get_wellness_claim(conn, claim_id)
    assert claim["status"] == "cancelled"
    assert claim["cancel_reason"] == "urgent work came up"


def test_claim_validation_rules(wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    with pytest.raises(ValueError):                                 # bad date
        st.create_wellness_claim(conn, emp, "not-a-date", "09:00", "exercise")
    with pytest.raises(ValueError):                                 # bad time
        st.create_wellness_claim(conn, emp, d, "25:99", "exercise")
    with pytest.raises(ValueError):                                 # bad activity
        st.create_wellness_claim(conn, emp, d, "09:00", "smoking")
    with pytest.raises(ValueError):                                 # backdated
        st.create_wellness_claim(
            conn, emp, (tz.today_local() - timedelta(days=8)).isoformat(),
            "09:00", "exercise")
    st.create_wellness_claim(conn, emp, d, "09:00", "exercise")
    with pytest.raises(ValueError, match="already have a break"):   # duplicate slot
        st.create_wellness_claim(conn, emp, d, "09:00", "hobby")


def test_break_needs_minimum_lead_time(wellness_org, monkeypatch):
    """Spot Feedback #73: a break must be filed at least
    WELLNESS_MIN_LEAD_MINUTES before it starts. Pin 'now' to a fixed weekday
    morning so the window is deterministic regardless of when the suite runs."""
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    fixed = datetime.combine(_monday_this_week(), time(10, 0))   # a Monday, 10:00
    monkeypatch.setattr(tz, "now_local", lambda: fixed)
    monkeypatch.setattr(tz, "today_local", lambda: fixed.date())
    day = fixed.date().isoformat()
    lead = st.WELLNESS_MIN_LEAD_MINUTES
    # too soon: only (lead - 15) minutes ahead of the start
    soon = (fixed + timedelta(minutes=lead - 15)).strftime("%H:%M")
    with pytest.raises(ValueError, match="minutes before"):
        st.create_wellness_claim(conn, emp, day, soon, "exercise")
    # a start already passed today is caught by the same rule (date-only check
    # used to let it through)
    with pytest.raises(ValueError, match="minutes before"):
        st.create_wellness_claim(conn, emp, day, "09:00", "exercise")
    # comfortably ahead (lead + 60 min) is accepted
    ok = (fixed + timedelta(minutes=lead + 60)).strftime("%H:%M")
    assert st.create_wellness_claim(conn, emp, day, ok, "meditation")


# ---------------------------------------------------------- approver routing

def test_chief_resolved_through_directory(wellness_org):
    conn = wellness_org["conn"]
    assert st.get_wellness_approvers(conn, wellness_org["emp"]) == [wellness_org["chief"]]
    # no directory link -> nobody resolved; admins are the fallback
    assert st.get_wellness_approvers(conn, wellness_org["loner"]) == []
    u_admin = {"id": 0, "role": "admin"}
    assert st.is_wellness_approver(conn, u_admin, wellness_org["loner"])


def test_nobody_decides_their_own_claim(wellness_org):
    conn, chief = wellness_org["conn"], wellness_org["chief"]
    u_chief = {"id": chief, "username": "w-chief@test.local", "role": "section_manager"}
    d = _future_weekday().isoformat()
    own = st.create_wellness_claim(conn, chief, d, "09:00", "exercise")
    assert not st.is_wellness_approver(conn, u_chief, chief)
    with pytest.raises(PermissionError):
        st.decide_wellness_claim(conn, own, u_chief, "approve")
    # the chief's own request lands in the admins' queue instead
    u_admin = {"id": 0, "role": "admin"}
    assert own in [c["id"] for c in st.list_wellness_approvals(conn, u_admin)]


def test_chief_queue_holds_only_their_unit(wellness_org):
    conn, emp, chief, loner = (wellness_org["conn"], wellness_org["emp"],
                               wellness_org["chief"], wellness_org["loner"])
    u_chief = {"id": chief, "username": "w-chief@test.local", "role": "section_manager"}
    d = _future_weekday().isoformat()
    inside = st.create_wellness_claim(conn, emp, d, "09:00", "exercise")
    outside = st.create_wellness_claim(conn, loner, d, "09:00", "exercise")
    ids = [c["id"] for c in st.list_wellness_approvals(conn, u_chief)]
    assert inside in ids and outside not in ids


# -------------------------------------------------- escalation up the full chain

@pytest.fixture()
def wellness_chain(app_mod):
    """A full division -> section -> directorate slice with an account at every
    tier, so escalation can be walked end to end (Spot Feedback #76). Crucially
    the Director carries their rank in `position` ('Director IV'), NOT in
    `designation` — mirroring the real directory, where an exact designation
    match found nobody."""
    conn = st.connect(app_mod.DB)
    st.init_db(conn)

    def _wipe():
        # SCOPED to this fixture's namespace. The Director tier is searched
        # GLOBALLY (no org filter), so a leaked 'Director IV' account would
        # pollute other files' escalation tests — clean up at BOTH ends.
        conn.execute("DELETE FROM wellness_claim WHERE user_id IN "
                     "(SELECT id FROM admin_user WHERE username LIKE 'wc-%@test.local')")
        conn.execute("DELETE FROM admin_user WHERE username LIKE 'wc-%@test.local'")
        conn.execute("DELETE FROM resource_person WHERE name LIKE 'WC-%'")
        conn.execute("DELETE FROM org_unit WHERE path LIKE 'TC/%' OR path = 'TC'")
        conn.commit()
    _wipe()

    def unit(name, short, level, path):
        return conn.execute(
            "INSERT INTO org_unit(name, short_name, level, path) VALUES (?,?,?,?)",
            (name, short, level, path)).lastrowid

    def person(name, desig, pos, org):
        return conn.execute(
            "INSERT INTO resource_person(name, is_internal, org_unit_id, designation, position) "
            "VALUES (?,1,?,?,?)", (name, org, desig, pos)).lastrowid

    unit("Test Bureau", "TC", "bureau", "TC")
    u_od = unit("Office of the Director", "OD", "office", "TC/OD")
    u_div = unit("Test Division", "TDIV", "division", "TC/TDIV")
    u_sec = unit("Test Section", "TSEC", "section", "TC/TDIV/TSEC")

    # Director's leadership lives in `position`, designation left blank on purpose.
    rp_dir = person("WC-Director", None, "Director IV", u_od)
    rp_div = person("WC-DivChief", "Division Chief", "Medical Officer V", u_div)
    rp_sec = person("WC-SecHead", "Section Head", "Chief Health Program Officer", u_sec)
    rp_emp = person("WC-Employee", None, "Health Program Officer II", u_sec)

    h, salt = "x", "y"
    mk = lambda uname, role, rp: st.create_admin_user(conn, uname, h, salt, role, rp_id=rp)
    director = mk("wc-dir@test.local", "section_manager", rp_dir)
    divchief = mk("wc-div@test.local", "section_manager", rp_div)
    sechead = mk("wc-sec@test.local", "section_manager", rp_sec)
    emp = mk("wc-emp@test.local", "viewer", rp_emp)
    conn.commit()
    yield {"conn": conn, "director": director, "divchief": divchief,
           "sechead": sechead, "emp": emp, "rp_sec": rp_sec}
    _wipe()
    conn.close()


def test_escalation_walks_each_tier(wellness_chain):
    """staff -> section head -> division chief -> Director, one hop per rank."""
    conn, w = wellness_chain["conn"], wellness_chain
    assert st.get_wellness_approvers(conn, w["emp"]) == [w["sechead"]]
    assert st.get_wellness_approvers(conn, w["sechead"]) == [w["divchief"]]
    # the Director is matched by POSITION ('Director IV'), not designation —
    # the exact-match bug this fix corrects.
    assert st.get_wellness_approvers(conn, w["divchief"]) == [w["director"]]
    # the Director sits at the top: nobody here, admins are the fallback.
    assert st.get_wellness_approvers(conn, w["director"]) == []
    u_admin = {"id": 0, "role": "admin"}
    assert st.is_wellness_approver(conn, u_admin, w["director"])


def test_director_and_chiefs_see_the_queue(wellness_chain):
    """Everyone who can approve SOMEONE reaches the approvals queue — including
    the Director, whose rank is carried in `position`."""
    conn = wellness_chain["conn"]
    for key in ("sechead", "divchief", "director"):
        u = {"id": wellness_chain[key], "role": "section_manager"}
        assert st.user_is_wellness_chief(conn, u), key
    # a plain employee does not
    assert not st.user_is_wellness_chief(
        conn, {"id": wellness_chain["emp"], "role": "viewer"})


def test_title_names_the_seat_even_without_an_account(wellness_chain):
    """The apply-page copy names the org-correct approver as long as that seat is
    filled in the directory — it must not fall back to 'an administrator' just
    because the chief has no app account yet (the reported symptom)."""
    conn, w = wellness_chain["conn"], wellness_chain
    assert st.wellness_approver_title(conn, w["emp"]) == "your section chief"
    assert st.wellness_approver_title(conn, w["sechead"]) == "your division chief"
    assert st.wellness_approver_title(conn, w["divchief"]) == "the Director"
    assert st.wellness_approver_title(conn, w["director"]) == "an administrator"

    # Drop the division chief's ACCOUNT but keep the person in the directory:
    # the section head's copy still reads 'your division chief' (the seat is
    # filled), while the actual claim now escalates to the admin fallback.
    conn.execute("DELETE FROM admin_user WHERE id=?", (w["divchief"],))
    conn.commit()
    assert st.wellness_approver_title(conn, w["sechead"]) == "your division chief"
    assert st.get_wellness_approvers(conn, w["sechead"]) == []


# ------------------------------------------- the full flow + gamification

def test_approve_complete_reflect_awards_points_and_badges(app_mod, wellness_org):
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    u_chief = {"id": chief, "username": "w-chief@test.local", "role": "section_manager"}
    monday = _monday_this_week().isoformat()     # always a weekday <= today
    # inserted directly (create_wellness_claim only files today-or-later; the
    # completion path is what's under test here)
    ws = st.wellness_week_start(date.fromisoformat(monday))
    cid = conn.execute(
        "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
        " hours, activity, status, created_at) VALUES (?,?,?,?,1,'exercise','pending',?)",
        (emp, ws, monday, "09:00", tz.now_iso())).lastrowid
    conn.commit()
    st.decide_wellness_claim(conn, cid, u_chief, "approve")
    claim, badges = st.complete_wellness_claim(
        conn, cid, emp, 2, 4, "brisk walk", "felt lighter after")
    assert claim["status"] == "completed"
    assert "First break" in badges
    stats = st.wellness_stats(conn, emp)
    # +10 break, +5 reflection, +2 for the 1-week streak
    assert stats["points"] == 17
    assert stats["streak"] == 1
    assert stats["stage"] == 1 and stats["level"] == 1
    assert [b["key"] for b in stats["badges"]] == ["first_break"]
    # a second completed hour that week earns the full-week bonus + badge
    cid2 = conn.execute(
        "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
        " hours, activity, status, created_at) VALUES (?,?,?,?,1,'hobby','approved',?)",
        (emp, ws, monday, "15:00", tz.now_iso())).lastrowid
    conn.commit()
    _, badges2 = st.complete_wellness_claim(conn, cid2, emp, 3, 5, "sketching", None)
    assert "Full week" in badges2
    stats = st.wellness_stats(conn, emp)
    # 2 breaks (20) + 2 reflections (10) + full week (5) + streak 1 (2)
    assert stats["points"] == 37 and stats["level"] == 2


def test_completion_guards(wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    cid = st.create_wellness_claim(conn, emp, d, "09:00", "exercise")
    with pytest.raises(ValueError, match="approved"):   # still pending
        st.complete_wellness_claim(conn, cid, emp, 3, 4, "walk", None)
    with pytest.raises(ValueError):                     # not yours
        st.complete_wellness_claim(conn, cid, wellness_org["loner"], 3, 4, "walk", None)


# ----------------------------------------------------------------- HTTP layer

def test_http_apply_flow_and_hub(app_mod, wellness_org):
    c = _client_for_id(app_mod, wellness_org["emp"])
    r = c.get("/wellness")
    # the companion now renders via the shared avatar macro (the mascot is the
    # adopted pet); the .avwrap wrapper is its structural marker
    assert r.status_code == 200 and "avwrap" in r.text
    r = c.get("/wellness/apply")
    assert r.status_code == 200
    # Phase 6a: the activity selector is now illustrated cards — inline SVG icons
    # behind the SAME radio inputs (name="activity"), so the POST below is unchanged.
    assert "<svg" in r.text
    assert 'type="radio" name="activity" value="exercise"' in r.text
    # the break-length picker offers both 1- and 2-hour options
    assert 'name="hours" value="1"' in r.text and 'name="hours" value="2"' in r.text
    d = _future_weekday().isoformat()
    # take the whole 2-hour entitlement in a single claim
    r = c.post("/wellness/apply", data={"break_date": d, "start_time": "10:00",
                                        "activity": "exercise", "note": "jog",
                                        "hours": "2"},
               follow_redirects=False)
    assert r.status_code == 303
    # no hours left this week: the form re-renders with the entitlement error
    r = c.post("/wellness/apply", data={"break_date": d, "start_time": "16:00",
                                        "activity": "hobby"})
    assert r.status_code == 400 and "entitlement" in r.text


def test_http_approve_and_reflect(app_mod, wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    monday = _monday_this_week().isoformat()
    ws = st.wellness_week_start(date.fromisoformat(monday))
    cid = conn.execute(
        "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
        " hours, activity, status, created_at) VALUES (?,?,?,?,1,'exercise','pending',?)",
        (emp, ws, monday, "09:00", tz.now_iso())).lastrowid
    conn.commit()
    chief_c = _client_for_id(app_mod, wellness_org["chief"])
    r = chief_c.get("/wellness/approvals")
    assert r.status_code == 200 and "W-Test Employee" in r.text
    r = chief_c.post(f"/wellness/claims/{cid}/decide",
                     data={"action": "approve"}, follow_redirects=False)
    assert r.status_code == 303
    emp_c = _client_for_id(app_mod, emp)
    r = emp_c.get(f"/wellness/claims/{cid}/complete")
    assert r.status_code == 200
    r = emp_c.post(f"/wellness/claims/{cid}/complete",
                   data={"mood_before": "2", "mood_after": "4",
                         "activity_done": "brisk walk", "takeaway": "nice"},
                   follow_redirects=False)
    assert r.status_code == 303 and "cheer=" in r.headers["location"]
    conn2 = st.connect(app_mod.DB)
    assert st.get_wellness_claim(conn2, cid)["status"] == "completed"
    conn2.close()


def test_chief_pages_are_gated(app_mod, wellness_org, viewer, admin):
    # a viewer with no chief designation: hub yes, chief pages no
    assert viewer.get("/wellness").status_code == 200
    for path in ("/wellness/approvals", "/wellness/dashboard", "/wellness/export.xlsx"):
        assert viewer.get(path).status_code == 403, path
    # admins and directory-resolved chiefs both pass
    chief_c = _client_for_id(app_mod, wellness_org["chief"])
    for c in (admin, chief_c):
        for path in ("/wellness/approvals", "/wellness/dashboard"):
            assert c.get(path).status_code == 200, path
    # deciding someone else's claim from a non-approver account bounces
    d = _future_weekday().isoformat()
    conn = wellness_org["conn"]
    cid = st.create_wellness_claim(conn, wellness_org["emp"], d, "11:00", "exercise")
    loner_c = _client_for_id(app_mod, wellness_org["loner"])
    assert loner_c.post(f"/wellness/claims/{cid}/decide",
                        data={"action": "approve"}).status_code == 403


def test_register_export(app_mod, wellness_org, admin):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    monday = _monday_this_week().isoformat()
    ws = st.wellness_week_start(date.fromisoformat(monday))
    cid = conn.execute(
        "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
        " hours, activity, status, created_at) VALUES (?,?,?,?,1,'meditation','approved',?)",
        (emp, ws, monday, "13:00", tz.now_iso())).lastrowid
    conn.commit()
    st.complete_wellness_claim(conn, cid, emp, 1, 5, "breathing exercises", "calm",
                               usefulness=4)
    # a still-pending break must NOT reach the register — only consumed breaks count
    conn.execute(
        "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
        " hours, activity, status, created_at) VALUES (?,?,?,?,2,'exercise','pending',?)",
        (emp, ws, monday, "09:00", tz.now_iso()))
    conn.commit()

    r = admin.get("/wellness/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    from openpyxl import load_workbook
    sh = load_workbook(io.BytesIO(r.content)).active
    grid = [[c.value for c in row] for row in sh.iter_rows()]
    text = "\n".join("|".join("" if v is None else str(v) for v in row) for row in grid)
    # the Annex A masthead and its seven columns
    assert "Annex A. Weekly Wellness Break Log" in text
    assert "Bureau of Local Health Systems Development" in text
    for col in ("No.", "Name of Personnel", "Wellness Break Availed (Date/Time)",
                "Duration (Minutes)", "Usefulness Rating (1-5)", "Remarks/ Observations"):
        assert col in text, col
    # exactly the one consumed break, mapped into the Annex A layout
    ours = [row for row in grid if (row[1] or "") == "W-Test Employee"]
    assert len(ours) == 1
    row = ours[0]
    assert row[3] == "Relaxation"   # meditation -> Annex A "Relaxation"
    assert row[4] == 60             # 1 hour -> 60 minutes
    assert row[5] == 4             # usefulness rating
    assert row[6] == "calm"        # takeaway -> Remarks/Observations
    # the pending 09:00 break never appears on the register
    assert "09:00" not in text
    # the dashboard shows the same week
    r = admin.get(f"/wellness/dashboard?week={monday}")
    assert r.status_code == 200 and "WSec" in r.text
    # the dashboard's download control carries the date-range filter
    assert 'action="/wellness/export.xlsx"' in r.text
    assert 'name="date_from"' in r.text and 'name="date_to"' in r.text


def test_register_export_date_range(app_mod, wellness_org, admin):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    monday = _monday_this_week()
    ws = st.wellness_week_start(monday)
    early = monday.isoformat()
    late = (monday + timedelta(days=2)).isoformat()
    for d, t in ((early, "10:00"), (late, "15:00")):
        conn.execute(
            "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
            " hours, activity, status, created_at) VALUES (?,?,?,?,1,'exercise','completed',?)",
            (emp, ws, d, t, tz.now_iso()))
    conn.commit()

    from openpyxl import load_workbook

    def register(**params):
        r = admin.get("/wellness/export.xlsx", params=params)
        assert r.status_code == 200
        sh = load_workbook(io.BytesIO(r.content)).active
        return "\n".join("|".join("" if c.value is None else str(c.value) for c in row)
                         for row in sh.iter_rows())

    # bound the range to the early day only -> the late break is filtered out
    text = register(date_from=early, date_to=early)
    assert early in text and late not in text
    # the requested bounds drive the "Weekly Covered" masthead
    assert f"Weekly Covered: {app_mod._fmt_reg_date(early)}" in text
    # widening the upper bound brings the late break back in
    text = register(date_from=early, date_to=late)
    assert early in text and late in text
    # an open lower bound (date_to only) still filters the far side
    assert late not in register(date_to=early)


# ---------------------------------------- "Others" activity needs a specification

def test_other_activity_requires_and_stores_specification(wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    d = _future_weekday().isoformat()
    # "other" with no specification is rejected
    with pytest.raises(ValueError, match="Others"):
        st.create_wellness_claim(conn, emp, d, "09:00", "other")
    with pytest.raises(ValueError, match="Others"):
        st.create_wellness_claim(conn, emp, d, "09:00", "other", activity_other="  ")
    # with a specification it is stored and surfaced in the label
    cid = st.create_wellness_claim(conn, emp, d, "10:00", "other",
                                   activity_other="community gardening")
    claim = st.get_wellness_claim(conn, cid)
    assert claim["activity_other"] == "community gardening"
    assert claim["activity_label"] == "Others — community gardening"
    # a non-"other" activity never keeps a stray specification
    cid2 = st.create_wellness_claim(conn, emp, d, "11:00", "exercise",
                                    activity_other="ignored")
    assert st.get_wellness_claim(conn, cid2)["activity_other"] is None


# ---------------------------------------------- rolling 4-week request window

def test_request_allowed_across_the_month(wellness_org):
    """A break can be filed for a future week within the window; each week keeps
    its own 2-hour cap, independent of the current week."""
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    next_week = _next_weekday(_monday_this_week() + timedelta(days=7))
    d = next_week.isoformat()
    cid = st.create_wellness_claim(conn, emp, d, "09:00", "exercise", hours=2)
    ws = st.wellness_week_start(next_week)
    assert st.wellness_hours_used(conn, emp, ws) == 2
    # the current week is untouched by next week's claim
    assert st.wellness_hours_used(conn, emp, st.wellness_week_start()) == 0


def test_hub_link_stays_enabled_once_a_future_week_is_open(app_mod, wellness_org):
    """Using up this week's 2 hours must not lock the 'Request a break' link —
    the apply flow still accepts a claim for any open week in the rolling
    window, so the hub link should only disable when EVERY week in the
    window is full."""
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    monday = _monday_this_week().isoformat()
    ws = st.wellness_week_start(date.fromisoformat(monday))
    conn.execute(
        "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
        " hours, activity, status, created_at) VALUES (?,?,?,?,2,'exercise','approved',?)",
        (emp, ws, monday, "09:00", tz.now_iso()))
    conn.commit()
    assert st.wellness_hours_used(conn, emp, ws) == st.WELLNESS_WEEKLY_HOURS
    c = _client_for_id(app_mod, emp)
    r = c.get("/wellness")
    assert r.status_code == 200
    m = re.search(r'<a class="btn" href="/wellness/apply"[^>]*>', r.text)
    assert m and "aria-disabled" not in m.group(0)
    # a next-week claim is still accepted through the real HTTP flow
    next_week = _next_weekday(_monday_this_week() + timedelta(days=7))
    r = c.post("/wellness/apply", data={"break_date": next_week.isoformat(),
                                        "start_time": "10:00", "activity": "exercise",
                                        "note": "jog", "hours": "1"},
               follow_redirects=False)
    assert r.status_code == 303


def test_request_beyond_window_is_rejected(wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    # a workday past the last week of the rolling window
    beyond = _next_weekday(_monday_this_week()
                           + timedelta(days=st.WELLNESS_WINDOW_WEEKS * 7))
    assert beyond > st.wellness_window_last_date()
    with pytest.raises(ValueError, match="weeks ahead"):
        st.create_wellness_claim(conn, emp, beyond.isoformat(), "09:00", "exercise")


# ------------------------------------------- richer post-activity report fields

def test_completion_records_post_activity_fields(app_mod, wellness_org):
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    u_chief = {"id": chief, "username": "w-chief@test.local", "role": "section_manager"}
    monday = _monday_this_week().isoformat()
    ws = st.wellness_week_start(date.fromisoformat(monday))
    cid = conn.execute(
        "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
        " hours, activity, status, created_at) VALUES (?,?,?,?,1,'exercise','approved',?)",
        (emp, ws, monday, "09:00", tz.now_iso())).lastrowid
    conn.commit()
    st.complete_wellness_claim(
        conn, cid, emp, 2, 4, "brisk walk", "felt lighter",
        usefulness="5", felt_refreshed="1", managed_stress="0",
        focus_impact="improved", suggestions="more shaded areas")
    claim = st.get_wellness_claim(conn, cid)
    assert claim["usefulness"] == 5
    assert claim["felt_refreshed"] == 1 and claim["managed_stress"] == 0
    assert claim["focus_impact"] == "improved"
    assert claim["suggestions"] == "more shaded areas"


def test_completion_validates_new_fields(wellness_org):
    conn, emp = wellness_org["conn"], wellness_org["emp"]
    monday = _monday_this_week().isoformat()
    ws = st.wellness_week_start(date.fromisoformat(monday))

    def _fresh(start):
        return conn.execute(
            "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
            " hours, activity, status, created_at) VALUES (?,?,?,?,1,'exercise','approved',?)",
            (emp, ws, monday, start, tz.now_iso())).lastrowid

    c1 = _fresh("09:00")
    with pytest.raises(ValueError, match="Usefulness"):
        st.complete_wellness_claim(conn, c1, emp, 2, 4, "walk", None, usefulness="9")
    c2 = _fresh("10:00")
    with pytest.raises(ValueError, match="focus"):
        st.complete_wellness_claim(conn, c2, emp, 2, 4, "walk", None, focus_impact="great")
    # the new fields are optional — the legacy 6-arg call still works
    c3 = _fresh("11:00")
    st.complete_wellness_claim(conn, c3, emp, 2, 4, "walk", None)
    assert st.get_wellness_claim(conn, c3)["status"] == "completed"


# --------------------------------------------------- approver named, not generic

def test_approver_name_resolves_seat_holder(wellness_chain):
    conn, w = wellness_chain["conn"], wellness_chain
    # staff -> section chief and section chief -> division chief are org-scoped,
    # so they name exactly the one seat holder in that unit.
    assert st.wellness_approver_name(conn, w["emp"]) == "WC-SecHead"
    assert st.wellness_approver_name(conn, w["sechead"]) == "WC-DivChief"
    # division chief -> the Director is searched bureau-wide (the top tier has no
    # org scope), so just assert our Director is named among them.
    assert "WC-Director" in (st.wellness_approver_name(conn, w["divchief"]) or "")
    # the Director sits at the top — no named approver, callers fall back to title
    assert st.wellness_approver_name(conn, w["director"]) is None


def test_apply_page_names_the_approver(app_mod, wellness_org):
    c = _client_for_id(app_mod, wellness_org["emp"])
    r = c.get("/wellness/apply")
    assert r.status_code == 200
    # the section chief's actual name appears on the submit button
    assert "W-Test Chief" in r.text
    # the "Others" specify field is present for progressive reveal
    assert 'name="activity_other"' in r.text


def test_home_shows_the_two_doors(viewer):
    r = viewer.get("/")
    assert r.status_code == 200
    assert "Wellness — my 2-hour break" in r.text
    assert 'href="/wellness"' in r.text


# ---------------------------------------------- monthly register (PDF) --------

def _seed_availed(conn, uid, chief, break_date, start_time, status, activity="exercise"):
    """Insert a break in a known state, decided by `chief` — the register only
    lists availed (approved/completed) breaks, so tests need real decided rows."""
    ws = st.wellness_week_start(date.fromisoformat(break_date))
    conn.execute(
        "INSERT INTO wellness_claim(user_id, week_start, break_date, start_time,"
        " hours, activity, status, decided_by, decided_at, created_at)"
        " VALUES (?,?,?,?,1,?,?,?,?,?)",
        (uid, ws, break_date, start_time, activity, status, chief,
         tz.now_iso(), tz.now_iso()))
    conn.commit()


def test_availed_months_and_person_month(wellness_org):
    import wellness_register as wreg
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    _seed_availed(conn, emp, chief, "2026-03-03", "10:00", "completed")
    _seed_availed(conn, emp, chief, "2026-03-17", "14:00", "approved", activity="hobby")
    _seed_availed(conn, emp, chief, "2026-04-07", "09:00", "completed")
    # a denied break must NOT count as availed
    _seed_availed(conn, emp, chief, "2026-03-24", "11:00", "denied")

    months = st.wellness_availed_months(conn, emp)
    yms = [m["ym"] for m in months]
    assert yms == ["2026-04", "2026-03"]          # newest first
    mar = next(m for m in months if m["ym"] == "2026-03")
    assert mar["count"] == 2 and mar["label"] == "March 2026"

    data = st.wellness_person_month(conn, emp, 2026, 3)
    assert data["name"] == "W-Test Employee" and data["office"] == "WSec"
    assert len(data["entries"]) == 2              # the denied one is excluded
    assert {e["date"] for e in data["entries"]} == {"2026-03-03", "2026-03-17"}
    assert all(e["approver"] == "W-Test Chief" for e in data["entries"])

    pdf = wreg.render_month_pdf(data)
    assert pdf[:4] == b"%PDF"
    assert wreg.filename(data) == "Wellness-Register-Employee-2026-03.pdf"


def test_register_pdf_download_and_hub_control(app_mod, wellness_org):
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    _seed_availed(conn, emp, chief, "2026-03-10", "10:00", "completed")
    c = _client_for_id(app_mod, emp)

    # the hub offers the monthly-register download once a break is availed
    r = c.get("/wellness")
    assert r.status_code == 200
    assert "/wellness/register.pdf" in r.text and "March 2026" in r.text

    # a real month yields a PDF attachment
    r = c.get("/wellness/register.pdf?month=2026-03")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert "Wellness-Register-Employee-2026-03.pdf" in r.headers["content-disposition"]
    assert r.content[:4] == b"%PDF"

    # a month with nothing availed bounces back to the hub with a message
    r = c.get("/wellness/register.pdf?month=2025-01", follow_redirects=False)
    assert r.status_code == 303 and "err=" in r.headers["location"]


# A real (tiny) PNG — the pad exports PNG, and the register only embeds what
# reportlab can actually open.
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")


def test_person_month_carries_the_approvers_signature(wellness_org):
    """The register data must hand the renderer the DECIDING chief's signature —
    read once per approver, not once per break."""
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    chief_rp = st.get_admin_user_by_id(conn, chief)["rp_id"]
    st.set_rp_signature(conn, chief_rp, _PNG, "image/png")
    _seed_availed(conn, emp, chief, "2026-03-03", "10:00", "completed")
    _seed_availed(conn, emp, chief, "2026-03-10", "10:00", "approved")

    data = st.wellness_person_month(conn, emp, 2026, 3)
    assert len(data["entries"]) == 2
    for e in data["entries"]:
        assert e["approver_signature"] == _PNG
        assert e["approver_signature_mime"] == "image/png"
    # one blob object shared by both entries: the read is per approver, not per break
    assert data["entries"][0]["approver_signature"] is data["entries"][1]["approver_signature"]


def test_person_month_falls_back_to_the_account_signature(wellness_org):
    """A chief with no directory row signs from their account instead."""
    conn, emp, loner = wellness_org["conn"], wellness_org["emp"], wellness_org["loner"]
    st.set_user_signature(conn, loner, _PNG, "image/png")
    _seed_availed(conn, emp, loner, "2026-05-06", "10:00", "completed")

    data = st.wellness_person_month(conn, emp, 2026, 5)
    assert data["entries"][0]["approver_signature"] == _PNG


def test_person_month_without_a_signature_carries_none(wellness_org):
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    _seed_availed(conn, emp, chief, "2026-03-03", "10:00", "completed")
    data = st.wellness_person_month(conn, emp, 2026, 3)
    assert data["entries"][0]["approver_signature"] is None


def test_day_cell_signs_on_the_day_of_the_break(wellness_org):
    """The signature belongs to the DAY, not the week: the flowable sits in the
    weekday cell, and the week's left band is only the form's label."""
    import wellness_register as wreg
    from reportlab.platypus import Image
    entry = {"date": "2026-03-03", "start_time": "10:00", "hours": 1,
             "activity_label": "Physical exercise", "status": "completed",
             "approver": "W-Test Chief",
             "approver_signature": _PNG, "approver_signature_mime": "image/png"}
    flows = wreg.build_day_cell([entry], 78, 20)
    assert any(isinstance(f, Image) for f in flows)
    text = " ".join(getattr(f, "text", "") for f in flows)
    assert "W-Test Chief" in text and "Digitally approved and signed" in text
    assert wreg._approver_html([entry]) == "Approved By:"

    # no signature on file: same cell, image dropped, stamp intact
    entry = dict(entry, approver_signature=None, approver_signature_mime=None)
    flows = wreg.build_day_cell([entry], 78, 20)
    assert not any(isinstance(f, Image) for f in flows)
    text = " ".join(getattr(f, "text", "") for f in flows)
    assert "W-Test Chief" in text and "Digitally approved and signed" in text


def test_register_pdf_embeds_the_signature_on_the_day(wellness_org):
    """End to end: a signed approval puts an image in the PDF; an unsigned one
    still prints the name over the stamp, and the download never fails."""
    import wellness_register as wreg
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    chief_rp = st.get_admin_user_by_id(conn, chief)["rp_id"]
    _seed_availed(conn, emp, chief, "2026-03-03", "10:00", "completed")

    unsigned = wreg.render_month_pdf(st.wellness_person_month(conn, emp, 2026, 3))
    assert unsigned[:4] == b"%PDF" and b"/Subtype /Image" not in unsigned

    st.set_rp_signature(conn, chief_rp, _PNG, "image/png")
    signed = wreg.render_month_pdf(st.wellness_person_month(conn, emp, 2026, 3))
    assert signed[:4] == b"%PDF" and b"/Subtype /Image" in signed


def test_register_pdf_survives_a_corrupt_signature(wellness_org):
    """A blob that won't open must degrade to the text stamp, not 500 the
    download — the register is the employee's own record, not the chief's."""
    import wellness_register as wreg
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    chief_rp = st.get_admin_user_by_id(conn, chief)["rp_id"]
    st.set_rp_signature(conn, chief_rp, b"not an image at all", "image/png")
    _seed_availed(conn, emp, chief, "2026-03-03", "10:00", "completed")

    pdf = wreg.render_month_pdf(st.wellness_person_month(conn, emp, 2026, 3))
    assert pdf[:4] == b"%PDF" and b"/Subtype /Image" not in pdf


def test_register_pdf_is_self_service_only(app_mod, wellness_org):
    # a viewer with no availed breaks still reaches the route (no chief gate),
    # they just get redirected because there is nothing to render
    c = _client_for_id(app_mod, wellness_org["loner"])
    r = c.get("/wellness/register.pdf", follow_redirects=False)
    assert r.status_code == 303


# ---------------------------------------------------------------------------
# UX friction #7 — the "this week" card was a dead target
# ---------------------------------------------------------------------------
# 30-day telemetry: 33 dead clicks from 9 users on the card body (friction rank
# 7), with 10 matching Spot Feedback comments. Only the trailing button was
# clickable; people clicked the status text and the hour pips. The card is now a
# stretched-link, which must NOT introduce a second tab stop, and the disabled
# state must stop being a real anchor.

def _week_card(html):
    """The 'This week (…)' panel, sliced out of the hub so assertions can't be
    satisfied by some other link elsewhere on the page."""
    start = html.index("This week (")
    start = html.rindex('<div class="panel entitle', 0, start)
    return html[start:html.index("</div>", html.index("Request a break"))]


def test_week_card_is_one_stretched_link(app_mod, wellness_org):
    c = _client_for_id(app_mod, wellness_org["emp"])
    card = _week_card(c.get("/wellness").text)

    # the card advertises itself as clickable...
    assert "entitle tappable" in card
    # ...via exactly ONE anchor — the stretched-link overlay must not add a
    # second focusable control over the same destination
    assert card.count('href="/wellness/apply"') == 1
    assert card.count("<a ") == 1


def test_week_card_disabled_state_is_not_a_link(app_mod, wellness_org):
    """With every hour in the rolling window claimed there is nothing to request.
    The control used to stay an <a href> carrying aria-disabled: the CSS
    pointer-events:none stopped a mouse, but it was still focusable and Enter
    still navigated, so keyboard and screen-reader users were told it was
    disabled and then taken to the form anyway."""
    conn, emp, chief = wellness_org["conn"], wellness_org["emp"], wellness_org["chief"]
    for ws in st.wellness_window_weeks():
        for hour in ("09:00", "14:00"):                 # 2 h = the whole week
            _seed_availed(conn, emp, chief, ws, hour, "approved")

    card = _week_card(_client_for_id(app_mod, emp).get("/wellness").text)

    assert 'href="/wellness/apply"' not in card         # nothing to activate
    assert 'aria-disabled="true"' in card               # still announced as off
    assert "tappable" not in card                       # and not advertised as clickable
