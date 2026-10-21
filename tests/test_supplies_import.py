"""S1.3 synthetic workbook parsing, staged evidence and inventory isolation."""
from datetime import datetime
from io import BytesIO
import json
import sqlite3
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
import pytest
from starlette.testclient import TestClient

import auth
import storage as st
import supplies as s
import supplies_import as imports
import supplies_stock as stock
from supplies_storage import Conflict, Denied, Disabled, Invalid
from tests.test_supplies import supplies_db, setup_catalog, create, key, hidden
from tests.conftest import session_token


def rewrite(data, mutate):
    out = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(out, 'w', ZIP_DEFLATED) as dest:
        for name in source.namelist():
            dest.writestr(name, mutate(name, source.read(name)))
    return out.getvalue()


def workbook(*, caches=True, title='Test Card', unit='Ream', summary_formula=None, latest_formula=None,
             opening_formula=None, gap=False, external=False, long_text=False):
    w = openpyxl.Workbook()
    summary = w.active
    summary.title = 'Summary Sheet'
    for c, value in [(1,'ARTICLE'),(6,'UNIT OF MEASURE'),(9,'BALANCE'),(16,'Critical Value')]:
        summary.cell(5,c,value)
    summary['A7'] = 'Synthetic paper' if not long_text else 'x'*4001
    summary['F7'] = 'Ream'
    summary['I7'] = summary_formula or ("='" + title.replace("'", "''") + "'!F27")
    summary['P7'] = 0
    card = w.create_sheet(title)
    card['A2'],card['A4'],card['A6'] = 'STOCK CARD','NAME OF ITEM: Synthetic paper','UNIT: '+unit
    for c,value in [(1,'DATE'),(2,'RECEIVED'),(4,'ISSUED'),(6,'BALANCE')]:
        card.cell(22,c,value)
    card['A23'],card['F23'] = datetime(2026,1,1), opening_formula or 10
    r = 25 if gap else 24
    card.cell(r,1,datetime(2026,1,2)); card.cell(r,2,5)
    card.cell(r,6,latest_formula or f'=IF(A{r}="","",SUM(F{r-1}+B{r}-D{r}))')
    card['F26'] = '=IF(A26="","",SUM(F25+B26-D26))'
    card['A27'],card['B27'],card['F27'] = 'Total Balance','=SUMIF(J23:J26,"<=1",B23:B26)','=SUM(F23+B27-D27)'
    if external:
        summary['Q7'].hyperlink = 'https://invalid.example/never-follow'
    stream = BytesIO(); w.save(stream); w.close()
    ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    def cache(name, data):
        if not caches or name not in ('xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml'):
            return data
        root = ET.fromstring(data)
        values = {'I7':'10'} if name.endswith('sheet1.xml') else {f'F{r}':'15','F26':'999','F27':'10'}
        for cell in root.findall('.//s:c',ns):
            if cell.attrib['r'] in values:
                v = cell.find('s:v',ns)
                if v is None: v=ET.SubElement(cell,'{'+ns['s']+'}v')
                v.text = values[cell.attrib['r']]
        return ET.tostring(root,encoding='utf-8')
    return rewrite(stream.getvalue(),cache)


def codes(row):
    return {x['code'] for x in row['issues'] + ((row['card'] or {}).get('issues',[]))}


def test_filtered_summary_latest_cache_and_independent_arithmetic():
    report = imports.parse_workbook(workbook(), 'synthetic.xlsx')
    assert len(report['rows']) == 1 and report['sheet_count'] == 2
    row = report['rows'][0]
    assert row['summary_balance'] == row['reference_balance'] == 10
    assert row['card']['latest_cell'] == 'F24'
    assert row['card']['latest_balance'] == row['card']['calculated_balance'] == 15
    assert len(row['card']['transactions']) == 2  # F26's cached 999 is never a transaction.
    assert {'filtered_total','summary_not_latest_row','summary_latest_mismatch','summary_arithmetic_mismatch','zero_threshold'} <= codes(row)
    assert any(i['code']=='physical_count_required' for i in report['issues'])


@pytest.mark.parametrize('settings,expected', [
    ({'unit':'Pack'}, 'unit_mismatch'),
    ({'caches':False}, 'missing_summary_cache'),
    ({'latest_formula':'=WEBSERVICE("https://invalid.example")'}, 'unsupported_balance_formula'),
    ({'summary_formula':'=SUM([outside.xlsx]Card!F27)'}, 'unsupported_summary_reference'),
    ({'opening_formula':'=10'}, 'unknown_opening'),
    ({'gap':True}, 'balance_chain_gap'),
])
def test_source_differences_and_unsupported_formulas_require_review(settings, expected):
    row = imports.parse_workbook(workbook(**settings),'synthetic.xlsx')['rows'][0]
    assert expected in codes(row)
    if settings.get('opening_formula'):
        assert row['card']['calculated_balance'] is None


def test_missing_caches_do_not_prevent_independent_literal_arithmetic():
    row = imports.parse_workbook(workbook(caches=False),'synthetic.xlsx')['rows'][0]
    assert row['card']['latest_balance'] is None
    assert row['card']['calculated_balance'] == 15


def test_error_cache_is_preserved_without_becoming_a_quantity():
    def error_cache(name,data):
        if name!='xl/worksheets/sheet2.xml':return data
        ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        root=ET.fromstring(data)
        for cell in root.findall('.//s:c',ns):
            if cell.attrib.get('r')=='F24':
                cell.set('t','e'); cell.find('s:v',ns).text='#VALUE!'
        return ET.tostring(root,encoding='utf-8')
    row=imports.parse_workbook(rewrite(workbook(),error_cache),'synthetic.xlsx')['rows'][0]
    assert row['card']['latest_balance'] is None and row['card']['calculated_balance']==15
    assert row['card']['transactions'][-1]['balance']['cached']=='#VALUE!'
    assert 'missing_or_invalid_card_cache' in codes(row)


def test_escaped_sheet_names_and_external_links_are_not_followed():
    report = imports.parse_workbook(workbook(title="Test's Card",external=True),'synthetic.xlsx')
    assert report['rows'][0]['card']['sheet'] == "Test's Card"
    assert any(i['code']=='external_links_ignored' for i in report['issues'])
    assert report['rows'][0]['card']['calculated_balance'] == 15


def test_case_insensitive_sheet_reference_and_missing_summary_name():
    row=imports.parse_workbook(workbook(summary_formula="='test card'!F27"),'synthetic.xlsx')['rows'][0]
    assert row['card']['sheet']=='Test Card'
    w=openpyxl.load_workbook(BytesIO(workbook()))
    w['Summary Sheet']['F8']='Ream';w['Summary Sheet']['I8']=3;w['Summary Sheet']['P8']=1
    out=BytesIO();w.save(out);w.close()
    report=imports.parse_workbook(out.getvalue(),'synthetic.xlsx')
    assert len(report['rows'])==2 and 'invalid_item_identity' in codes(report['rows'][1])


@pytest.mark.parametrize('kind', ['notzip','xml','dtd','utf16','macro','duplicates','oversizepart','bounds','text','extension'])
def test_invalid_or_unsafe_archives_fail_cleanly(kind):
    data=workbook(); name='synthetic.xlsx'
    if kind=='notzip': data=b'not a workbook'
    if kind=='xml': data=rewrite(data,lambda n,b: b'<' if n=='xl/workbook.xml' else b)
    if kind=='dtd': data=rewrite(data,lambda n,b: b'<!DOCTYPE x [<!ENTITY a "x">]><x>&a;</x>' if n=='xl/workbook.xml' else b)
    if kind=='utf16': data=rewrite(data,lambda n,b: '<!DOCTYPE x [<!ENTITY a "x">]><x>&a;</x>'.encode('utf-16') if n=='xl/workbook.xml' else b)
    if kind=='macro':
        out=BytesIO(data)
        with ZipFile(out,'a') as z:z.writestr('xl/vbaProject.bin',b'macro')
        data=out.getvalue()
    if kind=='duplicates':
        out=BytesIO(data)
        with pytest.warns(UserWarning):
            with ZipFile(out,'a') as z:z.writestr('xl/workbook.xml',b'<x/>')
        data=out.getvalue()
    if kind=='oversizepart': data=rewrite(data,lambda n,b:b'x'*(imports.MAX_XML+1) if n=='xl/workbook.xml' else b)
    if kind=='bounds': data=rewrite(data,lambda n,b:b.replace(b'r="A23"',b'r="AG2001"') if n.endswith('sheet2.xml') else b)
    if kind=='text':data=workbook(long_text=True)
    if kind=='extension':name='synthetic.xlsm'
    with pytest.raises(Invalid): imports.parse_workbook(data,name)


def test_workbook_limits_and_unknown_layout(monkeypatch):
    data=workbook()
    with monkeypatch.context() as patch:
        patch.setattr(imports,'MAX_SHEETS',1)
        with pytest.raises(Invalid,match='sheets'): imports.parse_workbook(data,'synthetic.xlsx')
    with monkeypatch.context() as patch:
        patch.setattr(imports,'MAX_CELLS',5)
        with pytest.raises(Invalid,match='cells'): imports.parse_workbook(data,'synthetic.xlsx')
    malformed=rewrite(data,lambda n,b:b.replace(b'ARTICLE',b'UNKNOWN') if n.endswith('sheet1.xml') else b)
    with pytest.raises(Invalid,match='summary sheet'): imports.parse_workbook(malformed,'synthetic.xlsx')
    with pytest.raises(Invalid,match='8 MiB'): imports.parse_workbook(b'x'*(imports.MAX_UPLOAD+1),'synthetic.xlsx')


def test_staged_preview_idempotency_and_inventory_isolation(supplies_db):
    c,u,_,_=supplies_db; setup_catalog(c,u); item=create(c,u[1])
    stock.post_opening(c,u[1],item['id'],{'quantity':'10','effective_date':'2026-10-19','reason':'Synthetic count'},expected_revision=1,operation_key=key())
    before={t:c.execute(f'SELECT * FROM {t}').fetchall() for t in ('supply_item','supply_receipt','supply_receipt_line','supply_stock_movement')}
    op=key(); data=workbook(unit='Pack')
    first=imports.preview_import(c,u[1],data,'synthetic.xlsx',operation_key=op)
    assert imports.preview_import(c,u[1],data,'synthetic.xlsx',operation_key=op)==first
    with pytest.raises(Conflict): imports.preview_import(c,u[1],workbook(),'synthetic.xlsx',operation_key=op)
    assert c.execute('SELECT COUNT(*) FROM supply_import_batch').fetchone()[0]==1
    preview=imports.get_preview(c,u[1],first['id'])
    assert 'unit_mismatch' in codes(preview['rows'][0]) and preview['status']=='preview'
    assert before=={t:c.execute(f'SELECT * FROM {t}').fetchall() for t in before}
    assert stock.stock_totals(c,u[2],item['id'])['on_hand']==10
    st.init_db(c)
    assert imports.get_preview(c,u[1],first['id'])['row_count']==1
    for sql in ("UPDATE supply_import_row SET raw_json='{}'",'DELETE FROM supply_import_row',
                "UPDATE supply_import_batch SET source_sha256='changed'",'DELETE FROM supply_import_batch'):
        with pytest.raises(sqlite3.IntegrityError,match='immutable'):c.execute(sql)
        c.rollback()


def test_authority_disabled_and_audit_failure(supplies_db,monkeypatch):
    c,u,_,_=supplies_db; grant=setup_catalog(c,u); data=workbook()
    result=imports.preview_import(c,u[1],data,'synthetic.xlsx',operation_key=key())
    s.grant_role(c,u[0],u[3]['id'],'auditor',operation_key=key())
    assert imports.get_preview(c,u[3],result['id'])['row_count']==1
    for who in (u[0],u[2],u[3]):
        with pytest.raises(Denied):imports.preview_import(c,who,data,'synthetic.xlsx',operation_key=key())
    with pytest.raises(Denied):imports.get_preview(c,u[2],result['id'])
    s.set_enabled(c,u[0],'0',expected_revision=1,operation_key=key())
    assert imports.get_preview(c,u[1],result['id'])['row_count']==1
    with pytest.raises(Disabled):imports.preview_import(c,u[1],data,'synthetic.xlsx',operation_key=key())
    s.set_enabled(c,u[0],'1',expected_revision=2,operation_key=key())
    def fail(*args,**kwargs):raise RuntimeError('Synthetic audit failure')
    with monkeypatch.context() as patch:
        patch.setattr(st,'audit',fail)
        with pytest.raises(RuntimeError):imports.preview_import(c,u[1],data,'synthetic.xlsx',operation_key=key())
    assert c.execute('SELECT COUNT(*) FROM supply_import_batch').fetchone()[0]==1
    s.revoke_role(c,u[0],grant['id'],expected_revision=1,operation_key=key())
    with pytest.raises(Denied):imports.get_preview(c,u[1],result['id'])


@pytest.mark.parametrize('change',['revoke','disable'])
def test_authority_rechecked_after_parsing(supplies_db,monkeypatch,change):
    c,u,_,_=supplies_db; grant=setup_catalog(c,u)
    original=imports.parse_workbook
    def changed(data,name):
        report=original(data,name)
        if change=='revoke':s.revoke_role(c,u[0],grant['id'],expected_revision=1,operation_key=key())
        else:s.set_enabled(c,u[0],'0',expected_revision=1,operation_key=key())
        return report
    monkeypatch.setattr(imports,'parse_workbook',changed)
    with pytest.raises(Denied if change=='revoke' else Disabled):
        imports.preview_import(c,u[1],workbook(),'synthetic.xlsx',operation_key=key())
    assert c.execute('SELECT COUNT(*) FROM supply_import_batch').fetchone()[0]==0


def test_preview_pagination_and_html_source_escaping(app_mod,supplies_db):
    c,u,_,path=supplies_db; setup_catalog(c,u)
    w=openpyxl.load_workbook(BytesIO(workbook()))
    summary=w['Summary Sheet']
    for r in range(7,33):
        summary.cell(r,1,'<script>bad()</script>' if r==7 else f'Synthetic {r}')
        summary.cell(r,6,'Ream');summary.cell(r,9,"='Test Card'!F27");summary.cell(r,16,5)
    out=BytesIO();w.save(out);w.close()
    result=imports.preview_import(c,u[1],out.getvalue(),'synthetic.xlsx',operation_key=key())
    assert len(imports.get_preview(c,u[1],result['id'])['rows'])==25
    assert len(imports.get_preview(c,u[1],result['id'],page=2)['rows'])==1
    old=app_mod.DB;app_mod.DB=path
    try:
        with TestClient(app_mod.app) as client:
            client.cookies.set(auth.COOKIE_NAME,session_token(u[1]))
            page=client.get(f"/supplies/import/{result['id']}")
            assert page.status_code==200 and '&lt;script&gt;bad()&lt;/script&gt;' in page.text
            assert '<script>bad()</script>' not in page.text
    finally:app_mod.DB=old


def test_chunked_upload_is_bounded_without_content_length(app_mod,supplies_db,monkeypatch):
    c,u,_,path=supplies_db;setup_catalog(c,u)
    old=app_mod.DB;app_mod.DB=path
    monkeypatch.setattr(imports,'MAX_UPLOAD',64)
    try:
        with TestClient(app_mod.app) as client:
            client.cookies.set(auth.COOKIE_NAME,session_token(u[1]))
            response=client.post('/supplies/import/preview',content=iter([b'x'*(65536+65)]),
                                 headers={'Content-Type':'multipart/form-data; boundary=synthetic'})
            assert response.status_code==422 and 'upload exceeds' in response.text
            assert c.execute('SELECT COUNT(*) FROM supply_import_batch').fetchone()[0]==0
    finally:app_mod.DB=old


def test_http_upload_preview_access_csrf_and_unreviewed_apply_rejected(app_mod,supplies_db):
    c,u,_,path=supplies_db; setup_catalog(c,u)
    old=app_mod.DB; app_mod.DB=path
    try:
        with TestClient(app_mod.app) as client:
            client.cookies.set(auth.COOKIE_NAME,session_token(u[1]))
            page=client.get('/supplies/import'); assert page.status_code==200
            token=hidden(page.text,'csrf'); op=key()
            data={'csrf':token,'operation_key':op}
            files={'workbook':('synthetic.xlsx',workbook(unit='Pack'),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
            assert client.post('/supplies/import/preview',data={**data,'csrf':'bad'},files=files).status_code==403
            assert client.post('/supplies/import/preview',data=data,files=files,headers={'Origin':'https://foreign.example'}).status_code==403
            result=client.post('/supplies/import/preview',data=data,files=files,follow_redirects=False)
            assert result.status_code==303
            target=result.headers['location']; report=client.get(target)
            assert report.status_code==200 and 'Summary and stock-card units differ' in report.text
            assert 'No inventory balances have changed' in report.text and 'Independent arithmetic' in report.text
            assert client.post('/supplies/import/preview',data=data,files=files,follow_redirects=False).headers['location']==target
            assert client.post(target+'/apply',data={**data,'operation_key':key(),'revision':'1','selection_sha256':'a'*64}).status_code==409
            invalid=client.post('/supplies/import/preview',data={**data,'operation_key':key()},files={'workbook':('bad.xlsx',b'notzip')})
            assert invalid.status_code==422 and 'readable XLSX' in invalid.text
            assert c.execute('SELECT COUNT(*) FROM supply_stock_movement').fetchone()[0]==0
            client.cookies.set(auth.COOKIE_NAME,session_token(u[2]))
            assert client.get('/supplies/import').status_code==403
            assert client.get(target).status_code==403
            assert client.post('/supplies/import/preview',data=data,files=files).status_code==403
    finally: app_mod.DB=old
