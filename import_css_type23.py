"""
import_css_type23.py — load the "CSS.xlsx" Google-Forms export into css_is.db.

The workbook has two sheets, one per survey family:

  * "Type 2"  ->  Activity / event feedback  (survey_type 2, form F2)
  * "Type 3"  ->  Resource-person feedback    (survey_type 3, form F3)

Both are the SERVQUAL capture form the app already knows: an expectation item
(A, 1-7), nine experience items (B1..B9, 1-7 with 0 = N/A), and one overall
service-quality item (C, 1-4). Sex and DOH affiliation are kept per response as
the SEX / AFFIL item_remarks — exactly what a walk-in web submission stores —
and the free-text "D" column becomes the response's overall remark. Rows are
written through storage.record_response(), so an imported response is
indistinguishable from one typed into the live form.

INTERNAL STAFF ONLY (Type 3)
----------------------------
Per the request, Type-3 feedback is imported ONLY for resource persons who are
INTERNAL staff. Each raw name in the sheet is fuzzy-matched (storage.names_match,
the same matcher the app uses) against the resource_person registry; a cluster is
imported only when it resolves to an existing person with is_internal=1. Anything
that resolves to an external guest — or to nobody at all — is SKIPPED and listed
in the report, never invented. Type-2 activity feedback is keyed by activity, not
by a resource person, so it is imported in full.

Because the classification comes from the LIVE registry, run --apply on the same
database the platform serves (pull it first with pull_backup.py if you are
working locally). If a person is missing or mis-classified, fix them in the
directory and re-run — the import is idempotent.

IDEMPOTENT
----------
Every form_instance this script creates is tagged in its qr_token:
    Type 2 -> "CSS-T2-IMPORT-<YYYY-MM>-<nn>"
    Type 3 -> "CSS-T3-IMPORT-<YYYY-MM>-RP<id>"
Re-running purges every CSS-T2/T3-IMPORT-* row (and its responses/answers/
remarks) first, then reloads — so a second run never doubles the data.

Usage
-----
    python import_css_type23.py PATH.xlsx              # DRY RUN: show the plan, write nothing
    python import_css_type23.py PATH.xlsx --apply      # take a backup, then import
    python import_css_type23.py --verify [DB]          # show what is already loaded
    python import_css_type23.py PATH.xlsx --self-test  # end-to-end check on a throwaway DB

    Options: --db DB (default css_is.db)
"""
from __future__ import annotations

import argparse
import collections
import datetime
import re
import sys

import openpyxl

import storage as st
import harmonize as h
import tzutil as tz

# ── constants ─────────────────────────────────────────────────────────────────
DEFAULT_DB = "css_is.db"
TAG_T2 = "CSS-T2-IMPORT-"
TAG_T3 = "CSS-T3-IMPORT-"
SHEET_T2 = "Type 2"
SHEET_T3 = "Type 3"

# A/B answers are on the 1-7 SERVQUAL scale; C is the 1-4 overall-quality scale.
B_ITEMS = ["B1", "B2", "B3", "B4", "B5", "B6", "B7", "B8", "B9"]
TITLE_CLUSTER_THRESHOLD = 0.72   # activity titles: looser than names (long strings)


# ── small parsers ─────────────────────────────────────────────────────────────
def parse_ts(raw):
    """('YYYY-MM-DDTHH:MM:SS', 'YYYY-MM') from a Google-Forms timestamp.

    Accepts a real datetime or the 'M/D/YYYY H:M:S' text the export uses; returns
    (None, None) if it cannot be read."""
    if isinstance(raw, datetime.datetime):
        return raw.strftime("%Y-%m-%dT%H:%M:%S"), raw.strftime("%Y-%m")
    if isinstance(raw, str):
        s = raw.strip()
        for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y",
                    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                dt = datetime.datetime.strptime(s, fmt)
                return dt.strftime("%Y-%m-%dT%H:%M:%S"), dt.strftime("%Y-%m")
            except ValueError:
                pass
    return None, None


def to_int(v):
    """Coerce a cell to an int rating, or None if blank/unreadable."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def sex_code(raw):
    """'Male (Lalaki)'/'Female (Babae)' -> 'M'/'F' (the SEX item_remark)."""
    s = (raw or "").strip().lower()
    if s.startswith("male") or "lalaki" in s:
        return "M"
    if s.startswith("female") or "babae" in s:
        return "F"
    return None


def affil_code(raw):
    """'Are you a DOH employee?' Yes/No -> 'DOH'/'NON-DOH' (the AFFIL remark)."""
    s = (raw or "").strip().lower()
    if s.startswith("yes") or "oo" in s:
        return "DOH"
    if s.startswith("no") or "hindi" in s:
        return "NON-DOH"
    return None


def clean_name(s):
    return re.sub(r"\s+", " ", (s or "").strip(" .,"))


def is_roster(name):
    """A single cell naming 2+ people (e.g. 'Ross Rosales, Neil Coroza, ...')."""
    parts = [p for p in re.split(r"\s*,\s*", name or "") if p.strip()]
    multi = [p for p in parts if len(re.findall(r"[A-Za-z]{2,}", p)) >= 2]
    return len(multi) >= 2


def split_roster(name):
    return [clean_name(p) for p in re.split(r"\s*,\s*", name or "")
            if len(re.findall(r"[A-Za-z]{2,}", p)) >= 2]


# ── generic greedy fuzzy clustering (uses the app's own matcher) ───────────────
def cluster(counter, threshold=None):
    """Group near-duplicate strings. counter = {raw: frequency}. Returns
    {representative_raw: [member_raw, ...]}, most frequent first. The
    representative is the highest-frequency variant, which becomes the canonical
    label the operator reviews."""
    reps = {}   # rep -> [members]
    for raw, _ in sorted(counter.items(), key=lambda kv: (-kv[1], kv[0])):
        placed = False
        for rep in reps:
            ok = (st.names_match(raw, rep) if threshold is None
                  else _ratio(raw, rep) >= threshold)
            if ok:
                reps[rep].append(raw)
                placed = True
                break
        if not placed:
            reps[raw] = [raw]
    return reps


def _ratio(a, b):
    import difflib
    na, nb = st._norm_name(a), st._norm_name(b)
    if not na or not nb:
        return 0.0
    return difflib.SequenceMatcher(None, na, nb).ratio()


# ── load the workbook ─────────────────────────────────────────────────────────
def _iter_rows(ws, header_row=1):
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in vals):
            continue
        yield vals


def load_type2(path):
    """Rows -> dicts. Columns (0-based): 0 ts, 4 sex, 5 doh, 7 activity,
    8 A, 9..17 B1..B9, 18 C, 21 comment."""
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET_T2]
    out = []
    for v in _iter_rows(ws):
        iso, month = parse_ts(v[0])
        out.append({
            "iso": iso, "month": month,
            "activity": (v[7] or "").strip(),
            "sex": sex_code(v[4]), "affil": affil_code(v[5]),
            "A": to_int(v[8]),
            "B": [to_int(v[9 + i]) for i in range(9)],
            "C": to_int(v[18]),
            "comment": (v[21] or "").strip() if len(v) > 21 else "",
        })
    return out


def load_type3(path):
    """Rows -> dicts. Columns (0-based): 0 ts, 3 sex, 4 doh, 6 RP name,
    7 A, 8..16 B1..B9, 17 C, 18 comment."""
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET_T3]
    out = []
    for v in _iter_rows(ws):
        iso, month = parse_ts(v[0])
        out.append({
            "iso": iso, "month": month,
            "rp_raw": clean_name(v[6]),
            "sex": sex_code(v[3]), "affil": affil_code(v[4]),
            "A": to_int(v[7]),
            "B": [to_int(v[8 + i]) for i in range(9)],
            "C": to_int(v[17]),
            "comment": (v[18] or "").strip() if len(v) > 18 else "",
        })
    return out


# ── registry matching (internal-staff filter for Type 3) ──────────────────────
def registry(conn):
    return conn.execute(
        "SELECT id, name, COALESCE(is_internal,0) FROM resource_person").fetchall()


def match_person(name, people):
    """Best resource_person match for `name`, preferring internal staff.
    Returns (id, name, is_internal, ratio) or None."""
    best = None
    for rid, rname, is_int in people:
        if not st.names_match(name, rname):
            continue
        rt = _ratio(name, rname)
        # rank: internal first, then higher ratio
        key = (is_int, rt)
        if best is None or key > best[0]:
            best = (key, (rid, rname, is_int, rt))
    return best[1] if best else None


# ── the plan (pure: computes what WOULD be written, no DB writes) ─────────────
def build_type2_plan(rows):
    """Group Type-2 rows into (canonical activity, month) form_instances."""
    titles = collections.Counter(r["activity"] or "(untitled)" for r in rows)
    clusters = cluster(titles, threshold=TITLE_CLUSTER_THRESHOLD)
    raw_to_rep = {m: rep for rep, members in clusters.items() for m in members}

    groups = collections.OrderedDict()   # (rep, month) -> [rows]
    for r in rows:
        rep = raw_to_rep[r["activity"] or "(untitled)"]
        key = (rep, r["month"] or "unknown")
        groups.setdefault(key, []).append(r)
    return clusters, groups


def build_type3_plan(rows, people):
    """Resolve each Type-3 name to an internal resource person; group the
    IMPORTABLE rows by (rp_id, month). Returns (groups, decisions) where
    decisions maps each name-cluster rep -> a human-readable outcome."""
    names = collections.Counter(r["rp_raw"] or "(blank)" for r in rows)
    clusters = cluster(names)   # names_match / 0.85
    raw_to_rep = {m: rep for rep, members in clusters.items() for m in members}

    decisions = {}      # rep -> {"action","detail","rp_id"?}
    for rep, members in clusters.items():
        n = sum(names[m] for m in members)
        if is_roster(rep):
            parts = split_roster(rep)
            hits = [match_person(p, people) for p in parts]
            internal = [m for m in hits if m and m[2] == 1]
            decisions[rep] = {
                "action": "roster", "n": n,
                "detail": f"roster of {len(parts)} -> {len(internal)} internal",
                "targets": internal}
            continue
        m = match_person(rep, people)
        if m is None:
            decisions[rep] = {"action": "skip-unmatched", "n": n,
                              "detail": "no registry match"}
        elif m[2] == 1:
            decisions[rep] = {"action": "import", "n": n, "rp_id": m[0],
                              "detail": f"-> internal #{m[0]} {m[1]!r} (r={m[3]:.2f})"}
        else:
            decisions[rep] = {"action": "skip-external", "n": n,
                              "detail": f"-> external #{m[0]} {m[1]!r} (r={m[3]:.2f})"}

    groups = collections.OrderedDict()   # (rp_id, month) -> [(row, rp_id)]
    for r in rows:
        rep = raw_to_rep[r["rp_raw"] or "(blank)"]
        d = decisions[rep]
        if d["action"] == "import":
            groups.setdefault((d["rp_id"], r["month"] or "unknown"), []).append(r)
        elif d["action"] == "roster":
            for (rid, _, _, _) in d["targets"]:
                groups.setdefault((rid, r["month"] or "unknown"), []).append(r)
    return groups, decisions


# ── writers ───────────────────────────────────────────────────────────────────
def _answers(r):
    return [("A", r["A"])] + [(B_ITEMS[i], r["B"][i]) for i in range(9)] + [("C", r["C"])]


def _item_remarks(r):
    ir = {}
    if r["sex"]:
        ir["SEX"] = r["sex"]
    if r["affil"]:
        ir["AFFIL"] = r["affil"]
    return ir


def _insert_form_instance(conn, survey_type, form, token, event_date, *,
                          topic=None, rp_id=None, lead_org_id=None):
    cur = conn.execute(
        """INSERT INTO form_instance(survey_type, form, qr_token, event_date,
               lead_org_id, rp_id, topic, created_at)
           VALUES (?,?,?,?,?,?,?,?)""",
        (survey_type, form, token, event_date, lead_org_id, rp_id, topic, tz.now_iso()))
    return cur.lastrowid


def purge(conn):
    ids = [r[0] for r in conn.execute(
        "SELECT id FROM form_instance WHERE qr_token LIKE ? OR qr_token LIKE ?",
        (TAG_T2 + "%", TAG_T3 + "%"))]
    if not ids:
        return 0, 0
    ph = ",".join("?" * len(ids))
    resp = [r[0] for r in conn.execute(
        f"SELECT id FROM response WHERE form_instance_id IN ({ph})", ids)]
    if resp:
        rp = ",".join("?" * len(resp))
        conn.execute(f"DELETE FROM answer      WHERE response_id IN ({rp})", resp)
        conn.execute(f"DELETE FROM item_remark WHERE response_id IN ({rp})", resp)
        conn.execute(f"DELETE FROM event_form  WHERE form_instance_id IN ({ph})", ids)
        conn.execute(f"DELETE FROM response    WHERE id IN ({rp})", resp)
    conn.execute(f"DELETE FROM form_instance WHERE id IN ({ph})", ids)
    conn.commit()
    return len(ids), len(resp)


def _min_date(group):
    ds = sorted(r["iso"] for r in group if r["iso"])
    return (ds[0][:10] if ds else tz.now_iso()[:10])


def apply_import(conn, t2_rows, t3_rows):
    lead_org_id = st.get_or_create_blhsd(conn)[0]
    people = registry(conn)

    _, t2_groups = build_type2_plan(t2_rows)
    t3_groups, t3_decisions = build_type3_plan(t3_rows, people)

    fi_purged, resp_purged = purge(conn)

    counts = {"t2_fi": 0, "t2_resp": 0, "t3_fi": 0, "t3_resp": 0}

    # Type 2 — one form_instance per (activity, month)
    for i, ((rep, month), group) in enumerate(t2_groups.items(), start=1):
        token = f"{TAG_T2}{month}-{i:02d}"
        topic = None if rep == "(untitled)" else rep
        fi = _insert_form_instance(conn, 2, "F2", token, _min_date(group),
                                   topic=topic, lead_org_id=lead_org_id)
        counts["t2_fi"] += 1
        for r in group:
            st.record_response(conn, fi, _answers(r), submitted_at=r["iso"],
                               remarks=r["comment"], item_remarks=_item_remarks(r))
            counts["t2_resp"] += 1

    # Type 3 — one form_instance per (internal rp, month)
    for (rp_id, month), group in t3_groups.items():
        token = f"{TAG_T3}{month}-RP{rp_id}"
        fi = _insert_form_instance(conn, 3, "F3", token, _min_date(group),
                                   rp_id=rp_id, lead_org_id=lead_org_id)
        counts["t3_fi"] += 1
        for r in group:
            st.record_response(conn, fi, _answers(r), submitted_at=r["iso"],
                               remarks=r["comment"], item_remarks=_item_remarks(r))
            counts["t3_resp"] += 1

    st.audit(conn, "(import:css_type23)", "bulk_import",
             target_type="form_instance", target_id="CSS-T2/T3",
             summary=(f"Imported CSS.xlsx: T2 {counts['t2_resp']} responses in "
                      f"{counts['t2_fi']} activity form(s); T3 {counts['t3_resp']} "
                      f"internal-staff responses in {counts['t3_fi']} speaker form(s). "
                      f"(purged {fi_purged} prior import form(s)/{resp_purged} resp.)"))
    conn.commit()
    return counts, t3_decisions


# ── reporting ─────────────────────────────────────────────────────────────────
def print_plan(conn, t2_rows, t3_rows):
    people = registry(conn)
    t2_clusters, t2_groups = build_type2_plan(t2_rows)
    t3_groups, t3_decisions = build_type3_plan(t3_rows, people)

    print(f"\nRegistry: {sum(1 for _,_,i in people if i==1)} internal / "
          f"{sum(1 for _,_,i in people if i==0)} external resource persons on file.")

    print(f"\n=== TYPE 2 — activity feedback ({len(t2_rows)} responses) ===")
    print(f"{len(t2_clusters)} activity/-ies after clustering; "
          f"{len(t2_groups)} form_instance(s) will be created:")
    titles = collections.Counter(r["activity"] or "(untitled)" for r in t2_rows)
    for rep, members in t2_clusters.items():
        n = sum(titles[m] for m in members)
        print(f"   {n:3d}  {rep!r}"
              + (f"   (+{len(members)-1} spelling variant(s))" if len(members) > 1 else ""))

    print(f"\n=== TYPE 3 — resource-person feedback ({len(t3_rows)} responses) ===")
    imp = [d for d in t3_decisions.values() if d["action"] == "import"]
    n_imp = sum(len(g) for g in t3_groups.values())
    print(f"IMPORT (internal staff): {len(imp)} person(s), "
          f"{n_imp} response(s) in {len(t3_groups)} form_instance(s)")
    for rep, d in sorted(t3_decisions.items(), key=lambda kv: -kv[1]["n"]):
        tag = {"import": "IMPORT ", "skip-external": "SKIP ext",
               "skip-unmatched": "SKIP  ??", "roster": "ROSTER "}[d["action"]]
        print(f"   [{tag}] {d['n']:3d}  {rep!r}  {d['detail']}")
    skipped = sum(d["n"] for d in t3_decisions.values()
                  if d["action"].startswith("skip"))
    print(f"\nType-3 responses skipped (external / unmatched): {skipped}")


def verify(conn):
    for tag, label in ((TAG_T2, "Type 2"), (TAG_T3, "Type 3")):
        fi = conn.execute("SELECT COUNT(*) FROM form_instance WHERE qr_token LIKE ?",
                          (tag + "%",)).fetchone()[0]
        resp = conn.execute(
            "SELECT COUNT(*) FROM response r JOIN form_instance fi "
            "ON r.form_instance_id=fi.id WHERE fi.qr_token LIKE ?",
            (tag + "%",)).fetchone()[0]
        ans = conn.execute(
            "SELECT COUNT(*) FROM answer a JOIN response r ON a.response_id=r.id "
            "JOIN form_instance fi ON r.form_instance_id=fi.id WHERE fi.qr_token LIKE ?",
            (tag + "%",)).fetchone()[0]
        print(f"{label}: form_instances={fi}  responses={resp}  answers={ans}")


# ── self-test (throwaway DB; proves the internal-only rule end-to-end) ─────────
def self_test(path):
    import tempfile, os
    tmp = tempfile.mktemp(suffix=".db")
    conn = st.connect(tmp)
    st.init_db(conn)
    st.get_or_create_blhsd(conn)
    # seed the registry: some of the sheet's people INTERNAL, one EXTERNAL
    st.add_rp(conn, "Maria Teresa Santiago", is_internal=True)
    st.add_rp(conn, "RayJustin Ventura", is_internal=True)
    st.add_rp(conn, "Dr. Ross Jacob Rosales", is_internal=True)
    st.add_rp(conn, "Vlanche April Dulfo", is_internal=False)   # external -> must skip
    t2 = load_type2(path)
    t3 = load_type3(path)
    print_plan(conn, t2, t3)
    counts, decisions = apply_import(conn, t2, t3)
    print("\nself-test counts:", counts)
    verify(conn)
    # assertions
    ext = decisions.get("Vlanche April Dulfo")
    assert ext and ext["action"] == "skip-external", "external RP was not skipped!"
    santiago = next((d for k, d in decisions.items()
                     if k == "Maria Teresa Santiago"), None)
    assert santiago and santiago["action"] == "import", "internal RP not imported!"
    # every imported answer belongs to an internal person
    bad = conn.execute("""
        SELECT COUNT(*) FROM form_instance fi
        JOIN resource_person rp ON rp.id = fi.rp_id
        WHERE fi.qr_token LIKE ? AND COALESCE(rp.is_internal,0)=0""",
        (TAG_T3 + "%",)).fetchone()[0]
    assert bad == 0, "external speaker form_instance was created!"
    assert counts["t2_resp"] == len(t2), "not all Type-2 rows imported"
    # idempotency: a second run yields identical counts
    counts2, _ = apply_import(conn, t2, t3)
    assert counts2 == counts, "import is not idempotent!"
    conn.close()
    os.remove(tmp)
    print("\nSELF-TEST PASSED ✔  (external skipped, internal imported, idempotent)")


# ── entry point ───────────────────────────────────────────────────────────────
def main(argv=None):
    ap = argparse.ArgumentParser(description="Import CSS.xlsx (Type 2 + Type 3).")
    ap.add_argument("xlsx", nargs="?", help="path to the CSS workbook")
    ap.add_argument("--db", default=DEFAULT_DB)
    ap.add_argument("--apply", action="store_true", help="write to the DB (default: dry run)")
    ap.add_argument("--verify", action="store_true", help="show already-loaded import counts")
    ap.add_argument("--self-test", action="store_true", help="run the end-to-end check")
    args = ap.parse_args(argv)

    if args.self_test:
        if not args.xlsx:
            ap.error("--self-test needs the xlsx path")
        return self_test(args.xlsx)

    if args.verify:
        verify(st.connect(args.db))
        return

    if not args.xlsx:
        ap.error("give the path to CSS.xlsx (or use --verify)")

    t2 = load_type2(args.xlsx)
    t3 = load_type3(args.xlsx)
    conn = st.connect(args.db)
    print_plan(conn, t2, t3)

    if not args.apply:
        print("\nDRY RUN — nothing written. Re-run with --apply on the live DB to import.")
        return

    bk = h.backup_db(args.db, tag="pre_css_type23_import")
    counts, _ = apply_import(conn, t2, t3)
    print(f"\nBackup: {bk}")
    print("Import complete:", counts)
    verify(conn)


if __name__ == "__main__":
    main()
