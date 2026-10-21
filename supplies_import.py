"""Bounded stock-workbook evidence preview. Never evaluates formulas or posts stock."""
from datetime import date, datetime
import hashlib
from io import BytesIO
import json
import math
from pathlib import PurePosixPath
import re
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile
import zlib

import openpyxl

import supplies
import supplies_permissions as permissions
from supplies_storage import Invalid, Missing, operation, record, rows, transaction
import tzutil as tz


MAX_UPLOAD = 8 * 1024 * 1024
MAX_EXPANDED = 32 * 1024 * 1024
MAX_XML = 4 * 1024 * 1024
MAX_ROWS, MAX_COLS, MAX_SHEETS, MAX_CELLS = 2000, 32, 150, 300000
MAX_ITEMS = 500
_REF = re.compile(r"^=(?:'((?:[^']|'')+)'|([A-Za-z_][\w.]*))!\$?F\$?([1-9][0-9]*)$", re.I)


def _plain(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _norm(value):
    return ' '.join(str(value or '').split()).casefold()


def _number(value):
    if type(value) is int and 0 <= value <= 2147483647:
        return value
    if type(value) is float and value.is_integer() and 0 <= value <= 2147483647:
        return int(value)
    return None


def _issue(issues, code, message):
    entry = {'code': code, 'message': message}
    if entry not in issues:
        issues.append(entry)


def _preflight(data):
    if not isinstance(data, bytes) or not 0 < len(data) <= MAX_UPLOAD:
        raise Invalid('Choose a nonempty XLSX stock workbook no larger than 8 MiB.')
    links = False
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            names = [m.filename for m in members]
            if len(members) > 1024 or len(set(names)) != len(names) or sum(m.file_size for m in members) > MAX_EXPANDED:
                raise Invalid('The workbook archive exceeds the supported limits or has duplicate members.')
            if not {'[Content_Types].xml', 'xl/workbook.xml'} <= set(names):
                raise Invalid('This is not a supported XLSX workbook.')
            cells = 0
            for member in members:
                name = member.filename
                if '..' in PurePosixPath(name).parts or name.startswith('/') or member.flag_bits & 1:
                    raise Invalid('Encrypted or invalid workbook archive members are not supported.')
                if 'vbaproject' in name.lower():
                    raise Invalid('Macro-enabled workbooks are not supported. Supply an XLSX copy.')
                if member.file_size > MAX_XML:
                    raise Invalid('A workbook archive part exceeds the supported size limit.')
                xml = archive.read(member)
                is_xml = name.lower().endswith(('.xml', '.rels')) or xml.lstrip(b'\xef\xbb\xbf \r\n\t').startswith(b'<')
                if xml.startswith((b'\xff\xfe', b'\xfe\xff')):
                    raise Invalid('Workbook XML must use UTF-8 encoding.')
                if not is_xml:
                    continue
                if b'\x00' in xml:
                    raise Invalid('Workbook XML must use UTF-8 encoding.')
                lowered = xml.lower()
                if b'<!doctype' in lowered or b'<!entity' in lowered:
                    raise Invalid('Workbook XML declarations with document types or entities are not supported.')
                if name == '[Content_Types].xml' and (b'macroenabled' in lowered or b'vbaproject' in lowered):
                    raise Invalid('Macro-enabled workbooks are not supported. Supply an XLSX copy.')
                root = ET.fromstring(xml)
                for node in root.iter():
                    tag = node.tag.rsplit('}', 1)[-1]
                    if tag == 'Relationship' and node.attrib.get('TargetMode', '').lower() == 'external':
                        links = True
                    if tag in ('t', 'f') and len(node.text or '') > 4000:
                        raise Invalid('Workbook cell text or formulas exceed 4,000 characters.')
                    if tag == 'c' and root.tag.rsplit('}', 1)[-1] == 'worksheet':
                        cells += 1
                        address = node.attrib.get('r', '')
                        match = re.fullmatch(r'([A-Z]+)([1-9][0-9]*)', address)
                        if not match or len(match[2]) > 4:
                            raise Invalid('Workbook cell addresses exceed the supported worksheet bounds.')
                        column = openpyxl.utils.column_index_from_string(match[1])
                        if column > MAX_COLS or int(match[2]) > MAX_ROWS:
                            raise Invalid('Worksheets must fit within 2,000 rows and 32 columns.')
            if cells > MAX_CELLS:
                raise Invalid('The workbook has too many populated cells.')
    except (BadZipFile, ET.ParseError, RuntimeError, KeyError, ValueError, zlib.error) as exc:
        if isinstance(exc, Invalid):
            raise
        raise Invalid('The file is not a readable XLSX workbook.') from exc
    return links


def _cell(sheet, cached, row, col):
    cell = sheet.cell(row, col)
    return {'cell': cell.coordinate, 'value': _plain(cell.value),
            'cached': _plain(cached.cell(row, col).value), 'formula': cell.data_type == 'f'}


def _numeric(cell):
    return _number(cell['cached'])


def _date_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        try:
            return date.fromisoformat(value).isoformat()
        except ValueError:
            pass
    return None


def _card(sheet, cached):
    issues = []
    result = {'sheet': sheet.title, 'name': None, 'unit': None, 'latest_balance': None,
              'calculated_balance': None, 'latest_cell': None, 'last_transaction_date': None,
              'transactions': [], 'issues': issues}
    header = None
    for r in range(1, min(sheet.max_row, 50) + 1):
        a = str(sheet.cell(r, 1).value or '').strip()
        if a.upper().startswith('NAME OF ITEM:'):
            result['name'] = a.split(':', 1)[1].strip()
        if re.match(r'^UNIT\s*:', a, re.I):
            result['unit'] = a.split(':', 1)[1].strip()
        if [_norm(sheet.cell(r, c).value) for c in (1, 2, 4, 6)] == ['date', 'received', 'issued', 'balance']:
            header = r
            break
    if not header:
        _issue(issues, 'unknown_card_layout', 'The stock-card date/received/issued/balance columns are not recognized.')
        return result
    footer = next((r for r in range(header + 1, sheet.max_row + 1)
                   if _norm(sheet.cell(r, 1).value) == 'total balance'), sheet.max_row + 1)
    result['footer_row'] = footer if footer <= sheet.max_row else None
    if footer <= sheet.max_row:
        result['footer_cells'] = [_cell(sheet, cached, footer, c) for c in (2, 4, 6)]
        _issue(issues, 'total_row_reference', 'The card has a separate total row; it is not a dated transaction balance.')
        if any('SUMIF' in str(sheet.cell(footer, c).value).upper() for c in (2, 4, 6)):
            _issue(issues, 'filtered_total', 'The total row uses a conditional subtotal and may cover a selected period instead of all transactions.')
    calculated, trustworthy, previous = None, True, None
    for r in range(header + 1, footer):
        raw_date = sheet.cell(r, 1).value
        received, issued, balance = [_cell(sheet, cached, r, c) for c in (2, 4, 6)]
        # Ignore prefilled balance formulas on otherwise blank rows, even if they have a cached value.
        if all(v in (None, '') for v in (raw_date, received['value'], issued['value'])) and (balance['formula'] or balance['value'] in (None, '')):
            continue
        day = _date_value(raw_date)
        tx = {'row': r, 'date': _plain(raw_date), 'received': received, 'issued': issued, 'balance': balance}
        result['transactions'].append(tx)
        if not day:
            trustworthy = False
            _issue(issues, 'invalid_transaction_date', f'Row {r} has movement/balance data without a recognized date.')
        else:
            result['last_transaction_date'] = day
        amounts = []
        for field in (received, issued):
            amount = 0 if field['value'] in (None, '') else _number(field['value'])
            if field['formula'] or amount is None:
                trustworthy = False
                _issue(issues, 'invalid_movement', f'Row {r} has a received/issued value that is not a literal nonnegative whole quantity.')
            amounts.append(amount or 0)
        if previous is None:
            if balance['formula'] or _number(balance['value']) is None:
                trustworthy = False
                _issue(issues, 'unknown_opening', 'The first dated balance must be a literal opening quantity to calculate stock independently.')
            calculated = _number(balance['value'])
            if any(amounts):
                trustworthy = False
                _issue(issues, 'ambiguous_opening', 'The first balance also contains received/issued quantities; its opening basis needs review.')
        elif calculated is not None:
            calculated += amounts[0] - amounts[1]
            if not 0 <= calculated <= 2147483647:
                trustworthy = False
                _issue(issues, 'invalid_calculated_balance', f'Independent arithmetic reaches an unsupported or negative balance at row {r}.')
        if balance['formula']:
            compact = re.sub(r'\s+', '', str(balance['value'])).upper()
            expected = f'=IF(A{r}="","",SUM(F{r-1}+B{r}-D{r}))'
            if compact != expected:
                _issue(issues, 'unsupported_balance_formula', f'Balance formula in F{r} is not the recognized running-balance pattern; it was not evaluated.')
            elif previous is not None and previous != r - 1:
                _issue(issues, 'balance_chain_gap', f'F{r} refers to a preceding blank transaction row; its cached balance may be broken.')
        last = _numeric(balance)
        if last is None:
            _issue(issues, 'missing_or_invalid_card_cache', f'F{r} has no usable cached whole-number balance.')
        elif trustworthy and calculated is not None and last != calculated:
            _issue(issues, 'card_cache_mismatch', f'F{r} cached balance differs from independent receipt/issue arithmetic.')
        result['latest_balance'], result['latest_cell'] = last, balance['cell']
        previous = r
    if not result['transactions']:
        _issue(issues, 'no_transactions', 'No dated opening/transaction evidence was found on this card.')
    result['calculated_balance'] = calculated if trustworthy else None
    return result


def parse_workbook(data, filename):
    if not isinstance(filename, str) or not filename.lower().endswith('.xlsx'):
        raise Invalid('Supply an XLSX workbook; macros and legacy XLS files are not supported.')
    external = _preflight(data)
    books = []
    try:
        for cached in (False, True):
            books.append(openpyxl.load_workbook(BytesIO(data), data_only=cached, keep_links=False))
        formulas, values = books
        if len(formulas.worksheets) > MAX_SHEETS:
            raise Invalid('A stock workbook may contain at most 150 sheets.')
        if any(s.max_row > MAX_ROWS or s.max_column > MAX_COLS for s in formulas):
            raise Invalid('Worksheets must fit within 2,000 rows and 32 columns.')
        if sum(s.max_row * s.max_column for s in formulas) > MAX_CELLS:
            raise Invalid('The workbook worksheet area exceeds the supported cell limit.')
        for sheet in formulas:
            for row in sheet:
                if any(isinstance(c.value, str) and len(c.value) > 4000 for c in row):
                    raise Invalid('Workbook cell text exceeds 4,000 characters.')
        candidates = []
        for sheet in formulas:
            for r in range(1, min(sheet.max_row, 10) + 1):
                if [_norm(sheet.cell(r, c).value) for c in (1, 6, 9)] == ['article', 'unit of measure', 'balance']:
                    candidates.append((sheet, r))
        if len(candidates) != 1:
            raise Invalid('Expected one summary sheet with ARTICLE, UNIT OF MEASURE and BALANCE in columns A, F and I.')
        summary, heading = candidates[0]
        cached_summary = values[summary.title]
        report = {'source_sha256': hashlib.sha256(data).hexdigest(), 'sheet_count': len(formulas.worksheets),
                  'summary_sheet': summary.title, 'issues': [], 'rows': []}
        if external:
            _issue(report['issues'], 'external_links_ignored', 'External workbook links/hyperlinks were ignored and never followed.')
        _issue(report['issues'], 'physical_count_required', 'Saved workbook values may reflect an earlier period. Check against a dated physical count before adopting any opening stock.')
        cards, used, seen_names = {}, set(), set()
        for r in range(heading + 1, summary.max_row + 1):
            if all(summary.cell(r, c).value in (None, '') for c in (1,6,9)):
                continue
            issues = []
            name, unit, balance, critical = [_cell(summary, cached_summary, r, c) for c in (1, 6, 9, 16)]
            row = {'source_sheet': summary.title, 'source_row': r, 'name': name['value'], 'unit': unit['value'],
                   'summary_balance': _numeric(balance), 'critical_level': _numeric(critical),
                   'source_cells': [name, unit, balance, critical], 'issues': issues, 'card': None,
                   'reference_balance': None}
            if not isinstance(name['value'], str) or name['formula'] or not isinstance(unit['value'], str) or not unit['value'].strip() or unit['formula']:
                _issue(issues, 'invalid_item_identity', 'Item name/unit must be literal nonempty text before mapping it to the catalog.')
            if _norm(name['value']) in seen_names:
                _issue(issues, 'duplicate_summary_name', 'This name occurs more than once in the summary; do not merge the rows automatically.')
            seen_names.add(_norm(name['value']))
            if row['summary_balance'] is None:
                _issue(issues, 'missing_summary_cache', 'Summary balance has no usable cached nonnegative whole quantity.')
            if row['critical_level'] is None or critical['formula']:
                _issue(issues, 'invalid_critical_level', 'Critical level needs a reviewed literal nonnegative whole value.')
            if row['critical_level'] == 0:
                _issue(issues, 'zero_threshold', 'Critical level is zero: positive low-stock warnings are disabled; zero stock will still be marked Out of stock.')
            match = _REF.fullmatch(str(balance['value']).strip()) if balance['formula'] else None
            target = None
            if match:
                target = (match[1].replace("''", "'") if match[1] is not None else match[2], int(match[3]))
                if target[1] > MAX_ROWS:
                    target = None
            if target is None:
                _issue(issues, 'unsupported_summary_reference', 'Summary balance is not a direct local stock-card F-cell reference; no formula was evaluated.')
                matches = [s.title for s in formulas if s.title != summary.title and _norm(s.title) == _norm(row['name'])]
                if len(matches) == 1:
                    target = (matches[0], None)
            if target:
                names = [name for name in formulas.sheetnames if name.casefold() == target[0].casefold()]
                if len(names) == 1:
                    target = (names[0], target[1])
            if target and target[0] in formulas.sheetnames and target[0] != summary.title:
                title, ref_row = target
                used.add(title)
                if title not in cards:
                    cards[title] = _card(formulas[title], values[title])
                row['card'] = cards[title]
                if ref_row:
                    reference = _cell(formulas[title], values[title], ref_row, 6)
                    row['reference_cell'] = reference
                    row['reference_balance'] = _numeric(reference)
                    if row['reference_balance'] != row['summary_balance']:
                        _issue(issues, 'summary_reference_mismatch', 'The summary cache and referenced card-cell cache disagree.')
                    if cards[title]['latest_cell'] and reference['cell'] != cards[title]['latest_cell']:
                        _issue(issues, 'summary_not_latest_row', 'The summary points to a different row from the last meaningful dated card row.')
                if _norm(unit['value']) != _norm(cards[title]['unit']):
                    _issue(issues, 'unit_mismatch', 'Summary and stock-card units differ. Confirm the unit of measure against the physical count.')
                if _norm(name['value']) != _norm(cards[title]['name']):
                    _issue(issues, 'name_mismatch', 'Summary and card item descriptions differ; confirm that they describe the same supply.')
                for field, code in [('latest_balance', 'summary_latest_mismatch'), ('calculated_balance', 'summary_arithmetic_mismatch')]:
                    if cards[title][field] is not None and cards[title][field] != row['summary_balance']:
                        _issue(issues, code, 'The summary balance differs from ' + ('the latest dated card cache.' if field == 'latest_balance' else 'independent opening + receipts − issues arithmetic.'))
            else:
                _issue(issues, 'missing_stock_card', 'No unambiguous referenced stock card was found.')
            report['rows'].append(row)
            if len(report['rows']) > MAX_ITEMS:
                raise Invalid('The preview supports at most 500 summary item rows.')
        if not report['rows']:
            raise Invalid('The summary sheet contains no item rows.')
        report['unmapped_sheets'] = [s.title for s in formulas if s.title != summary.title and s.title not in used]
        if report['unmapped_sheets']:
            _issue(report['issues'], 'unmapped_sheets', 'Some sheets are not associated with summary rows; review them separately.')
        if len(json.dumps(report).encode()) > MAX_UPLOAD:
            raise Invalid('The preview evidence exceeds the supported size limit.')
        return report
    except Invalid:
        raise
    except (ValueError, TypeError, KeyError, IndexError, AttributeError, OverflowError, OSError,
            BadZipFile, ET.ParseError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise Invalid('The workbook could not be read as the supported stock-card layout.') from exc
    finally:
        for book in books:
            book.close()


def preview_import(conn, user, data, filename, *, operation_key):
    permissions.require_custodian(conn, user)
    supplies.require_enabled(conn)
    filename = str(filename or '').replace('\\', '/').rsplit('/', 1)[-1]
    filename = supplies.text(filename, 'Workbook name', 240)
    report = parse_workbook(data, filename)
    with transaction(conn):
        user = permissions.require_custodian(conn, user)
        supplies.require_enabled(conn)
        payload = {'source_sha256': report['source_sha256'], 'filename': filename}
        key, digest, prior = operation(conn, user, operation_key, 'import_preview', payload)
        if prior is not None:
            return prior
        metadata = {k:v for k,v in report.items() if k != 'rows'}
        batch_id = conn.execute('''INSERT INTO supply_import_batch
          (source_sha256,filename,sheet_count,metadata_json,created_by,created_at,operation_key)
          VALUES (?,?,?,?,?,?,?)''', (report['source_sha256'], filename, report['sheet_count'],
           json.dumps(metadata), user['id'], tz.now_iso(), key)).lastrowid
        for row in report['rows']:
            conn.execute('INSERT INTO supply_import_row(batch_id,source_sheet,source_row,raw_json) VALUES (?,?,?,?)',
                         (batch_id, row['source_sheet'], row['source_row'], json.dumps(row)))
        result = {'id': batch_id, 'source_sha256': report['source_sha256'], 'row_count': len(report['rows'])}
        return record(conn, user, key, digest, 'import_preview', 'import', result, result)


def get_preview(conn, user, batch_id, *, page=1):
    permissions.require_stock_reader(conn, user)
    page = supplies.integer(page, 'Page', 1)
    found = rows(conn, 'SELECT * FROM supply_import_batch WHERE id=?', (batch_id,))
    if not found:
        raise Missing('Workbook preview not found or unavailable to you.')
    result = found[0]
    result['metadata'] = json.loads(result.pop('metadata_json'))
    result['row_count'] = conn.execute('SELECT COUNT(*) FROM supply_import_row WHERE batch_id=?', (batch_id,)).fetchone()[0]
    result['rows'] = [dict(json.loads(r['raw_json']), staging_id=r['id']) for r in rows(conn,
        'SELECT id,raw_json FROM supply_import_row WHERE batch_id=? ORDER BY id LIMIT 25 OFFSET ?', (batch_id, (page-1)*25))]
    result['page'], result['has_next'] = page, page*25 < result['row_count']
    return result


def list_previews(conn, user, *, page=1):
    permissions.require_stock_reader(conn, user)
    page = supplies.integer(page, 'Page', 1)
    found = rows(conn, 'SELECT id,filename,created_at,status FROM supply_import_batch ORDER BY id DESC LIMIT 26 OFFSET ?', ((page-1)*25,))
    return {'batches': found[:25], 'page': page, 'has_next': len(found)>25}
