"""Editable Word drafts based on a permitted outgoing document record."""
from io import BytesIO
from datetime import date
from docx import Document
from docx.shared import Inches, Pt


KINDS = {'internal': 'Internal memorandum', 'response': 'Response memorandum'}


def build(doc, fields):
    kind = str(fields.get('kind') or '')
    if kind not in KINDS:
        raise ValueError('Choose an available memorandum template.')
    values = {}
    for name, limit in [('to', 500), ('from', 500), ('subject', 500),
                        ('reference', 500), ('body', 12000), ('prepared_by', 200)]:
        value = str(fields.get(name) or '').strip()
        if len(value) > limit:
            raise ValueError(f'{name.replace("_", " ").title()} is too long.')
        values[name] = value
    if not values['to'] or not values['from'] or not values['subject'] or not values['body']:
        raise ValueError('Complete To, From, Subject, and Body before downloading.')
    try:
        memo_date = date.fromisoformat(str(fields.get('date') or '')).isoformat()
    except ValueError:
        raise ValueError('Choose a valid memorandum date.') from None
    paper = Document()
    section = paper.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    normal = paper.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11)
    paper.add_heading(KINDS[kind].upper(), 0)
    paper.add_paragraph('Working draft for review', style='Subtitle')
    for label, value in [('To', values['to']), ('From', values['from']),
                         ('Date', memo_date), ('Subject', values['subject']),
                         ('Reference', values['reference'] or doc['number'])]:
        paragraph = paper.add_paragraph()
        paragraph.add_run(label + ': ').bold = True
        paragraph.add_run(value)
    paper.add_paragraph('')
    for block in values['body'].split('\n'):
        paper.add_paragraph(block)
    if values['prepared_by']:
        paper.add_paragraph('')
        paragraph = paper.add_paragraph()
        paragraph.add_run('Prepared by: ').bold = True
        paragraph.add_run(values['prepared_by'])
    paper.add_paragraph('Document record: ' + doc['number'])
    paper.add_paragraph('This draft is not an approval, signature, or dispatch.')
    output = BytesIO()
    paper.save(output)
    return output.getvalue()
