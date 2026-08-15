"""
ANNEX C — Internal Work Arrangement Schedule, in the two forms the bureau
submits: an Excel workbook (a sheet per half-month period, the shape of the
file it has always circulated) and a print-ready PDF.

Both are built from the same storage.work_arrangement_annex_data() payload and
render the SAME masthead, the same six columns, the same verbatim legend and the
same Prepared-by / Approved-by block — which is why they live in one module
rather than one each. A workbook and a printout of the same period that disagree
would be worse than either alone.

Only APPROVED arrangements carry a value; anyone pending, denied or who never
filed appears with the Work Arrangement and Remarks cells blank. That is
faithful to the real form: the blanks are part of the record, and the system
surfaces them on screen before the deadline rather than hiding them here.

This module is pure formatting — it holds no SQL and makes no decisions about
what belongs on the form.
"""
from __future__ import annotations

import io
from xml.sax.saxutils import escape

# The six columns of ANNEX C, in the order the source sheets use.
HEADERS = ["No.", "Name", "Position", "Division/Unit", "Work Arrangement",
           "Remarks"]
# Column widths (Excel character units) and the PDF's proportional split.
_XLSX_WIDTHS = [5, 30, 34, 20, 18, 30]
_PDF_WEIGHTS = [4, 24, 26, 17, 13, 22]

_TITLE = "ANNEX C"
_SUBTITLE = "INTERNAL WORK ARRANGEMENT SCHEDULE"

# The house table treatment, matched to the wellness register workbook so the
# bureau's two generated forms look like they came from the same system.
_HEADER_FILL = "1F4E79"
_BORDER_RGB = "9AA1AC"


def period_title(data):
    """'August 16-31, 2026' — also the worksheet tab name."""
    return data.get("period_label") or data.get("period_start") or "ANNEX C"


def filename(period_start, ext, month=None):
    """A safe download name: ANNEX-C-Work-Arrangement-<period|month>.<ext>."""
    key = month or period_start or "schedule"
    return f"ANNEX-C-Work-Arrangement-{key}.{ext}"


# ---------------------------------------------------------------- Excel ------

def _sheet(ws, data):
    """Lay one period out on one worksheet: masthead, headers, body, legend,
    signatures."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color=_BORDER_RGB)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ncols = len(HEADERS)

    masthead = [
        (_TITLE, Font(bold=True, size=12)),
        (_SUBTITLE, Font(bold=True, size=14)),
        (data["office_unit"], Font(bold=True, size=11)),
        (period_title(data), Font(size=11)),
        ("", None),
    ]
    for i, (text, font) in enumerate(masthead, start=1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ncols)
        c = ws.cell(i, 1, text)
        if font:
            c.font = font
        c.alignment = ctr

    hrow = len(masthead) + 1
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(hrow, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = ctr
        c.border = border

    r = hrow + 1
    for row in data["rows"]:
        cells = [row["no"], row["name"], row["position"], row["division_unit"],
                 row["arrangement"], row["remarks"]]
        for col, val in enumerate(cells, start=1):
            c = ws.cell(r, col, val)
            c.alignment = ctr if col in (1, 5) else left
            c.border = border
        r += 1
    if not data["rows"]:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.cell(r, 1, "No staff on this period's schedule.").alignment = ctr
        r += 1

    lrow = r + 1
    ws.merge_cells(start_row=lrow, start_column=1, end_row=lrow, end_column=ncols)
    c = ws.cell(lrow, 1, data["legend"])
    c.font = Font(size=9, italic=True, color="404756")
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Signature block: two columns, "Prepared by" left and "Approved by" right,
    # each a label / NAME / title stack like the printed form.
    prepared, approved = data["prepared_by"], data["approved_by"]
    srow = lrow + 2
    blocks = [
        ("Prepared by:", "Approved by:", None),
        ((prepared["name"] or "").upper(), (approved["name"] or "").upper(),
         Font(bold=True)),
        (prepared["title"], approved["title"], Font(size=9)),
    ]
    for i, (a, b, font) in enumerate(blocks):
        ws.merge_cells(start_row=srow + i, start_column=1,
                       end_row=srow + i, end_column=3)
        ws.merge_cells(start_row=srow + i, start_column=4,
                       end_row=srow + i, end_column=ncols)
        for col, val in ((1, a), (4, b)):
            cell = ws.cell(srow + i, col, val)
            cell.alignment = left
            if font:
                cell.font = font

    for i, w in enumerate(_XLSX_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{hrow + 1}"
    ws.print_title_rows = f"{hrow}:{hrow}"


def render_xlsx(periods):
    """One workbook, one worksheet per period — so a whole month goes out as the
    two-sheet file the bureau already circulates. `periods` is a list of
    storage.work_arrangement_annex_data() dicts."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for data in periods:
        # Excel caps tab names at 31 chars and forbids []:*?/\ — the period
        # labels are short and clean, but truncate defensively.
        ws = wb.create_sheet(title=period_title(data)[:31])
        _sheet(ws, data)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF ------

def render_pdf(periods):
    """The same annex, print-ready. A4 LANDSCAPE — six columns do not fit
    portrait, which is where this diverges from wellness_register.py's layout.
    One page (or more, if a period is long) per period."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    page = landscape(A4)
    title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=11,
                           leading=14, alignment=TA_CENTER)
    subtitle = ParagraphStyle("s", fontName="Helvetica-Bold", fontSize=13,
                              leading=16, alignment=TA_CENTER)
    office = ParagraphStyle("o", fontName="Helvetica-Bold", fontSize=10,
                            leading=13, alignment=TA_CENTER)
    period = ParagraphStyle("p", fontName="Helvetica", fontSize=10, leading=13,
                            alignment=TA_CENTER)
    cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=10)
    cellc = ParagraphStyle("cc", parent=cell, alignment=TA_CENTER)
    head = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=8, leading=10,
                          alignment=TA_CENTER, textColor=colors.white)
    small = ParagraphStyle("sm", fontName="Helvetica-Oblique", fontSize=7.5,
                           leading=10, textColor=colors.Color(.25, .28, .34))
    sig = ParagraphStyle("sig", fontName="Helvetica", fontSize=9, leading=13)

    usable = page[0] - 30 * mm
    widths = [usable * w / sum(_PDF_WEIGHTS) for w in _PDF_WEIGHTS]

    story = []
    for n, data in enumerate(periods):
        if n:
            story.append(PageBreak())
        story += [
            Paragraph(_TITLE, title),
            Paragraph(_SUBTITLE, subtitle),
            Paragraph(escape(data["office_unit"]), office),
            Paragraph(escape(period_title(data)), period),
            Spacer(1, 8),
        ]
        table_data = [[Paragraph(escape(h), head) for h in HEADERS]]
        for row in data["rows"]:
            table_data.append([
                Paragraph(str(row["no"]), cellc),
                Paragraph(escape(row["name"] or ""), cell),
                Paragraph(escape(row["position"] or ""), cell),
                Paragraph(escape(row["division_unit"] or ""), cell),
                Paragraph(escape(row["arrangement"] or ""), cellc),
                Paragraph(escape(row["remarks"] or ""), cell),
            ])
        if not data["rows"]:
            table_data.append([Paragraph("No staff on this period's schedule.",
                                         cellc)] + [""] * (len(HEADERS) - 1))
        table = Table(table_data, colWidths=widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#" + _BORDER_RGB)),
            ("BOX", (0, 0), (-1, -1), 1.0, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + _HEADER_FILL)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story += [table, Spacer(1, 6),
                  Paragraph(escape(data["legend"]), small), Spacer(1, 14)]

        prepared, approved = data["prepared_by"], data["approved_by"]

        def _stack(label, who):
            name = escape((who["name"] or "").upper())
            title_ = escape(who["title"] or "")
            return Paragraph(
                f"{label}<br/><br/><b>{name}</b><br/>{title_}", sig)

        sig_table = Table(
            [[_stack("Prepared by:", prepared), _stack("Approved by:", approved)]],
            colWidths=[usable / 2.0] * 2)
        sig_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ]))
        story.append(sig_table)

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=page, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm,
        title=f"ANNEX C — {period_title(periods[0]) if periods else ''}",
    ).build(story)
    return buf.getvalue()
