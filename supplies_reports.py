"""Read-only Supplies reports. Historical filters use posting date, never backdated effective dates."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import supplies as s
import supplies_permissions as permissions
import supplies_stock as stock
import tzutil as tz
from supplies_storage import Invalid, rows


KINDS = {
    'inventory': 'Current inventory',
    'fulfillment': 'Current request fulfillment',
    'consumption': 'Issue consumption by posting date',
    'replenishment': 'Current outstanding replenishment',
    'stock-card': 'Stock card',
}
HEADERS = {
    'inventory': [('code', 'Item code'), ('description', 'Description'), ('unit', 'Unit'),
                  ('on_hand', 'On hand'), ('reserved', 'Reserved'), ('available', 'Available'),
                  ('incoming', 'Approved incoming'), ('critical_level', 'Critical level')],
    'fulfillment': [('control_number', 'Request'), ('requester_snapshot', 'Requester'),
                    ('unit_snapshot', 'Office'), ('status', 'Status'), ('requested', 'Requested'),
                    ('approved', 'Approved'), ('issued', 'Net issued'), ('cancelled', 'Cancelled'),
                    ('reserved', 'Reserved'), ('awaiting', 'Awaiting stock')],
    'consumption': [('posted_date', 'Posting date'), ('effective_date', 'Effective date'),
                    ('code', 'Item code'), ('description', 'Description'), ('unit', 'Unit'),
                    ('release_number', 'Release'), ('request_number', 'Request'),
                    ('issued', 'Issued'), ('reversed', 'Reversed'), ('net_issued', 'Net issued'),
                    ('usable_returned', 'Usable returned'), ('unusable_returned', 'Unusable returned')],
    'replenishment': [('control_number', 'Replenishment'), ('status', 'Status'),
                      ('code', 'Item code'), ('description', 'Description'), ('unit', 'Unit'),
                      ('approved', 'Approved'), ('net_received', 'Net accepted'),
                      ('cancelled', 'Cancelled'), ('incoming', 'Outstanding incoming'),
                      ('estimated_unit_cost', 'Draft estimated unit cost')],
    'stock-card': [('movement_id', 'Posting ID'), ('posted_at', 'Posted at'),
                   ('effective_date', 'Effective date'), ('kind', 'Transaction'),
                   ('quantity_delta', 'Change'), ('balance', 'Running balance'),
                   ('reference', 'Reference'), ('supplier', 'Supplier'), ('dr_number', 'DR number'),
                   ('source_unit_cost', 'Receipt source unit cost'), ('actor_snapshot', 'Recorded by'),
                   ('reason', 'Reason')],
}


def _as_of(value):
    if value in ('', None):
        return None
    day = stock._date(value, 'Posting date')
    if day > tz.today_local().isoformat():
        raise Invalid('Posting date cannot be in the future.')
    return day


def report(conn, user, kind, *, as_of=None, item_id=None):
    if kind not in KINDS:
        raise Invalid('Choose a supported Supplies report.')
    day = _as_of(as_of)
    if day and kind not in ('stock-card', 'consumption'):
        raise Invalid('Posting-date filters apply only to stock cards and issue consumption; other reports show the current snapshot.')
    with stock._snapshot(conn):
        permissions.require_stock_reader(conn, user)
        item = None
        if kind == 'stock-card':
            if item_id in ('', None):
                raise Invalid('Choose a stock-card item.')
            item = s.get_item(conn, user, s.integer(item_id, 'Item', 1))
        elif item_id not in ('', None):
            raise Invalid('Item selection applies only to a stock card.')
        if kind == 'inventory':
            data = rows(conn, '''SELECT i.id AS item_id,i.code,i.description,i.unit,i.critical_level,
                COALESCE(m.on_hand,0) AS on_hand,COALESCE(r.reserved,0) AS reserved,
                COALESCE(m.on_hand,0)-COALESCE(r.reserved,0) AS available,
                COALESCE(c.incoming,0) AS incoming
                FROM supply_item i
                LEFT JOIN (SELECT item_id,SUM(quantity_delta) on_hand FROM supply_stock_movement GROUP BY item_id) m ON m.item_id=i.id
                LEFT JOIN (SELECT item_id,SUM(reserved_quantity) reserved FROM supply_reservation GROUP BY item_id) r ON r.item_id=i.id
                LEFT JOIN (SELECT l.item_id,SUM(c.approved_quantity-c.received_quantity-c.cancelled_quantity) incoming
                  FROM supply_replenishment_commitment c JOIN supply_replenishment_line l ON l.id=c.line_id
                  JOIN supply_replenishment p ON p.id=l.replenishment_id
                  WHERE p.status IN ('approved','ordered','partially_received') GROUP BY l.item_id) c ON c.item_id=i.id
                ORDER BY i.code,i.id''')
        elif kind == 'fulfillment':
            data = rows(conn, '''SELECT q.id AS request_id,q.control_number,q.requester_snapshot,q.unit_snapshot,
                q.approval_status AS status,COALESCE(SUM(l.requested_quantity),0) requested,
                COALESCE(SUM(d.approved_quantity),0) approved,
                COALESCE(SUM(r.effective_issued_quantity),0) issued,
                COALESCE(SUM(r.cancelled_quantity),0) cancelled,
                COALESCE(SUM(r.reserved_quantity),0) reserved,
                COALESCE(SUM(d.approved_quantity-r.effective_issued_quantity-r.cancelled_quantity-r.reserved_quantity),0) awaiting
                FROM supply_request q JOIN supply_request_revision v
                  ON v.request_id=q.id AND v.revision=q.current_revision
                JOIN supply_request_revision_line l ON l.revision_id=v.id
                LEFT JOIN supply_request_decision_line d ON d.revision_line_id=l.id
                LEFT JOIN supply_reservation r ON r.revision_line_id=l.id
                GROUP BY q.id ORDER BY q.id''')
        elif kind == 'consumption':
            data = rows(conn, '''SELECT i.id AS issue_id,i.release_number,q.control_number AS request_number,
                substr(i.posted_at,1,10) AS posted_date,i.effective_date,
                l.code_snapshot AS code,l.description_snapshot AS description,l.unit_snapshot AS unit,
                l.issued_quantity AS issued,
                COALESCE(c.reversed,0) AS reversed,
                l.issued_quantity-COALESCE(c.reversed,0) AS net_issued,
                COALESCE(c.usable_returned,0) AS usable_returned,
                COALESCE(c.unusable_returned,0) AS unusable_returned
                FROM supply_issue_line l JOIN supply_issue i ON i.id=l.issue_id
                JOIN supply_request q ON q.id=i.request_id
                LEFT JOIN (SELECT issue_line_id,
                  SUM(CASE WHEN kind='issue_reversal' THEN quantity ELSE 0 END) reversed,
                  SUM(CASE WHEN kind='return' AND condition='usable' THEN quantity ELSE 0 END) usable_returned,
                  SUM(CASE WHEN kind='return' AND condition='unusable' THEN quantity ELSE 0 END) unusable_returned
                  FROM supply_correction WHERE (? IS NULL OR substr(posted_at,1,10)<=?)
                  GROUP BY issue_line_id) c ON c.issue_line_id=l.id
                WHERE (? IS NULL OR substr(i.posted_at,1,10)<=?)
                ORDER BY i.posted_at,i.id,l.id''', (day, day, day, day))
        elif kind == 'replenishment':
            data = rows(conn, '''SELECT p.id AS replenishment_id,p.control_number,p.status,
                l.code_snapshot AS code,l.description_snapshot AS description,l.unit_snapshot AS unit,
                c.approved_quantity AS approved,c.received_quantity AS net_received,
                c.cancelled_quantity AS cancelled,
                c.approved_quantity-c.received_quantity-c.cancelled_quantity AS incoming,
                l.estimated_unit_cost
                FROM supply_replenishment_commitment c JOIN supply_replenishment_line l ON l.id=c.line_id
                JOIN supply_replenishment p ON p.id=l.replenishment_id
                WHERE p.status IN ('approved','ordered','partially_received')
                  AND c.approved_quantity>c.received_quantity+c.cancelled_quantity
                ORDER BY p.id,l.id''')
        else:
            data = rows(conn, '''SELECT m.id AS movement_id,m.posted_at,m.effective_date,m.kind,
                m.quantity_delta,SUM(m.quantity_delta) OVER (ORDER BY m.id ROWS UNBOUNDED PRECEDING) AS balance,
                m.actor_snapshot,m.reason,r.supplier,r.dr_number,l.unit_cost AS source_unit_cost,
                CASE WHEN m.kind='receipt' THEN 'RCV-'||printf('%06d',m.source_id)
                  WHEN m.kind='issue' THEN COALESCE(i.release_number,'Issue '||m.source_id)
                  WHEN m.kind IN ('return','adjustment','reversal') THEN 'Correction '||m.source_id
                  ELSE 'Opening count' END AS reference
                FROM supply_stock_movement m
                LEFT JOIN supply_receipt_line l ON l.id=m.receipt_line_id
                LEFT JOIN supply_receipt r ON r.id=l.receipt_id
                LEFT JOIN supply_issue i ON i.id=m.source_id AND m.kind='issue'
                WHERE m.item_id=? AND (? IS NULL OR substr(m.posted_at,1,10)<=?)
                ORDER BY m.id''', (item['id'], day, day))
        summary = []
        if kind == 'consumption':
            grouped = {}
            for entry in data:
                group = (entry['posted_date'], entry['code'], entry['description'], entry['unit'])
                if group not in grouped:
                    grouped[group] = {'posted_date': group[0], 'code': group[1],
                                      'description': group[2], 'unit': group[3],
                                      'issued': 0, 'reversed': 0, 'net_issued': 0,
                                      'usable_returned': 0, 'unusable_returned': 0}
                for name in ('issued', 'reversed', 'net_issued', 'usable_returned', 'unusable_returned'):
                    grouped[group][name] += entry[name]
            summary = [grouped[group] for group in sorted(grouped)]
        return {'kind': kind, 'title': KINDS[kind], 'headers': HEADERS[kind], 'rows': data,
                'summary': summary,
                'as_of': day, 'item': item, 'generated_at': tz.now_iso(),
                'date_note': ('Posting date through ' + day if day else 'All saved postings') if kind in ('stock-card','consumption') else 'Current saved snapshot'}


def _safe_cell(value):
    if value is None:
        return ''
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def xlsx_bytes(result):
    book = Workbook()
    sheet = book.active
    sheet.title = result['title'][:31]
    sheet.append([_safe_cell(result['title'])])
    sheet.append([result['date_note']])
    sheet.append(['Generated at', result['generated_at']])
    sheet.append(['Source costs are recorded receipt/draft estimates, not official accounting valuation.'])
    sheet.append([])
    sheet.append([label for _, label in result['headers']])
    for cell in sheet[6]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='294764')
    for entry in result['rows']:
        sheet.append([_safe_cell(entry.get(name)) for name, _ in result['headers']])
    sheet.freeze_panes = 'A7'
    sheet.auto_filter.ref = f'A6:{sheet.cell(sheet.max_row, len(result["headers"])).coordinate}'
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = min(45, max(12, max(len(str(c.value or '')) for c in col) + 2))
    if result['kind'] == 'consumption':
        summary = book.create_sheet('By item and posting date')
        names = ('posted_date', 'code', 'description', 'unit', 'issued', 'reversed',
                 'net_issued', 'usable_returned', 'unusable_returned')
        summary.append(['Posting date', 'Item code', 'Description', 'Unit', 'Issued',
                        'Reversed', 'Net issued', 'Usable returned', 'Unusable returned'])
        for entry in result['summary']:
            summary.append([_safe_cell(entry[name]) for name in names])
        summary.freeze_panes = 'A2'
        summary.auto_filter.ref = summary.dimensions
        for col in summary.columns:
            summary.column_dimensions[col[0].column_letter].width = min(45, max(12, max(len(str(c.value or '')) for c in col) + 2))
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
